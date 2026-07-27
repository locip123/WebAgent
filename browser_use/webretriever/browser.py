"""Playwright-only browser runtime for the WebRetriever challenge.

The competition supplies an already-running browser over CDP.  Connection
ownership deliberately stays with the runner; :class:`BrowserRuntime` owns
only the pages it creates inside the supplied ``BrowserContext``.
"""

from __future__ import annotations

import asyncio
import base64
import contextlib
import copy
import hashlib
import json
import logging
import math
import os
import re
import secrets
import stat
import tempfile
import time
import unicodedata
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable, Coroutine, Literal, Mapping, Sequence, cast
from urllib.parse import parse_qsl, unquote, urlsplit
from xml.etree import ElementTree

from playwright.async_api import BrowserContext, Download, Frame, Locator, Page, Request, Response
from playwright.async_api import Error as PlaywrightError
from playwright.async_api import TimeoutError as PlaywrightTimeoutError
from pydantic import BaseModel, ConfigDict

if TYPE_CHECKING:
	from browser_use.webretriever.models import AgentDecision


__all__ = [
	'BrowserObservation',
	'BrowserRuntime',
	'ElementRef',
	'cdp_headers_for_url',
	'is_sec_url',
	'is_forbidden_search_url',
	'redact_cdp_url',
]


_CDP_TOKEN_KEY = 'access_token'
_REDACTED = '<redacted>'
_INTERACTIVE_ATTRIBUTE = 'data-webretriever-index'
_OVERLAY_CLASS = '__webretriever_overlay'
_MAX_ELEMENT_TEXT = 240
_MAX_PAGE_TEXT = 80_000
_MAX_RENDERED_TEXT = 100_000
_MAX_RESPONSE_BODY_BYTES = 128 * 1024
_NETWORK_SEARCH_TIMEOUT_SECONDS = 30.0
_NETWORK_SEARCH_NODE_HEAP_MIB = 512
_NETWORK_BODY_PAGE_CHARACTERS = 60_000
_MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
_MAX_DOWNLOAD_TEXT = 250_000
_MAX_ARCHIVE_FILES = 1_000
_DOWNLOAD_PLACEHOLDER_URL = ':'
_MAX_ARCHIVE_MEMBER_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_TOTAL_BYTES = 250 * 1024 * 1024
_MAX_ARCHIVE_WARNINGS = 100
_PNG_SIGNATURE = b'\x89PNG\r\n\x1a\n'
_DOCUMENT_MIME_EXTENSIONS = {
	'application/pdf': '.pdf',
	'application/vnd.openxmlformats-officedocument.wordprocessingml.document': '.docx',
	'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': '.xlsx',
	'application/vnd.ms-excel': '.xls',
	'text/csv': '.csv',
	'application/csv': '.csv',
}
_DOCUMENT_EXTENSIONS = frozenset(_DOCUMENT_MIME_EXTENSIONS.values())


class _ScreenshotFallbackError(RuntimeError):
	"""Raised when Playwright screenshot timeout recovery through CDP also fails."""


class _ArchiveResourceLimitError(RuntimeError):
	"""Raised when a ZIP member exceeds an extraction resource budget."""


class _ExtractedArchiveMember(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	archive_name: str
	relative_name: str
	path: Path
	size_bytes: int


class _ArchiveExtractionResult(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	status: Literal['success', 'partial', 'failed']
	members: list[_ExtractedArchiveMember]
	skipped_count: int
	warnings: list[str]


class _ArchiveMemberDownload(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	timestamp: float
	url: str
	suggested_filename: str
	filename: str
	path: str
	size_bytes: int
	source: Literal['archive_member']
	source_archive: str
	text: str
	text_truncated: bool
	text_extraction_error: str | None = None


def _consume_detached_task_result(task: asyncio.Future[Any]) -> None:
	"""Retrieve a late cancellation-resistant result without warning."""

	if task.cancelled():
		return
	try:
		task.exception()
	except BaseException:
		pass


def cdp_headers_for_url(cdp_url: str) -> dict[str, str]:
	"""Return the cloud-sandbox authentication header encoded in a CDP URL."""

	try:
		query = parse_qsl(urlsplit(cdp_url).query, keep_blank_values=True)
	except (TypeError, ValueError):
		return {}
	for key, value in query:
		if key == _CDP_TOKEN_KEY and value:
			return {'X-Access-Token': value}
	return {}


_TOKEN_RE = re.compile(r'([?&]access_token=)[^&#]*', re.IGNORECASE)


def redact_cdp_url(cdp_url: str) -> str:
	"""Redact CDP credentials without otherwise rewriting the URL."""

	return _TOKEN_RE.sub(rf'\1{_REDACTED}', str(cdp_url))


def _host_matches(host: str, domain: str) -> bool:
	return host == domain or host.endswith(f'.{domain}')


def is_sec_url(url: str) -> bool:
	"""Whether *url* belongs to SEC.gov or one of its subdomains."""

	if not isinstance(url, str) or not url.strip():
		return False
	candidate = url.strip()
	if '://' not in candidate and not candidate.startswith('//'):
		candidate = f'//{candidate}'
	try:
		host = (urlsplit(candidate).hostname or '').rstrip('.').lower()
	except ValueError:
		return False
	return _host_matches(host, 'sec.gov')


def is_forbidden_search_url(url: str) -> bool:
	"""Whether *url* targets a known public web-search/answer engine.

	The match is hostname based (never substring based), so a site URL such as
	``example.test/?next=google.com`` is not incorrectly rejected.
	"""

	if not isinstance(url, str) or not url.strip():
		return False
	candidate = url.strip()
	if '://' not in candidate and not candidate.startswith('//'):
		candidate = f'//{candidate}'
	try:
		parsed = urlsplit(candidate)
		host = (parsed.hostname or '').rstrip('.').lower()
	except ValueError:
		return False
	if not host:
		return False

	# Google properties such as Drive and Maps are not search engines.  Match
	# Google's country-domain search frontends only when their URL is search-like.
	if re.fullmatch(r'(?:www\.)?google\.[a-z.]{2,}', host):
		path = parsed.path.rstrip('/').lower() or '/'
		if path in {'/', '/search', '/webhp', '/custom', '/m'} or 'q=' in parsed.query.lower():
			return True

	blocked_domains = (
		'bing.com',
		'duckduckgo.com',
		'baidu.com',
		'yandex.com',
		'yandex.ru',
		'ask.com',
		'ecosia.org',
		'qwant.com',
		'startpage.com',
		'dogpile.com',
		'mojeek.com',
		'search.brave.com',
		'perplexity.ai',
		'you.com',
		'so.com',
		'sm.cn',
	)
	if any(_host_matches(host, domain) for domain in blocked_domains):
		return True
	if re.fullmatch(r'(?:search|www)\.yahoo\.[a-z.]{2,}', host):
		if host.startswith('search.') or parsed.path.lower().startswith(('/search', '/s')) or parsed.query:
			return True
	if _host_matches(host, 'search.aol.com') or (_host_matches(host, 'aol.com') and parsed.path.lower().startswith('/search')):
		return True
	if _host_matches(host, 'search.naver.com') or _host_matches(host, 'search.sogou.com'):
		return True
	if _host_matches(host, 'sogou.com') and parsed.path.lower().startswith(('/web', '/search')):
		return True
	return False


@dataclass(slots=True)
class ElementRef:
	"""A visible interactive element indexed in the latest observation."""

	index: int
	tag: str
	text: str = ''
	role: str = ''
	name: str = ''
	placeholder: str = ''
	href: str = ''
	input_type: str = ''
	frame_index: int = 0
	frame_url: str = ''
	x: float = 0.0
	y: float = 0.0
	width: float = 0.0
	height: float = 0.0
	selector: str = ''

	def render_text(self) -> str:
		parts = [f'[{self.index}]', self.tag]
		if self.role:
			parts.append(f'role={self.role}')
		if self.input_type:
			parts.append(f'type={self.input_type}')
		if self.name:
			parts.append(f'name={json.dumps(self.name, ensure_ascii=False)}')
		if self.text and self.text != self.name:
			parts.append(f'text={json.dumps(self.text, ensure_ascii=False)}')
		if self.placeholder:
			parts.append(f'placeholder={json.dumps(self.placeholder, ensure_ascii=False)}')
		if self.href:
			parts.append(f'href={json.dumps(self.href, ensure_ascii=False)}')
		if self.frame_index:
			parts.append(f'frame={self.frame_index}')
		parts.append(f'box=({self.x:.0f},{self.y:.0f},{self.width:.0f},{self.height:.0f})')
		return ' '.join(parts)

	def to_dict(self) -> dict[str, Any]:
		return asdict(self)


@dataclass(slots=True)
class BrowserObservation:
	"""Multimodal state made available to the decision model at one step."""

	screenshot: bytes
	url: str
	title: str
	tabs: list[dict[str, Any]]
	viewport_width: int
	viewport_height: int
	elements: list[ElementRef]
	page_text: str
	recent_network: list[dict[str, Any]]
	downloads: list[dict[str, Any]]
	step: int = 0
	screenshot_path: str = ''
	visual_screenshot_path: str = ''

	def render_text(self) -> str:
		"""Render bounded, model-friendly browser state (image bytes excluded)."""

		def clip(value: Any, limit: int) -> str:
			text = str(value)
			return text if len(text) <= limit else text[: limit - 1] + '…'

		def section(name: str, lines: list[str], budget: int) -> str:
			body = '\n'.join(lines) or '  (none)'
			return f'{name}:\n{clip(body, budget)}'

		tab_lines = [
			f'  {tab["index"]}: {clip(tab.get("title", ""), 300)!r} {clip(tab.get("url", ""), 1_500)}'
			+ (' [active]' if tab.get('active') else '')
			for tab in self.tabs
		]
		element_lines = [f'  {clip(element.render_text(), 1_200)}' for element in self.elements]
		network_lines = []
		for item in self.recent_network[-12:]:
			status = item.get('status', 'pending')
			network_lines.append(f'  {clip(item.get("method", ""), 20)} {clip(item.get("url", ""), 1_500)} [{status}]')
		archive_summaries = [item for item in self.downloads if 'extraction_status' in item][-4:]
		summary_paths = {str(item.get('path', '')) for item in archive_summaries}
		other_downloads = [item for item in self.downloads if str(item.get('path', '')) not in summary_paths]
		download_items = [*archive_summaries, *other_downloads[-(8 - len(archive_summaries)) :]]
		download_items.sort(key=lambda item: float(item.get('timestamp', 0)))
		download_lines = [
			f'  {clip(item.get("filename", item.get("suggested_filename", "download")), 300)}: {clip(item.get("text", ""), 800)}'
			for item in download_items
		]
		# Put compact, high-value evidence before potentially very long page text,
		# then budget the page section independently so downloads/network cannot be
		# cut off merely because a page has a huge body.
		sections = [
			f'URL: {clip(self.url, 4_000)}',
			f'Title: {clip(self.title, 1_000)}',
			f'Viewport: {self.viewport_width}x{self.viewport_height}',
			section('Tabs', tab_lines, 8_000),
			section('Interactive elements', element_lines, 40_000),
			section('Recent XHR/Fetch', network_lines, 14_000),
			section('Downloads', download_lines, 14_000),
		]
		prefix = '\n\n'.join(sections)
		page_budget = max(0, _MAX_RENDERED_TEXT - len(prefix) - len('\n\nPage text:\n'))
		return (prefix + '\n\nPage text:\n' + self.page_text[:page_budget])[:_MAX_RENDERED_TEXT]


@dataclass(slots=True)
class _ElementBinding:
	frame: Frame
	selector: str


class _VisibleTextParser(HTMLParser):
	def __init__(self) -> None:
		super().__init__()
		self.parts: list[str] = []
		self._ignored_depth = 0

	def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
		if tag in {'script', 'style', 'noscript'}:
			self._ignored_depth += 1

	def handle_endtag(self, tag: str) -> None:
		if tag in {'script', 'style', 'noscript'} and self._ignored_depth:
			self._ignored_depth -= 1

	def handle_data(self, data: str) -> None:
		if not self._ignored_depth and data.strip():
			self.parts.append(data.strip())


_MARK_ELEMENTS_JS = r"""
({start, markerAttribute, overlayClass, frameIndex, maxText}) => {
  document.querySelectorAll('.' + overlayClass).forEach((node) => node.remove());
  document.querySelectorAll('[' + markerAttribute + ']').forEach((node) => node.removeAttribute(markerAttribute));
  const selector = [
    'a[href]', 'button', 'input:not([type="hidden"])', 'textarea', 'select',
    '[contenteditable="true"]', 'summary', 'label[for]',
    '[role="button"]', '[role="link"]', '[role="checkbox"]', '[role="radio"]',
    '[role="menuitem"]', '[role="option"]', '[role="tab"]', '[role="switch"]',
    '[role="textbox"]', '[role="combobox"]', '[tabindex]:not([tabindex="-1"])'
  ].join(',');
  const candidates = Array.from(document.querySelectorAll(selector));
  const results = [];
  let next = start;
  for (const element of candidates) {
    if (!(element instanceof HTMLElement) && !(element instanceof SVGElement)) continue;
    if (element.matches(':disabled,[aria-disabled="true"],[inert]')) continue;
    const style = getComputedStyle(element);
    const rect = element.getBoundingClientRect();
    if (style.display === 'none' || style.visibility === 'hidden' || Number(style.opacity) === 0) continue;
    if (rect.width < 2 || rect.height < 2) continue;
    if (rect.bottom <= 0 || rect.right <= 0 || rect.top >= innerHeight || rect.left >= innerWidth) continue;
    if (typeof element.checkVisibility === 'function' && !element.checkVisibility({checkOpacity: true, checkVisibilityCSS: true})) continue;
    const index = next++;
    element.setAttribute(markerAttribute, String(index));
    const text = (element.innerText || element.value || element.getAttribute('aria-label') || '').replace(/\s+/g, ' ').trim();
    const name = (element.getAttribute('aria-label') || element.getAttribute('title') || text).replace(/\s+/g, ' ').trim();
    const overlay = document.createElement('div');
    overlay.className = overlayClass;
    Object.assign(overlay.style, {
      position: 'fixed', left: `${Math.max(0, rect.left)}px`, top: `${Math.max(0, rect.top)}px`,
      width: `${Math.max(1, Math.min(rect.width, innerWidth - Math.max(0, rect.left)))}px`,
      height: `${Math.max(1, Math.min(rect.height, innerHeight - Math.max(0, rect.top)))}px`,
      border: '2px solid #ff2d55', background: 'rgba(255,45,85,.04)', boxSizing: 'border-box',
      pointerEvents: 'none', zIndex: '2147483646'
    });
    const badge = document.createElement('span');
    badge.textContent = String(index);
    Object.assign(badge.style, {
      position: 'absolute', left: '-2px', top: '-18px', background: '#ff2d55', color: '#fff',
      font: 'bold 12px/16px sans-serif', padding: '0 3px', minWidth: '16px', textAlign: 'center'
    });
    overlay.appendChild(badge);
    (document.documentElement || document.body).appendChild(overlay);
    results.push({
      index, tag: element.tagName.toLowerCase(), text: text.slice(0, maxText),
      role: element.getAttribute('role') || '', name: name.slice(0, maxText),
      placeholder: (element.getAttribute('placeholder') || '').slice(0, maxText),
      href: String(element.href || '').slice(0, 2000), input_type: element.getAttribute('type') || '', frame_index: frameIndex,
      x: rect.x, y: rect.y, width: rect.width, height: rect.height
    });
  }
  return results;
}
"""


_CLEAR_MARKERS_JS = r"""
({markerAttribute, overlayClass, removeAttributes}) => {
  document.querySelectorAll('.' + overlayClass).forEach((node) => node.remove());
  if (removeAttributes) {
    document.querySelectorAll('[' + markerAttribute + ']').forEach((node) => node.removeAttribute(markerAttribute));
  }
}
"""


class BrowserRuntime:
	"""A small async runtime that performs every browser operation with Playwright."""

	def __init__(
		self,
		context: BrowserContext,
		task_dir: Path,
		logger: logging.Logger,
		*,
		navigation_timeout_ms: int = 60_000,
		action_timeout_ms: int = 30_000,
		screenshot_timeout_ms: int = 20_000,
		cdp_screenshot_timeout_ms: int = 20_000,
		max_response_body_bytes: int = _MAX_RESPONSE_BODY_BYTES,
		declared_user_agent: str | None = None,
	) -> None:
		self.context = context
		self.task_dir = Path(task_dir)
		self.logger = logger
		self.navigation_timeout_ms = navigation_timeout_ms
		self.action_timeout_ms = action_timeout_ms
		self.screenshot_timeout_ms = screenshot_timeout_ms
		self.cdp_screenshot_timeout_ms = cdp_screenshot_timeout_ms
		self.max_response_body_bytes = max(0, max_response_body_bytes)
		self.declared_user_agent = declared_user_agent

		self.page: Page | None = None
		self.website = ''
		self.visited_urls: list[str] = []
		self.all_requests: list[dict[str, Any]] = []
		# ``all_requests`` is the competition-compatible XHR/Fetch capture.  Keep a
		# separate all-resource cache for chart-request discovery so broad capture
		# does not change the evaluator-facing capture.json contract.
		self.network_requests: list[dict[str, Any]] = []
		self.downloads: list[dict[str, Any]] = []

		self._owned_pages: list[Page] = []
		self._page_handlers: dict[int, list[tuple[str, Callable[..., Any]]]] = {}
		self._context_handlers: list[tuple[str, Callable[..., Any]]] = []
		self._request_entries: dict[int, dict[str, Any]] = {}
		self._request_entries_by_id: dict[int, dict[str, Any]] = {}
		self._request_ids_by_entry: dict[int, int] = {}
		self._network_request_entries: dict[int, dict[str, Any]] = {}
		self._network_entries_by_id: dict[int, dict[str, Any]] = {}
		self._network_responses: dict[int, Response] = {}
		self._network_page_cursors: dict[str, tuple[int, int, str, int]] = {}
		self._next_network_request_id = 0
		self._page_ids: dict[int, int] = {}
		self._page_document_generations: dict[int, int] = {}
		self._element_bindings: dict[int, _ElementBinding] = {}
		self._last_safe_urls: dict[int, str] = {}
		self._background_tasks: set[asyncio.Task[Any]] = set()
		self._download_tasks: set[asyncio.Task[Any]] = set()
		self._policy_tasks: set[asyncio.Task[Any]] = set()
		self._screenshot_recovery_tasks: set[asyncio.Task[Any]] = set()
		self._rollback_pages: set[int] = set()
		self._captured_document_responses: set[tuple[str, int]] = set()
		self._download_urls_seen: set[str] = set()
		self._reserved_download_paths: set[Path] = set()
		self._started = False
		self._closed = False
		self._execute_lock = asyncio.Lock()
		self._archive_extraction_lock = asyncio.Lock()

		self.trajectory_dir = self.task_dir / 'trajectory'
		self.trajectory_visual_dir = self.task_dir / 'trajectory_visual'
		self.download_dir = self.task_dir / 'downloads'
		self.trajectory_dir.mkdir(parents=True, exist_ok=True)
		self.trajectory_visual_dir.mkdir(parents=True, exist_ok=True)
		self.download_dir.mkdir(parents=True, exist_ok=True)

	async def __aenter__(self) -> BrowserRuntime:
		return self

	async def __aexit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
		await self.close()

	async def start(self, website: str) -> Page:
		"""Create an owned page and navigate to the exact task start URL."""

		if self._closed:
			raise RuntimeError('BrowserRuntime is closed')
		if self._started:
			raise RuntimeError('BrowserRuntime.start() may only be called once')
		self._validate_navigation_url(website)
		if is_forbidden_search_url(website):
			raise ValueError('The designated start URL is a prohibited external search engine')

		self.website = website
		self._install_context_listeners()
		page = await self.context.new_page()
		self._register_page(page, make_active=True)
		await self._configure_owned_page(page)
		self._started = True
		self.logger.info('Opening exact task start URL: %s', redact_cdp_url(website))
		download_started = await self._goto_exact(page, website)
		self._record_url(website if download_started else page.url, unless_last=True)
		if not is_forbidden_search_url(page.url):
			self._last_safe_urls[id(page)] = page.url
		await self._enforce_search_policy()
		return page

	async def observe(self, step: int) -> BrowserObservation:
		"""Capture raw/annotated screenshots plus a cross-frame textual state."""

		self._ensure_started()
		await self._enforce_search_policy()
		await self._drain_download_tasks()
		page = await self._active_page_for_observation()
		await self._clear_markers(remove_attributes=True)

		step_name = self._safe_step_name(step)
		raw_path = self.trajectory_dir / f'{step_name}.png'
		visual_path = self.trajectory_visual_dir / f'{step_name}.png'
		raw_screenshot = await self._capture_screenshot(page, raw_path)

		page_text = await self._collect_page_text(page)
		elements = await self._collect_elements(page)
		try:
			screenshot = await self._capture_screenshot(page, visual_path)
		except _ScreenshotFallbackError:
			raise
		except Exception:
			self.logger.exception('Could not capture annotated screenshot for step %s', step_name)
			screenshot = raw_screenshot

		viewport = await self._viewport(page)
		tabs = await self._tabs()
		title = await self._page_title(page)
		return BrowserObservation(
			screenshot=screenshot,
			url=page.url,
			title=title,
			tabs=tabs,
			viewport_width=viewport['width'],
			viewport_height=viewport['height'],
			elements=elements,
			page_text=page_text,
			recent_network=copy.deepcopy(self.all_requests[-20:]),
			downloads=copy.deepcopy(self.downloads),
			step=int(step),
			screenshot_path=str(raw_path),
			visual_screenshot_path=str(visual_path),
		)

	async def _capture_screenshot(self, page: Page, path: Path) -> bytes:
		try:
			return await page.screenshot(
				path=str(path),
				type='png',
				animations='disabled',
				timeout=self.screenshot_timeout_ms,
			)
		except PlaywrightTimeoutError as screenshot_timeout:
			self.logger.warning('Playwright screenshot timed out for %s; falling back to CDP capture', path.name)
			try:
				screenshot = await self._capture_cdp_screenshot_before_deadline(page, path)
			except Exception as fallback_error:
				path.unlink(missing_ok=True)
				raise _ScreenshotFallbackError(
					f'Playwright screenshot timed out after {self.screenshot_timeout_ms}ms '
					f'({screenshot_timeout}) and CDP fallback failed: '
					f'{type(fallback_error).__name__}: {fallback_error}'
				) from fallback_error
			self.logger.info('Recovered screenshot %s through CDP capture', path.name)
			return screenshot

	async def _capture_cdp_screenshot_before_deadline(self, page: Page, path: Path) -> bytes:
		expired = asyncio.Event()
		task = asyncio.create_task(self._capture_cdp_screenshot(page, path, expired))
		timeout_seconds = self.cdp_screenshot_timeout_ms / 1000
		try:
			done, _ = await asyncio.wait({task}, timeout=timeout_seconds)
		except BaseException:
			expired.set()
			if not task.done():
				self._track_screenshot_recovery_task(task)
				task.cancel()
			raise
		if task in done or task.done():
			return task.result()
		expired.set()
		self._track_screenshot_recovery_task(task)
		task.cancel()
		raise TimeoutError(f'CDP screenshot recovery exceeded the {timeout_seconds:g}-second hard deadline')

	def _track_screenshot_recovery_task(self, task: asyncio.Task[Any]) -> None:
		self._screenshot_recovery_tasks.add(task)

		def consume_and_discard(finished: asyncio.Task[Any]) -> None:
			self._screenshot_recovery_tasks.discard(finished)
			_consume_detached_task_result(finished)

		task.add_done_callback(consume_and_discard)

	async def _capture_cdp_screenshot(self, page: Page, path: Path, expired: asyncio.Event) -> bytes:
		cdp_session = None
		try:
			cdp_session = await self.context.new_cdp_session(page)
			result = await cdp_session.send(
				'Page.captureScreenshot',
				{'format': 'png', 'captureBeyondViewport': False},
			)
			encoded = result.get('data')
			if not isinstance(encoded, str) or not encoded:
				raise ValueError('CDP screenshot response did not include image data')
			try:
				screenshot = base64.b64decode(encoded, validate=True)
			except (ValueError, TypeError) as exc:
				raise ValueError('CDP screenshot response was not valid base64') from exc
			if not screenshot.startswith(_PNG_SIGNATURE):
				raise ValueError('CDP screenshot response was not a PNG image')
			if expired.is_set():
				raise TimeoutError('CDP screenshot recovery expired before the image could be persisted')
			path.write_bytes(screenshot)
			return screenshot
		finally:
			if cdp_session is not None:
				await cdp_session.detach()

	async def execute(self, decision: AgentDecision | Mapping[str, Any]) -> str:
		"""Execute one model decision and return a compact action result."""

		self._ensure_started()
		async with self._execute_lock:
			data = self._decision_dict(decision)
			action, params = self._normalise_action(data)
			page = self._active_page()
			previous_page = page
			previous_url = page.url
			known_page_ids = {id(item) for item in self._live_owned_pages()}
			await self._clear_markers(remove_attributes=False)

			result = await self._perform_action(action, params)
			if action not in {'wait'} and self.page is not None and not self.page.is_closed():
				await self.page.wait_for_timeout(200)
			await self._enforce_search_policy(previous_page, previous_url, known_page_ids)
			return result

	def capture_payload(self) -> dict[str, Any]:
		"""Return the official ``capture.json`` envelope.

		The request keys from the reference implementation are retained verbatim;
		bounded response fields are additive and ignored by older evaluators.
		"""

		return {
			'capture_time': datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S'),
			'total_requests': len(self.all_requests),
			'all_requests': copy.deepcopy(self.all_requests),
		}

	async def save_capture(self, path: str | Path | None = None) -> Path:
		await self._drain_background_tasks()
		await self._drain_download_tasks()
		destination = Path(path) if path is not None else self.task_dir / 'capture.json'
		destination.parent.mkdir(parents=True, exist_ok=True)
		destination.write_text(json.dumps(self.capture_payload(), ensure_ascii=False, indent=2), encoding='utf-8')
		return destination

	async def close(self) -> None:
		"""Detach listeners, finish collectors, and close only runtime-owned pages."""

		if self._closed:
			return
		self._closed = True
		for event, handler in self._context_handlers:
			with contextlib.suppress(Exception):
				self.context.remove_listener(event, handler)
		self._context_handlers.clear()

		for page in list(self._owned_pages):
			for event, handler in self._page_handlers.pop(id(page), []):
				with contextlib.suppress(Exception):
					page.remove_listener(event, handler)
		await self._cancel_screenshot_recovery_tasks()
		await self._drain_background_tasks()
		await self._drain_download_tasks()
		await self._drain_policy_tasks()
		with contextlib.suppress(Exception):
			await self._clear_markers(remove_attributes=True)
		for page in reversed(self._owned_pages):
			with contextlib.suppress(Exception):
				if not page.is_closed():
					await page.close(run_before_unload=False)
		self._owned_pages.clear()
		self._element_bindings.clear()
		self._request_entries.clear()
		self._request_entries_by_id.clear()
		self._request_ids_by_entry.clear()
		self._network_request_entries.clear()
		self._network_entries_by_id.clear()
		self._network_responses.clear()
		self._network_page_cursors.clear()
		self._page_ids.clear()
		self._page_document_generations.clear()
		self._reserved_download_paths.clear()
		self.page = None

	async def _cancel_screenshot_recovery_tasks(self) -> None:
		tasks = tuple(self._screenshot_recovery_tasks)
		if not tasks:
			return
		for task in tasks:
			task.cancel()
		cleanup_timeout = min(1.0, max(0.0, self.cdp_screenshot_timeout_ms / 1000))
		_, pending = await asyncio.wait(tasks, timeout=cleanup_timeout)
		if pending:
			self.logger.warning('%d timed-out CDP screenshot recovery task(s) resisted bounded cleanup', len(pending))

	def _install_context_listeners(self) -> None:
		handlers: list[tuple[str, Callable[..., Any]]] = [
			('page', self._on_context_page),
			('request', self._on_request),
			('response', self._on_response),
			('requestfailed', self._on_request_failed),
		]
		for event, handler in handlers:
			cast(Any, self.context).on(event, handler)
		self._context_handlers.extend(handlers)

	def _register_page(self, page: Page, *, make_active: bool) -> None:
		if all(id(existing) != id(page) for existing in self._owned_pages):
			self._owned_pages.append(page)
			self._page_ids[id(page)] = len(self._page_ids)
			self._page_document_generations[id(page)] = 0
		page.set_default_timeout(self.action_timeout_ms)
		page.set_default_navigation_timeout(self.navigation_timeout_ms)
		if id(page) not in self._page_handlers:
			frame_handler = lambda frame, owner=page: self._on_frame_navigated(owner, frame)
			download_handler = self._on_download
			close_handler = lambda owner=page: self._on_page_closed(owner)
			page.on('framenavigated', frame_handler)
			page.on('download', download_handler)
			page.on('close', close_handler)
			self._page_handlers[id(page)] = [
				('framenavigated', frame_handler),
				('download', download_handler),
				('close', close_handler),
			]
		if make_active:
			self.page = page
		self._record_url(page.url, unless_last=True)

	async def _configure_owned_page(self, page: Page) -> None:
		"""Configure the User-Agent consistently before an owned page navigates."""

		if self.declared_user_agent and is_sec_url(self.website):
			await page.set_extra_http_headers({'User-Agent': self.declared_user_agent})
			return

		user_agent = await page.evaluate('navigator.userAgent')
		if not isinstance(user_agent, str):
			return
		normalized_user_agent = user_agent.replace('HeadlessChrome/', 'Chrome/')
		if normalized_user_agent == user_agent:
			return

		# Keep the HTTP header and JavaScript-visible value in sync for the page's
		# future documents. This must run before its first navigation.
		await page.set_extra_http_headers({'User-Agent': normalized_user_agent})
		await page.add_init_script(
			f"""
			(() => {{
				const userAgent = {json.dumps(normalized_user_agent)};
				Object.defineProperty(Navigator.prototype, 'userAgent', {{
					configurable: true,
					get: () => userAgent,
				}});
			}})();
			"""
		)

	def _capture_request_headers(self, headers: Mapping[str, str]) -> dict[str, str]:
		"""Copy request headers while keeping configured contact details private."""

		captured = dict(headers)
		if not self.declared_user_agent:
			return captured
		for name, value in tuple(captured.items()):
			if name.lower() == 'user-agent' and value == self.declared_user_agent:
				captured[name] = _REDACTED
		return captured

	def _on_context_page(self, page: Page) -> None:
		if self._closed:
			return
		self._register_page(page, make_active=True)
		# A target=_blank link or script-created popup is registered by the
		# browser, not by one of our explicit new-page actions. Apply the same
		# page-scoped User-Agent configuration for its follow-up requests.
		self._spawn_background(self._configure_owned_page(page))

	def _on_page_closed(self, page: Page) -> None:
		if self.page is page:
			live = [item for item in self._owned_pages if item is not page and not item.is_closed()]
			self.page = live[-1] if live else None

	def _on_frame_navigated(self, page: Page, frame: Frame) -> None:
		if frame is not page.main_frame:
			return
		self._page_document_generations[id(page)] = self._page_document_generations.get(id(page), 0) + 1
		url = frame.url
		if is_forbidden_search_url(url):
			if id(page) not in self._rollback_pages and not self._closed:
				self._rollback_pages.add(id(page))
				self._spawn_policy(self._rollback_forbidden_page(page))
		elif url and url not in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
			self._record_url(url)
			self._last_safe_urls[id(page)] = url

	def _network_request_owner(self, request: Request) -> tuple[int | None, int | None, str, str]:
		"""Resolve a request to its owning top-level page and document generation."""

		try:
			frame = request.frame
			page = frame.page
		except Exception:
			return None, None, '', ''
		page_key = id(page)
		document_generation = self._page_document_generations.get(page_key)
		# The main-document request fires before ``framenavigated`` advances the
		# generation. Attribute the entire redirect chain to the document it is
		# creating so the current-page snapshot includes its own HTML request.
		with contextlib.suppress(Exception):
			if request.is_navigation_request() and frame is page.main_frame and document_generation is not None:
				document_generation += 1
		return (
			self._page_ids.get(page_key),
			document_generation,
			str(page.url),
			str(frame.url),
		)

	def _on_request(self, request: Request) -> None:
		post_data: str | None = None
		raw_post_data: str | None = None
		json_data: Any = None
		with contextlib.suppress(Exception):
			post_data = request.post_data
		with contextlib.suppress(Exception):
			raw = request.post_data_buffer
			if raw:
				raw_post_data = raw.hex()
				if post_data is None:
					post_data = raw.decode('utf-8', errors='replace')
		if post_data:
			with contextlib.suppress(json.JSONDecodeError):
				json_data = json.loads(post_data)
		page_id, document_generation, page_url, frame_url = self._network_request_owner(request)
		request_id = self._next_network_request_id
		self._next_network_request_id += 1
		network_entry: dict[str, Any] = {
			'request_id': request_id,
			'timestamp': time.time(),
			'url': request.url,
			'method': request.method,
			'headers': self._capture_request_headers(request.headers),
			'resource_type': request.resource_type,
			'post_data': post_data,
			'post_text': post_data,
			'json_data': json_data,
			'page_id': page_id,
			'document_generation': document_generation,
			'page_url': page_url,
			'frame_url': frame_url,
			'response_body_state': 'pending',
		}
		if raw_post_data is not None:
			network_entry['raw_post_data'] = raw_post_data
		self.network_requests.append(network_entry)
		self._network_request_entries[id(request)] = network_entry
		self._network_entries_by_id[request_id] = network_entry

		if request.resource_type in {'xhr', 'fetch'}:
			# Preserve the historical capture schema exactly; page provenance and
			# request IDs belong only to the action-specific all-resource cache.
			entry = {
				'timestamp': network_entry['timestamp'],
				'url': network_entry['url'],
				'method': network_entry['method'],
				'headers': copy.deepcopy(network_entry['headers']),
				'resource_type': network_entry['resource_type'],
				'post_data': post_data,
				'post_text': post_data,
				'json_data': copy.deepcopy(json_data),
			}
			if raw_post_data is not None:
				entry['raw_post_data'] = raw_post_data
			self.all_requests.append(entry)
			self._request_entries[id(request)] = entry
			self._request_entries_by_id[request_id] = entry
			self._request_ids_by_entry[id(entry)] = request_id

	def _on_response(self, response: Response) -> None:
		network_entry = self._network_request_entries.get(id(response.request))
		if network_entry is not None:
			network_entry['status'] = response.status
			network_entry['response_headers'] = dict(response.headers)
			network_entry['response_body_state'] = 'available'
			self._network_responses[int(network_entry['request_id'])] = response
		entry = self._request_entries.get(id(response.request))
		if entry is not None:
			self._spawn_background(self._capture_response(response, entry))
		extension = self._document_response_extension(response)
		if extension is not None:
			key = (response.url, response.status)
			if key not in self._captured_document_responses:
				self._captured_document_responses.add(key)
				self._spawn_download(self._save_document_response(response, extension))

	def _on_request_failed(self, request: Request) -> None:
		network_entry = self._network_request_entries.get(id(request))
		if network_entry is not None:
			with contextlib.suppress(Exception):
				network_entry['failure'] = request.failure
			network_entry['response_body_state'] = 'failed'
		entry = self._request_entries.get(id(request))
		if entry is None:
			return
		with contextlib.suppress(Exception):
			entry['failure'] = request.failure

	async def _capture_response(self, response: Response, entry: dict[str, Any]) -> None:
		try:
			entry['status'] = response.status
			entry['response_headers'] = dict(response.headers)
			content_length = self._content_length(response.headers)
			if self.max_response_body_bytes <= 0:
				return
			if content_length is not None and content_length > self.max_response_body_bytes:
				entry['response_body_omitted'] = f'content-length {content_length} exceeds limit'
				entry['response_body_truncated'] = True
				return
			body = await asyncio.wait_for(response.body(), timeout=max(1, self.action_timeout_ms / 1000))
			truncated = len(body) > self.max_response_body_bytes
			bounded = body[: self.max_response_body_bytes]
			content_type = response.headers.get('content-type', '').lower()
			if self._is_textual_content(content_type, bounded):
				text = bounded.decode(self._charset(content_type), errors='replace')
				entry['response_body'] = text
				if 'json' in content_type:
					with contextlib.suppress(json.JSONDecodeError):
						entry['response_json'] = json.loads(text)
			else:
				entry['response_body_base64'] = base64.b64encode(bounded).decode('ascii')
				entry['response_body_encoding'] = 'base64'
			entry['response_body_bytes'] = len(body)
			entry['response_body_truncated'] = truncated
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			entry['response_error'] = f'{type(exc).__name__}: {exc}'[:500]

	async def settle_network_capture(self) -> None:
		"""Finish response collectors that were already scheduled by the browser."""

		await self._drain_background_tasks()

	def current_page_network_requests(self) -> list[dict[str, Any]]:
		"""Return all HTTP requests belonging to the active page's current document.

		Requests initiated by child frames retain the top-level page identifier, so
		Tableau and similar cross-origin embeds are included. Requests from previous
		documents and other tabs are excluded.
		"""

		page = self._active_page()
		page_key = id(page)
		page_id = self._page_ids.get(page_key)
		document_generation = self._page_document_generations.get(page_key)
		if page_id is None or document_generation is None:
			return []
		return [
			copy.deepcopy(entry)
			for entry in self.network_requests
			if entry.get('page_id') == page_id and entry.get('document_generation') == document_generation
		]

	async def materialize_network_request(
		self,
		request_id: int,
		*,
		max_body_bytes: int = _MAX_DOWNLOAD_BYTES,
	) -> dict[str, Any]:
		"""Return one all-resource request with its complete bounded response body."""

		entry = self._network_entries_by_id.get(int(request_id))
		if entry is None:
			raise ValueError(f'Unknown network request_id {request_id}')
		state = str(entry.get('response_body_state', 'pending'))
		if state in {'complete', 'body_too_large', 'failed', 'unavailable', 'error'}:
			return copy.deepcopy(entry)

		response = self._network_responses.get(int(request_id))
		if response is None:
			entry['response_body_state'] = 'unavailable' if entry.get('status') is not None else 'pending'
			return copy.deepcopy(entry)

		headers = entry.get('response_headers')
		if not isinstance(headers, Mapping):
			headers = dict(response.headers)
			entry['response_headers'] = dict(headers)
		content_length = self._content_length(headers)
		if content_length is not None and content_length > max_body_bytes:
			entry['response_body_state'] = 'body_too_large'
			entry['response_body_bytes'] = content_length
			entry['response_error'] = f'content-length {content_length} exceeds {max_body_bytes} byte limit'
			return copy.deepcopy(entry)

		try:
			body = await asyncio.wait_for(response.body(), timeout=max(1, self.action_timeout_ms / 1000))
			entry['response_body_bytes'] = len(body)
			entry['response_body_sha256'] = hashlib.sha256(body).hexdigest()
			if len(body) > max_body_bytes:
				entry['response_body_state'] = 'body_too_large'
				entry['response_error'] = f'response body {len(body)} exceeds {max_body_bytes} byte limit'
				return copy.deepcopy(entry)

			content_type = str(headers.get('content-type', '')).lower()
			if self._is_textual_content(content_type, body):
				text = body.decode(self._charset(content_type), errors='replace')
				entry['response_body'] = text
				if 'json' in content_type:
					with contextlib.suppress(json.JSONDecodeError):
						entry['response_json'] = json.loads(text)
			else:
				entry['response_body_base64'] = base64.b64encode(body).decode('ascii')
				entry['response_body_encoding'] = 'base64'
			entry['response_body_state'] = 'complete'
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			entry['response_body_state'] = 'error'
			entry['response_error'] = f'{type(exc).__name__}: {exc}'[:500]
		return copy.deepcopy(entry)

	def _on_download(self, download: Download) -> None:
		self._download_urls_seen.add(download.url)
		self._spawn_download(self._save_download(download))

	async def _save_download(self, download: Download) -> None:
		suggested = self._safe_download_name(download.suggested_filename)
		destination = self._unique_download_path(suggested)
		item: dict[str, Any] = {
			'timestamp': time.time(),
			'url': download.url,
			'suggested_filename': download.suggested_filename,
			'filename': destination.name,
			'path': str(destination),
		}
		self.downloads.append(item)
		try:
			await download.save_as(str(destination))
			failure = await download.failure()
			if failure:
				item['failure'] = failure
				return
			item['size_bytes'] = destination.stat().st_size
			if destination.suffix.lower() == '.zip':
				async with self._archive_extraction_lock:
					extraction = await asyncio.to_thread(self._extract_download_archive, destination)
				item['extraction_status'] = extraction.status
				item['extracted_count'] = len(extraction.members)
				item['skipped_count'] = extraction.skipped_count
				item['extraction_warnings'] = extraction.warnings
				item['text'] = self._archive_extraction_summary(extraction)
				item['text_truncated'] = False
				for member in extraction.members:
					text = ''
					truncated = False
					text_extraction_error: str | None = None
					if member.path.suffix.lower() != '.zip':
						try:
							text, truncated = await asyncio.to_thread(self._extract_download_text, member.path)
						except Exception as exc:
							text_extraction_error = f'{type(exc).__name__}: {exc}'[:500]
							text = f'[text extraction failed: {text_extraction_error}]'
					member_item = _ArchiveMemberDownload(
						timestamp=time.time(),
						url=download.url,
						suggested_filename=member.archive_name,
						filename=member.relative_name,
						path=str(member.path),
						size_bytes=member.size_bytes,
						source='archive_member',
						source_archive=destination.name,
						text=text,
						text_truncated=truncated,
						text_extraction_error=text_extraction_error,
					)
					self.downloads.append(member_item.model_dump(exclude_none=True))
			else:
				text, truncated = await asyncio.to_thread(self._extract_download_text, destination)
				item['text'] = text
				item['text_truncated'] = truncated
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			item['failure'] = f'{type(exc).__name__}: {exc}'[:500]

	async def _save_document_response(self, response: Response, extension: str) -> None:
		"""Persist inline PDF/Office documents whose viewer DOM has no useful text."""

		# A headless Chromium build can convert even an ``inline`` PDF response
		# into a Playwright Download.  Give that higher-fidelity event one loop
		# turn to arrive and avoid writing the same document through both paths.
		await asyncio.sleep(0.1)
		if response.url in self._download_urls_seen:
			return
		content_type = response.headers.get('content-type', '').split(';', 1)[0].strip().lower()
		name = self._response_filename(response, extension)
		destination = self._unique_download_path(name)
		item: dict[str, Any] = {
			'timestamp': time.time(),
			'url': response.url,
			'suggested_filename': name,
			'filename': destination.name,
			'path': str(destination),
			'mime_type': content_type,
			'source': 'document_response',
		}
		self.downloads.append(item)
		try:
			content_length = self._content_length(response.headers)
			if content_length is not None and content_length > _MAX_DOWNLOAD_BYTES:
				item['failure'] = f'document body exceeds {_MAX_DOWNLOAD_BYTES} byte extraction limit'
				item['text'] = ''
				item['text_truncated'] = True
				return
			body = await asyncio.wait_for(response.body(), timeout=max(1, self.navigation_timeout_ms / 1000))
			if len(body) > _MAX_DOWNLOAD_BYTES:
				item['failure'] = f'document body exceeds {_MAX_DOWNLOAD_BYTES} byte extraction limit'
				item['size_bytes'] = len(body)
				item['text'] = ''
				item['text_truncated'] = True
				return
			destination.write_bytes(body)
			item['size_bytes'] = len(body)
			text, truncated = await asyncio.to_thread(self._extract_download_text, destination)
			item['text'] = text
			item['text_truncated'] = truncated
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			item['failure'] = f'{type(exc).__name__}: {exc}'[:500]

	async def _perform_action(self, action: str, params: dict[str, Any]) -> str:
		page = self._active_page()
		if action == 'click':
			locator = self._target_locator(params)
			await locator.click(timeout=self.action_timeout_ms)
			return f'clicked element {self._target_index(params)}'
		if action == 'double_click':
			locator = self._target_locator(params)
			await locator.dblclick(timeout=self.action_timeout_ms)
			return f'double-clicked element {self._target_index(params)}'
		if action == 'type':
			locator = self._target_locator(params)
			text = str(self._first(params, 'text', 'value', 'input_text', default=''))
			try:
				await locator.fill(text, timeout=self.action_timeout_ms)
			except Exception:
				await locator.click(timeout=self.action_timeout_ms)
				await locator.press('ControlOrMeta+A', timeout=self.action_timeout_ms)
				await locator.type(text, timeout=self.action_timeout_ms)
			if self._as_bool(self._first(params, 'submit', default=False)):
				await locator.press('Enter', timeout=self.action_timeout_ms)
			return f'typed into element {self._target_index(params)}'
		if action == 'select':
			locator = self._target_locator(params)
			value = self._first(params, 'value', 'text', 'option')
			if value is None:
				raise ValueError('select requires value/text/option')
			try:
				selected = await locator.select_option(value=str(value), timeout=self.action_timeout_ms)
			except Exception:
				selected = await locator.select_option(label=str(value), timeout=self.action_timeout_ms)
			return f'selected {selected!r} on element {self._target_index(params)}'
		if action == 'press':
			key = str(self._first(params, 'key', 'text', 'value', default='Enter'))
			if self._has_target(params):
				await self._target_locator(params).press(key, timeout=self.action_timeout_ms)
			else:
				await page.keyboard.press(key)
			return f'pressed {key}'
		if action == 'scroll':
			return await self._scroll(params)
		if action == 'hover':
			await self._target_locator(params).hover(timeout=self.action_timeout_ms)
			return f'hovered element {self._target_index(params)}'
		if action == 'xy':
			x = float(self._required(params, 'x'))
			y = float(self._required(params, 'y'))
			button = cast(Literal['left', 'middle', 'right'], str(self._first(params, 'button', default='left')))
			if button not in {'left', 'middle', 'right'}:
				raise ValueError(f'Unsupported mouse button: {button!r}')
			await page.mouse.click(x, y, button=button)
			return f'clicked coordinates ({x:g}, {y:g})'
		if action == 'hover_xy':
			x = float(self._required(params, 'x'))
			y = float(self._required(params, 'y'))
			await page.mouse.move(x, y)
			return f'hovered coordinates ({x:g}, {y:g})'
		if action == 'drag':
			return await self._drag(params)
		if action == 'back':
			await page.go_back(wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			return f'navigated back to {page.url}'
		if action == 'navigate':
			url = str(self._required(params, 'url'))
			self._validate_navigation_url(url)
			if is_forbidden_search_url(url):
				raise ValueError('External search engines are prohibited by the challenge rules')
			new_tab = self._as_bool(self._first(params, 'new_tab', default=False))
			if new_tab:
				page = await self.context.new_page()
				self._register_page(page, make_active=True)
				await self._configure_owned_page(page)
			download_started = await self._goto_exact(page, url)
			self._record_url(url if download_started else page.url, unless_last=True)
			return f'downloaded document from {url}' if download_started else f'navigated to {page.url}'
		if action == 'wait':
			milliseconds = self._wait_milliseconds(params)
			await page.wait_for_timeout(milliseconds)
			return f'waited {milliseconds}ms'
		if action == 'tab':
			return await self._tab_action(params)
		if action == 'read':
			return await self._read(params)
		if action == 'find':
			return await self._find(params)
		if action == 'inspect_network':
			return await self._inspect_network(params)
		if action == 'calculate':
			return self._calculate(params)
		if action in {'done', 'finish', 'answer', 'noop'}:
			return str(self._first(params, 'answer', 'text', 'message', default=action))
		raise ValueError(f'Unsupported browser action: {action!r}')

	async def _scroll(self, params: dict[str, Any]) -> str:
		direction = str(self._first(params, 'direction', default='down')).lower()
		amount_value = self._first(params, 'amount', 'pixels', 'delta')
		if amount_value is None and self._first(params, 'pages') is not None:
			viewport = await self._viewport(self._active_page())
			axis_size = viewport['height'] if direction in {'up', 'down'} else viewport['width']
			amount_value = max(axis_size, 600) * float(self._first(params, 'pages'))
		amount = abs(float(amount_value if amount_value is not None else 600))
		delta_x = float(self._first(params, 'delta_x', 'dx', default=0))
		delta_y = float(self._first(params, 'delta_y', 'dy', default=0))
		if not delta_x and not delta_y:
			if direction in {'up', 'down'}:
				delta_y = -amount if direction == 'up' else amount
			else:
				delta_x = -amount if direction == 'left' else amount
		if self._has_target(params):
			result = await self._target_locator(params).evaluate(
				"""(element, delta) => {
					const permitsScroll = (node, axis) => {
						const style = getComputedStyle(node);
						const overflow = axis === 'x' ? style.overflowX : style.overflowY;
						const hasRange = axis === 'x'
							? node.scrollWidth > node.clientWidth + 1
							: node.scrollHeight > node.clientHeight + 1;
						return hasRange && /(auto|scroll|overlay)/.test(overflow);
					};
					let target = element;
					while (target && !permitsScroll(target, 'x') && !permitsScroll(target, 'y')) {
						target = target.parentElement;
					}
					target = target || document.scrollingElement || document.documentElement;
					const beforeX = target.scrollLeft;
					const beforeY = target.scrollTop;
					target.scrollBy(delta.x, delta.y);
					return {
						tag: target.tagName.toLowerCase(),
						id: target.id || '',
						className: typeof target.className === 'string' ? target.className : '',
						beforeX,
						beforeY,
						afterX: target.scrollLeft,
						afterY: target.scrollTop,
					};
				}""",
				{'x': delta_x, 'y': delta_y},
			)
			target_name = str(result.get('tag', 'element'))
			if result.get('id'):
				target_name += f'#{result["id"]}'
			return (
				f'scrolled {target_name} from ({result.get("beforeX", 0):g}, {result.get("beforeY", 0):g}) '
				f'to ({result.get("afterX", 0):g}, {result.get("afterY", 0):g}); '
				f'requested ({delta_x:g}, {delta_y:g})'
			)
		else:
			await self._active_page().mouse.wheel(delta_x, delta_y)
		return f'scrolled ({delta_x:g}, {delta_y:g})'

	async def _drag(self, params: dict[str, Any]) -> str:
		if self._first(params, 'source_index', 'from_index') is not None:
			source = self._locator_for_index(int(self._first(params, 'source_index', 'from_index')))
			target_value = self._first(params, 'target_index', 'to_index')
			if target_value is None:
				raise ValueError('element drag requires target_index/to_index')
			target = self._locator_for_index(int(target_value))
			await source.drag_to(target, timeout=self.action_timeout_ms)
			return f'dragged element {self._first(params, "source_index", "from_index")} to {target_value}'
		start_x = float(self._first(params, 'start_x', 'from_x', 'x1', 'x'))
		start_y = float(self._first(params, 'start_y', 'from_y', 'y1', 'y'))
		end_x = float(self._first(params, 'end_x', 'to_x', 'x2'))
		end_y = float(self._first(params, 'end_y', 'to_y', 'y2'))
		mouse = self._active_page().mouse
		await mouse.move(start_x, start_y)
		await mouse.down()
		await mouse.move(end_x, end_y, steps=12)
		await mouse.up()
		return f'dragged ({start_x:g}, {start_y:g}) to ({end_x:g}, {end_y:g})'

	async def _tab_action(self, params: dict[str, Any]) -> str:
		operation = str(self._first(params, 'operation', 'tab_action', 'command', default='switch')).lower()
		pages = self._live_owned_pages()
		if operation in {'new', 'open', 'create'}:
			page = await self.context.new_page()
			self._register_page(page, make_active=True)
			await self._configure_owned_page(page)
			url = self._first(params, 'url')
			if url is not None:
				url = str(url)
				self._validate_navigation_url(url)
				if is_forbidden_search_url(url):
					raise ValueError('External search engines are prohibited by the challenge rules')
				download_started = await self._goto_exact(page, url)
				if download_started:
					self._record_url(url, unless_last=True)
			return f'opened tab {len(self._live_owned_pages()) - 1}'
		if not pages:
			raise RuntimeError('No live tabs')
		if operation in {'close', 'remove'}:
			page = self._select_tab(params, pages)
			await page.close(run_before_unload=False)
			self.page = self._live_owned_pages()[-1] if self._live_owned_pages() else None
			return 'closed tab'
		if operation in {'next', 'previous', 'prev'}:
			current = pages.index(self._active_page())
			delta = 1 if operation == 'next' else -1
			page = pages[(current + delta) % len(pages)]
		else:
			page = self._select_tab(params, pages)
		self.page = page
		await page.bring_to_front()
		return f'switched to tab {pages.index(page)}: {page.url}'

	async def _read(self, params: dict[str, Any]) -> str:
		if self._has_target(params):
			text = await self._target_locator(params).inner_text(timeout=self.action_timeout_ms)
			return text[:_MAX_PAGE_TEXT]
		return await self._collect_page_text(self._active_page())

	async def _find(self, params: dict[str, Any]) -> str:
		await self._drain_download_tasks()
		query = str(self._first(params, 'query', 'text', 'value', default='')).strip()
		if not query:
			raise ValueError('find requires query/text')
		matches: list[str] = []
		revealed_control = False
		for frame_index, frame in enumerate(self._active_page().frames):
			try:
				locator = frame.get_by_text(query, exact=False)
				count = min(await locator.count(), 20)
				for match_index in range(count):
					item = locator.nth(match_index)
					if await item.is_visible():
						text = re.sub(r'\s+', ' ', await item.inner_text(timeout=self.action_timeout_ms)).strip()
						metadata = await item.evaluate(
							"""(element, options) => {
								const interactiveSelector = [
									'a[href]', 'button', 'input:not([type="hidden"])', 'textarea', 'select',
									'[contenteditable="true"]', '[role="button"]', '[role="link"]',
									'[role="tab"]', '[tabindex]:not([tabindex="-1"])'
								].join(',');
								const interactive = element.closest(interactiveSelector)
									|| element.querySelector(interactiveSelector);
								let rect = element.getBoundingClientRect();
								const wasInViewport = rect.bottom > 0 && rect.right > 0
									&& rect.top < window.innerHeight && rect.left < window.innerWidth;
								let revealed = false;
								if (options.reveal && interactive && !wasInViewport) {
									interactive.scrollIntoView({block: 'center', inline: 'nearest'});
									rect = interactive.getBoundingClientRect();
									revealed = true;
								}
								const anchor = (interactive && interactive.matches('a[href]') ? interactive : null)
									|| element.closest('a[href]') || element.querySelector('a[href]');
								return {
									tag: (interactive || element).tagName.toLowerCase(),
									id: (interactive || element).id || '',
									href: anchor ? anchor.href : '',
									interactiveIndex: (interactive || element).getAttribute('data-webretriever-index') || '',
									inViewport: rect.bottom > 0 && rect.right > 0
										&& rect.top < window.innerHeight && rect.left < window.innerWidth,
									revealed,
								};
							}""",
							{'reveal': not revealed_control},
						)
						if metadata.get('revealed'):
							revealed_control = True
						details = [
							f'tag={metadata.get("tag", "")}',
							f'in_viewport={str(bool(metadata.get("inViewport"))).lower()}',
						]
						if metadata.get('revealed'):
							details.append('revealed_for_next_observation=true')
						if metadata.get('id'):
							details.append(f'id={metadata["id"]}')
						if metadata.get('interactiveIndex'):
							details.append(f'element_id={metadata["interactiveIndex"]}')
						if metadata.get('href'):
							details.append(f'href={metadata["href"]}')
						matches.append(f'frame {frame_index}: {text[:500]} | {" | ".join(details)}')
						if len(matches) == 20:
							break
			except Exception:
				continue
			if len(matches) == 20:
				break
		query_lower = query.casefold()
		for download in self.downloads:
			text = self._download_text_for_search(download)
			start = text.casefold().find(query_lower)
			if start < 0:
				continue
			context_start = max(0, start - 500)
			context_end = min(len(text), start + len(query) + 1_500)
			filename = download.get('filename', download.get('suggested_filename', 'download'))
			matches.append(f'download {filename}: {text[context_start:context_end]}')
			if len(matches) == 20:
				break
		return '\n'.join(matches) if matches else f'No visible text matching {query!r}'

	def _download_text_for_search(self, download: Mapping[str, Any]) -> str:
		"""Return searchable download text, lazily reopening truncated text files.

		Observation previews stay bounded, but ``find_text`` must be able to locate
		rows near the end of a competition CSV.  The file itself was already saved
		through Playwright and capped by ``_MAX_DOWNLOAD_BYTES``.
		"""
		text = str(download.get('text', ''))
		if not download.get('text_truncated'):
			return text
		path_value = download.get('path')
		if not isinstance(path_value, str):
			return text
		path = Path(path_value)
		if path.suffix.lower() not in {'.csv', '.tsv', '.txt', '.json', '.xml'}:
			return text
		try:
			resolved = path.resolve(strict=True)
			if not resolved.is_relative_to(self.download_dir.resolve()):
				return text
			if resolved.stat().st_size > _MAX_DOWNLOAD_BYTES:
				return text
			return resolved.read_text(encoding='utf-8-sig', errors='replace')
		except OSError:
			return text

	async def _inspect_network(self, params: dict[str, Any]) -> str:
		request_id_value = self._first(params, 'request_id')
		if request_id_value is not None:
			return await self._inspect_network_request(
				int(request_id_value),
				cursor=self._first(params, 'network_cursor', 'cursor'),
			)

		query = str(self._first(params, 'query', 'text', default=''))
		limit = min(max(int(self._first(params, 'limit', default=10)), 1), 50)
		if not query:
			items = []
			for item in self.all_requests[-limit:]:
				output = copy.deepcopy(item)
				request_id = self._request_ids_by_entry.get(id(item))
				if request_id is not None:
					output['request_id'] = request_id
				items.append(output)
			return json.dumps(items, ensure_ascii=False, indent=2)[:_MAX_PAGE_TEXT]

		await self._settle_network_capture_bounded(_NETWORK_SEARCH_TIMEOUT_SECONDS)
		requests = []
		for item in self.all_requests:
			request_id = self._request_ids_by_entry.get(id(item))
			if request_id is None:
				continue
			requests.append(
				{
					'request_id': request_id,
					'timestamp': item.get('timestamp'),
					'url': item.get('url'),
					'method': item.get('method'),
					'status': item.get('status'),
					'resource_type': item.get('resource_type'),
					'post_data': item.get('post_data'),
					'json_data': item.get('json_data'),
					'response_headers': item.get('response_headers'),
					'response_body': item.get('response_body'),
					'response_body_truncated': item.get('response_body_truncated', False),
				}
			)
		payload = {'query': query, 'requests': requests}
		result = await self._run_network_search(payload)
		# Lunr tokenizes a query, so a response containing only a few generic
		# terms can rank highly without containing the complete requested phrase.
		# Count literal response-body matches separately so callers can distinguish
		# relevance from an exact hit. Request metadata is intentionally excluded:
		# a phrase echoed in a URL or request payload is not retrieved evidence.
		result = {
			'exact_match_count': self._exact_network_response_match_count(query, requests),
			**result,
		}
		return json.dumps(result, ensure_ascii=False, indent=2)[:_MAX_PAGE_TEXT]

	@staticmethod
	def _exact_network_response_match_count(query: str, requests: Sequence[Mapping[str, Any]]) -> int:
		"""Count captured response bodies containing ``query`` as one full phrase.

		The count is per captured request (rather than per textual occurrence), and
		uses Unicode NFKC normalization plus case folding. This keeps a title match
		stable across harmless compatibility variants while preserving whitespace and
		word boundaries: token-level or partial Lunr matches do not qualify.
		"""

		normalized_query = unicodedata.normalize('NFKC', query).casefold()
		if not normalized_query:
			return 0
		return sum(
			1
			for request in requests
			if isinstance(request.get('response_body'), str)
			and normalized_query in unicodedata.normalize('NFKC', request['response_body']).casefold()
		)

	async def _inspect_network_request(self, request_id: int, *, cursor: Any = None) -> str:
		if request_id not in self._request_entries_by_id:
			raise ValueError(f'Unknown inspect_network request_id {request_id}')

		metadata_included = cursor is None
		if cursor is None:
			await self._settle_network_capture_bounded(_NETWORK_SEARCH_TIMEOUT_SECONDS)
			snapshot = await self._materialize_inspect_request(request_id)
			offset = 0
			page_number = 1
		else:
			cursor_state = self._network_page_cursors.get(str(cursor))
			if cursor_state is None or cursor_state[0] != request_id:
				raise ValueError('inspect_network cursor is invalid for this request_id')
			snapshot = await self.materialize_network_request(request_id, max_body_bytes=_MAX_DOWNLOAD_BYTES)
			offset = cursor_state[1]
			page_number = cursor_state[3]

		body_state, body_encoding, body_data = self._inspect_response_body(snapshot, request_id)
		body_sha256 = str(snapshot.get('response_body_sha256', ''))
		if not body_sha256 and body_data:
			if body_encoding == 'base64':
				with contextlib.suppress(ValueError):
					body_sha256 = hashlib.sha256(base64.b64decode(body_data, validate=True)).hexdigest()
			else:
				body_sha256 = hashlib.sha256(body_data.encode('utf-8')).hexdigest()

		if cursor is not None and cursor_state[2] != body_sha256:
			raise ValueError('inspect_network cursor no longer matches the response body')

		response = {
			'status': snapshot.get('status'),
			'headers': copy.deepcopy(snapshot.get('response_headers', {})),
			'body_encoding': body_encoding,
			'body_bytes': snapshot.get('response_body_bytes'),
			'body_characters': len(body_data),
			'body_sha256': body_sha256 or None,
			'body_truncated': bool(snapshot.get('response_body_truncated')),
			'error': snapshot.get('response_error'),
		}
		output: dict[str, Any] = {
			'mode': 'request',
			'request_id': request_id,
			'url': snapshot.get('url'),
			'body_state': body_state,
			'response': response,
			'page': {
				'metadata_included': metadata_included,
				'number': page_number,
				'offset': offset,
				'data': '',
				'next_cursor': None,
			},
		}
		if metadata_included:
			response_only = {
				'status',
				'response_headers',
				'response_body',
				'response_json',
				'response_body_base64',
				'response_body_encoding',
				'response_body_bytes',
				'response_body_sha256',
				'response_body_state',
				'response_body_truncated',
				'response_body_omitted',
				'response_error',
				'failure',
			}
			output['request'] = {key: copy.deepcopy(value) for key, value in snapshot.items() if key not in response_only}

		page_length = self._fit_network_body_page(output, body_data, offset)
		page_data = body_data[offset : offset + page_length]
		next_offset = offset + len(page_data)
		next_cursor: str | None = None
		if next_offset < len(body_data):
			next_cursor = secrets.token_urlsafe(24)
			self._network_page_cursors[next_cursor] = (request_id, next_offset, body_sha256, page_number + 1)
		output['page']['data'] = page_data
		output['page']['next_cursor'] = next_cursor
		rendered = json.dumps(output, ensure_ascii=False, indent=2)
		if len(rendered) > _MAX_PAGE_TEXT:
			raise ValueError('inspect_network request metadata exceeds the action output limit')
		return rendered

	@staticmethod
	def _fit_network_body_page(output: dict[str, Any], body_data: str, offset: int) -> int:
		remaining = max(0, len(body_data) - offset)
		length = min(_NETWORK_BODY_PAGE_CHARACTERS, remaining)
		cursor_placeholder = 'x' * 32
		while True:
			output['page']['data'] = body_data[offset : offset + length]
			output['page']['next_cursor'] = cursor_placeholder if length < remaining else None
			rendered_length = len(json.dumps(output, ensure_ascii=False, indent=2))
			if rendered_length <= _MAX_PAGE_TEXT:
				if remaining and length == 0:
					raise ValueError('inspect_network request metadata leaves no room for response data')
				return length
			if length == 0:
				raise ValueError('inspect_network request metadata exceeds the action output limit')
			length = max(0, length - max(1, rendered_length - _MAX_PAGE_TEXT))

	async def _materialize_inspect_request(self, request_id: int) -> dict[str, Any]:
		deadline = asyncio.get_running_loop().time() + _NETWORK_SEARCH_TIMEOUT_SECONDS
		while True:
			snapshot = await self.materialize_network_request(request_id, max_body_bytes=_MAX_DOWNLOAD_BYTES)
			if snapshot.get('response_body_state') != 'pending':
				return snapshot
			remaining = deadline - asyncio.get_running_loop().time()
			if remaining <= 0:
				return snapshot
			await asyncio.sleep(min(0.1, remaining))

	def _inspect_response_body(self, snapshot: Mapping[str, Any], request_id: int) -> tuple[str, str, str]:
		state = str(snapshot.get('response_body_state', 'unavailable'))
		if isinstance(snapshot.get('response_body'), str):
			return ('complete' if state == 'complete' else 'partial', 'text', str(snapshot['response_body']))
		if isinstance(snapshot.get('response_body_base64'), str):
			return ('complete' if state == 'complete' else 'partial', 'base64', str(snapshot['response_body_base64']))

		captured = self._request_entries_by_id[request_id]
		if isinstance(captured.get('response_body'), str):
			return 'partial', 'text', str(captured['response_body'])
		if isinstance(captured.get('response_body_base64'), str):
			return 'partial', 'base64', str(captured['response_body_base64'])
		if state == 'body_too_large':
			return 'body_too_large', 'none', ''
		if state in {'error', 'failed'}:
			return 'error', 'none', ''
		return 'unavailable', 'none', ''

	async def _settle_network_capture_bounded(self, timeout_seconds: float) -> None:
		deadline = asyncio.get_running_loop().time() + max(0.0, timeout_seconds)
		while self._background_tasks:
			remaining = deadline - asyncio.get_running_loop().time()
			if remaining <= 0:
				return
			tasks = tuple(self._background_tasks)
			done, _ = await asyncio.wait(tasks, timeout=remaining)
			if not done:
				return
			self._background_tasks.difference_update(done)
			for task in done:
				with contextlib.suppress(asyncio.CancelledError, Exception):
					task.result()

	async def _run_network_search(self, payload: Mapping[str, Any]) -> dict[str, Any]:
		script = Path(__file__).with_name('lunr_network_search.js')
		try:
			process = await asyncio.create_subprocess_exec(
				'node',
				f'--max-old-space-size={_NETWORK_SEARCH_NODE_HEAP_MIB}',
				str(script),
				stdin=asyncio.subprocess.PIPE,
				stdout=asyncio.subprocess.PIPE,
				stderr=asyncio.subprocess.PIPE,
			)
		except OSError as exc:
			return self._substring_network_search(payload, f'{type(exc).__name__}: {exc}')
		input_bytes = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()
		try:
			stdout, stderr = await asyncio.wait_for(
				process.communicate(input_bytes),
				timeout=_NETWORK_SEARCH_TIMEOUT_SECONDS,
			)
		except asyncio.CancelledError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			raise
		except TimeoutError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			return self._substring_network_search(payload, 'Node search exceeded 30 seconds')
		if process.returncode != 0:
			reason = stderr.decode('utf-8', errors='replace')[:500] or f'Node search exited {process.returncode}'
			return self._substring_network_search(payload, reason)
		try:
			result = json.loads(stdout)
		except (json.JSONDecodeError, UnicodeDecodeError) as exc:
			return self._substring_network_search(payload, f'Invalid Node search output: {exc}')
		if not isinstance(result, dict):
			return self._substring_network_search(payload, 'Node search output was not an object')
		return result

	@staticmethod
	def _substring_network_search(payload: Mapping[str, Any], reason: str) -> dict[str, Any]:
		query = str(payload.get('query', ''))
		query_lower = query.casefold()
		requests = payload.get('requests')
		if not isinstance(requests, list):
			requests = []
		results = []
		for request in reversed(requests):
			if not isinstance(request, Mapping):
				continue
			serialized = json.dumps(request, ensure_ascii=False)
			start = serialized.casefold().find(query_lower)
			if start < 0:
				continue
			context_start = max(0, start - 500)
			context_end = min(len(serialized), start + len(query) + 1_000)
			results.append(
				{
					'rank': len(results) + 1,
					'request_id': request.get('request_id'),
					'score': None,
					'matched_query_terms': [query],
					'matched_fields': ['serialized_request'],
					'duplicate_count': 1,
					'request': {
						'timestamp': request.get('timestamp'),
						'url': request.get('url'),
						'method': request.get('method'),
						'status': request.get('status'),
						'resource_type': request.get('resource_type'),
						'post_data': request.get('post_data'),
					},
					'matched_chunks': [{'score': None, 'text': serialized[context_start:context_end]}],
				}
			)
			if len(results) == 10:
				break
		return {
			'search_mode': 'substring_fallback',
			'fallback_reason': reason,
			'query': query,
			'indexed_requests': len(requests),
			'pending_response_bodies': sum(
				1 for request in requests if isinstance(request, Mapping) and request.get('status') is None
			),
			'omitted_due_to_budget': 0,
			'results': results,
		}

	@classmethod
	def _calculate(cls, params: dict[str, Any]) -> str:
		"""Deterministically compare a browser-observed numeric series."""
		operation = str(cls._required(params, 'operation'))
		raw_series = cls._required(params, 'text')
		try:
			payload = json.loads(str(raw_series))
		except json.JSONDecodeError as exc:
			raise ValueError(f'calculate text must be valid JSON: {exc.msg}') from exc

		if isinstance(payload, dict):
			raw_points = list(payload.items())
		elif isinstance(payload, list):
			raw_points = []
			for position, item in enumerate(payload):
				if not isinstance(item, Mapping) or 'label' not in item or 'value' not in item:
					raise ValueError(f'calculate series item {position} must contain label and value')
				raw_points.append((item['label'], item['value']))
		else:
			raise ValueError('calculate text must encode an object or a list of label/value objects')
		if not 1 <= len(raw_points) <= 500:
			raise ValueError('calculate series must contain between 1 and 500 points')

		points: list[tuple[str, float]] = []
		for raw_label, raw_value in raw_points:
			if isinstance(raw_value, bool):
				raise ValueError(f'calculate value for {raw_label!r} must be numeric')
			try:
				value = float(raw_value)
			except (TypeError, ValueError) as exc:
				raise ValueError(f'calculate value for {raw_label!r} must be numeric') from exc
			if not math.isfinite(value):
				raise ValueError(f'calculate value for {raw_label!r} must be finite')
			points.append((str(raw_label), value))

		if operation in {'argmax', 'argmin'}:
			reverse = operation == 'argmax'
			ranked = sorted(points, key=lambda point: point[1], reverse=reverse)
			return json.dumps(
				{
					'operation': operation,
					'compared_count': len(points),
					'winner': {'label': ranked[0][0], 'value': ranked[0][1]},
					'top_candidates': [{'label': label, 'value': value} for label, value in ranked[:10]],
				},
				ensure_ascii=False,
				separators=(',', ':'),
			)

		if operation not in {'argmax_difference', 'argmin_difference', 'argmax_growth', 'argmin_growth'}:
			raise ValueError(f'unsupported calculation operation: {operation!r}')
		if len(points) < 2:
			raise ValueError(f'{operation} requires at least two points')

		try:
			ordered = sorted(points, key=lambda point: float(point[0]))
		except ValueError:
			ordered = points
		comparisons: list[dict[str, str | float]] = []
		for (previous_label, previous_value), (label, value) in zip(ordered, ordered[1:], strict=False):
			difference = value - previous_value
			if operation.endswith('growth'):
				if previous_value == 0:
					continue
				metric = difference / abs(previous_value)
			else:
				metric = difference
			comparisons.append(
				{
					'label': label,
					'value': value,
					'previous_label': previous_label,
					'previous_value': previous_value,
					'metric': metric,
				}
			)
		if not comparisons:
			raise ValueError(f'{operation} has no eligible comparisons')
		reverse = operation.startswith('argmax')
		ranked_comparisons = sorted(comparisons, key=lambda item: float(item['metric']), reverse=reverse)
		return json.dumps(
			{
				'operation': operation,
				'point_count': len(points),
				'compared_count': len(comparisons),
				'winner': ranked_comparisons[0],
				'top_candidates': ranked_comparisons[:10],
			},
			ensure_ascii=False,
			separators=(',', ':'),
		)

	async def _enforce_search_policy(
		self,
		previous_page: Page | None = None,
		previous_url: str | None = None,
		known_page_ids: set[int] | None = None,
	) -> None:
		await self._drain_policy_tasks()
		for page in list(self._live_owned_pages()):
			if not is_forbidden_search_url(page.url):
				if page.url and page.url not in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
					self._last_safe_urls[id(page)] = page.url
				continue
			if known_page_ids is not None and id(page) not in known_page_ids:
				self.logger.warning('Closing tab that escaped to prohibited search engine: %s', redact_cdp_url(page.url))
				with contextlib.suppress(Exception):
					await page.close(run_before_unload=False)
				continue
			fallback = self._last_safe_urls.get(id(page))
			if page is previous_page and previous_url and not is_forbidden_search_url(previous_url):
				fallback = previous_url
			await self._rollback_forbidden_page(page, fallback)
		self._active_page()

	async def _rollback_forbidden_page(self, page: Page, fallback: str | None = None) -> None:
		try:
			if page.is_closed() or not is_forbidden_search_url(page.url):
				return
			self.logger.warning('Rolling back prohibited external search-engine navigation: %s', redact_cdp_url(page.url))
			page_safe_url = self._last_safe_urls.get(id(page))
			other_safe_pages = [
				item for item in self._live_owned_pages() if item is not page and not is_forbidden_search_url(item.url)
			]
			if fallback is None and page_safe_url is None and other_safe_pages:
				# A newly opened search-engine popup has no meaningful history to
				# restore. Close it and return focus to the originating safe tab.
				await page.close(run_before_unload=False)
				self.page = other_safe_pages[-1]
				return
			fallback = fallback or page_safe_url or self.website
			with contextlib.suppress(Exception):
				await page.go_back(wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			if not page.is_closed() and is_forbidden_search_url(page.url) and fallback and not is_forbidden_search_url(fallback):
				await page.goto(fallback, wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			if not page.is_closed() and not is_forbidden_search_url(page.url):
				self.page = page
				self._last_safe_urls[id(page)] = page.url
		finally:
			self._rollback_pages.discard(id(page))

	async def _collect_elements(self, page: Page) -> list[ElementRef]:
		self._element_bindings.clear()
		result: list[ElementRef] = []
		next_index = 0
		for frame_index, frame in enumerate(page.frames):
			try:
				items = await frame.evaluate(
					_MARK_ELEMENTS_JS,
					{
						'start': next_index,
						'markerAttribute': _INTERACTIVE_ATTRIBUTE,
						'overlayClass': _OVERLAY_CLASS,
						'frameIndex': frame_index,
						'maxText': _MAX_ELEMENT_TEXT,
					},
				)
			except Exception as exc:
				self.logger.debug('Skipping inaccessible frame %s: %s', redact_cdp_url(frame.url), exc)
				continue
			for item in items or []:
				index = int(item['index'])
				selector = f'[{_INTERACTIVE_ATTRIBUTE}="{index}"]'
				ref = ElementRef(
					index=index,
					tag=str(item.get('tag', 'element')),
					text=str(item.get('text', '')),
					role=str(item.get('role', '')),
					name=str(item.get('name', '')),
					placeholder=str(item.get('placeholder', '')),
					href=str(item.get('href', '')),
					input_type=str(item.get('input_type', '')),
					frame_index=frame_index,
					frame_url=frame.url,
					x=float(item.get('x', 0)),
					y=float(item.get('y', 0)),
					width=float(item.get('width', 0)),
					height=float(item.get('height', 0)),
					selector=selector,
				)
				result.append(ref)
				self._element_bindings[index] = _ElementBinding(frame=frame, selector=selector)
				next_index = max(next_index, index + 1)
		return result

	async def _clear_markers(self, *, remove_attributes: bool) -> None:
		page = self.page
		if page is None or page.is_closed():
			return
		for frame in page.frames:
			with contextlib.suppress(Exception):
				await frame.evaluate(
					_CLEAR_MARKERS_JS,
					{
						'markerAttribute': _INTERACTIVE_ATTRIBUTE,
						'overlayClass': _OVERLAY_CLASS,
						'removeAttributes': remove_attributes,
					},
				)

	async def _collect_page_text(self, page: Page) -> str:
		parts: list[str] = []
		remaining = _MAX_PAGE_TEXT
		for frame_index, frame in enumerate(page.frames):
			if remaining <= 0:
				break
			try:
				text = await frame.locator('body').inner_text(timeout=min(self.action_timeout_ms, 5_000))
			except Exception:
				continue
			text = re.sub(r'[ \t]+', ' ', text).strip()
			if not text:
				continue
			prefix = '' if frame_index == 0 else f'\n[Frame {frame_index}: {frame.url}]\n'
			chunk = (prefix + text)[:remaining]
			parts.append(chunk)
			remaining -= len(chunk)
		return '\n'.join(parts)

	async def _viewport(self, page: Page) -> dict[str, int]:
		viewport = page.viewport_size
		if viewport:
			return {'width': int(viewport['width']), 'height': int(viewport['height'])}
		try:
			value = await page.evaluate('() => ({width: window.innerWidth, height: window.innerHeight})')
			return {'width': int(value['width']), 'height': int(value['height'])}
		except Exception:
			return {'width': 0, 'height': 0}

	async def _tabs(self) -> list[dict[str, Any]]:
		active = self._active_page()
		result: list[dict[str, Any]] = []
		for index, page in enumerate(self._live_owned_pages()):
			result.append(
				{
					'index': index,
					'url': page.url,
					'title': await self._page_title(page),
					'active': page is active,
				}
			)
		return result

	async def _page_title(self, page: Page) -> str:
		try:
			return await page.title()
		except Exception:
			return ''

	def _target_locator(self, params: Mapping[str, Any]) -> Locator:
		selector = self._first(params, 'selector')
		if selector is not None:
			return self._active_page().locator(str(selector)).first
		return self._locator_for_index(self._target_index(params))

	def _locator_for_index(self, index: int) -> Locator:
		binding = self._element_bindings.get(index)
		if binding is None:
			raise ValueError(f'Unknown element index {index}; call observe() before interacting')
		return binding.frame.locator(binding.selector).first

	def _target_index(self, params: Mapping[str, Any]) -> int:
		value = self._first(params, 'index', 'element_index', 'target_index', 'element_id', 'target')
		if isinstance(value, str):
			match = re.fullmatch(r'\s*\[?(\d+)\]?\s*', value)
			if match:
				value = match.group(1)
		if value is None:
			raise ValueError('Action requires an element index or selector')
		try:
			return int(value)
		except (TypeError, ValueError) as exc:
			raise ValueError(f'Invalid element index: {value!r}') from exc

	def _has_target(self, params: Mapping[str, Any]) -> bool:
		return self._first(params, 'selector', 'index', 'element_index', 'target_index', 'element_id', 'target') is not None

	def _select_tab(self, params: Mapping[str, Any], pages: Sequence[Page]) -> Page:
		value = self._first(params, 'tab_index', 'index', 'target', default=0)
		index = int(value)
		if not -len(pages) <= index < len(pages):
			raise ValueError(f'Tab index out of range: {index}')
		return pages[index]

	def _active_page(self) -> Page:
		if self.page is not None and not self.page.is_closed() and not is_forbidden_search_url(self.page.url):
			return self.page
		live = [page for page in self._live_owned_pages() if not is_forbidden_search_url(page.url)]
		if not live:
			raise RuntimeError('BrowserRuntime has no live safe page')
		self.page = live[-1]
		return self.page

	async def _active_page_for_observation(self) -> Page:
		page = self._active_page()
		if page.url != _DOWNLOAD_PLACEHOLDER_URL:
			return page

		opener: Page | None = None
		with contextlib.suppress(Exception):
			opener = await page.opener()
		if (
			opener is None
			or opener is page
			or opener.is_closed()
			or opener.url == _DOWNLOAD_PLACEHOLDER_URL
			or is_forbidden_search_url(opener.url)
		):
			opener = next(
				(
					candidate
					for candidate in reversed(self._live_owned_pages())
					if candidate is not page
					and candidate.url != _DOWNLOAD_PLACEHOLDER_URL
					and not is_forbidden_search_url(candidate.url)
				),
				None,
			)

		if opener is None:
			with contextlib.suppress(Exception):
				await page.close(run_before_unload=False)
			self.page = None
			raise RuntimeError("Download placeholder page ':' has no live safe opener or fallback page")

		self.logger.info(
			'Closing download placeholder page and restoring active page: %s',
			redact_cdp_url(opener.url),
		)
		self.page = opener
		with contextlib.suppress(Exception):
			await page.close(run_before_unload=False)
		return opener

	def _live_owned_pages(self) -> list[Page]:
		return [page for page in self._owned_pages if not page.is_closed()]

	def _ensure_started(self) -> None:
		if self._closed:
			raise RuntimeError('BrowserRuntime is closed')
		if not self._started or self.page is None:
			raise RuntimeError('Call BrowserRuntime.start(website) first')

	def _record_url(self, url: str, *, unless_last: bool = False) -> None:
		if not url or url in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
			return
		if unless_last and self.visited_urls and self.visited_urls[-1] == url:
			return
		self.visited_urls.append(url)

	def _spawn_background(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._background_tasks.add(task)

	def _spawn_download(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._download_tasks.add(task)

	def _spawn_policy(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._policy_tasks.add(task)

	async def _drain_policy_tasks(self) -> None:
		while self._policy_tasks:
			tasks = tuple(self._policy_tasks)
			self._policy_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _drain_background_tasks(self) -> None:
		while self._background_tasks:
			tasks = tuple(self._background_tasks)
			self._background_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _drain_download_tasks(self) -> None:
		while self._download_tasks:
			tasks = tuple(self._download_tasks)
			self._download_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _goto_exact(self, page: Page, url: str) -> bool:
		"""Navigate without rewriting *url*; return true for download navigations."""

		try:
			await page.goto(url, wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			return False
		except PlaywrightError as exc:
			if 'download is starting' not in str(exc).casefold():
				raise
			# The page itself stays at its previous URL, but the requested URL is
			# still valuable trajectory evidence and its Download event is handled
			# solely through Playwright.
			await page.wait_for_timeout(100)
			await self._drain_download_tasks()
			return True

	@staticmethod
	def _decision_dict(decision: AgentDecision | Mapping[str, Any]) -> dict[str, Any]:
		if isinstance(decision, Mapping):
			return dict(decision)
		model_dump = getattr(decision, 'model_dump', None)
		if callable(model_dump):
			dumped = model_dump(exclude_none=True)
			if not isinstance(dumped, Mapping):
				raise TypeError('AgentDecision.model_dump() did not return a mapping')
			return {str(key): value for key, value in dumped.items()}
		if hasattr(decision, '__dict__'):
			return {key: value for key, value in vars(decision).items() if not key.startswith('_') and value is not None}
		raise TypeError(f'Unsupported AgentDecision value: {type(decision).__name__}')

	@staticmethod
	def _normalise_action(data: dict[str, Any]) -> tuple[str, dict[str, Any]]:
		action_value = data.get('action', data.get('name', data.get('kind')))
		params = dict(data)
		if isinstance(action_value, Mapping):
			if len(action_value) != 1:
				raise ValueError('Structured action must contain exactly one operation')
			action_value, nested = next(iter(action_value.items()))
			if isinstance(nested, Mapping):
				params.update(nested)
		elif hasattr(action_value, 'value'):
			action_value = getattr(action_value, 'value')
		if action_value is None:
			raise ValueError('AgentDecision is missing action')
		action = str(action_value).strip().lower().replace('-', '_').replace(' ', '_')
		aliases = {
			'click_xy': 'xy',
			'coordinate_click': 'xy',
			'go_back': 'back',
			'goto': 'navigate',
			'open_url': 'navigate',
			'switch_tab': 'tab',
			'new_tab': 'tab',
			'close_tab': 'tab',
			'read_element': 'read',
			'find_text': 'find',
			'network': 'inspect_network',
		}
		original = action
		action = aliases.get(action, action)
		if original == 'switch_tab':
			params.setdefault('operation', 'switch')
		elif original == 'new_tab':
			params.setdefault('operation', 'new')
		elif original == 'close_tab':
			params.setdefault('operation', 'close')
		return action, params

	@staticmethod
	def _first(values: Mapping[str, Any], *names: str, default: Any = None) -> Any:
		for name in names:
			if name in values and values[name] is not None:
				return values[name]
		return default

	@classmethod
	def _required(cls, values: Mapping[str, Any], name: str) -> Any:
		value = cls._first(values, name)
		if value is None:
			raise ValueError(f'Action requires {name}')
		return value

	@staticmethod
	def _as_bool(value: Any) -> bool:
		if isinstance(value, str):
			return value.strip().lower() in {'1', 'true', 'yes', 'on'}
		return bool(value)

	def _wait_milliseconds(self, params: Mapping[str, Any]) -> int:
		if self._first(params, 'milliseconds', 'ms') is not None:
			milliseconds = float(self._first(params, 'milliseconds', 'ms'))
		else:
			seconds = float(self._first(params, 'seconds', 'duration', 'value', default=1))
			milliseconds = seconds * 1000
		return int(min(max(milliseconds, 0), self.action_timeout_ms))

	@staticmethod
	def _validate_navigation_url(url: str) -> None:
		if not isinstance(url, str) or not url:
			raise ValueError('Navigation URL must be a non-empty string')
		if any(character.isspace() or ord(character) < 32 for character in url):
			raise ValueError('Navigation URL must not contain whitespace or control characters')
		try:
			parsed = urlsplit(url)
		except ValueError as exc:
			raise ValueError(f'Invalid navigation URL: {url!r}') from exc
		if parsed.scheme not in {'http', 'https'} or not parsed.hostname:
			raise ValueError('Navigation URL must be an absolute http(s) URL')
		if parsed.username is not None or parsed.password is not None:
			raise ValueError('Navigation URL must not contain credentials')
		try:
			_ = parsed.port
		except ValueError as exc:
			raise ValueError('Navigation URL contains an invalid port') from exc

	@staticmethod
	def _safe_step_name(step: int) -> str:
		try:
			return str(max(0, int(step)))
		except (TypeError, ValueError) as exc:
			raise ValueError(f'Invalid observation step: {step!r}') from exc

	@staticmethod
	def _content_length(headers: Mapping[str, str]) -> int | None:
		try:
			return int(headers.get('content-length', ''))
		except (TypeError, ValueError):
			return None

	@staticmethod
	def _charset(content_type: str) -> str:
		match = re.search(r'charset=([\w.-]+)', content_type)
		return match.group(1) if match else 'utf-8'

	@staticmethod
	def _is_textual_content(content_type: str, body: bytes) -> bool:
		if any(marker in content_type for marker in ('text/', 'json', 'xml', 'javascript', 'x-www-form-urlencoded')):
			return True
		return b'\x00' not in body[:1024]

	@staticmethod
	def _safe_download_name(name: str) -> str:
		name = Path(str(name).replace('\\', '/')).name
		name = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', '_', name).strip(' .')
		return name[:180] or 'download'

	def _unique_download_path(self, filename: str) -> Path:
		candidate = self.download_dir / filename
		counter = 1
		while candidate.exists() or candidate in self._reserved_download_paths:
			candidate = self.download_dir / f'{Path(filename).stem}_{counter}{Path(filename).suffix}'
			counter += 1
		self._reserved_download_paths.add(candidate)
		return candidate

	@staticmethod
	def _archive_member_parts(filename: str) -> tuple[str, ...] | None:
		assert isinstance(filename, str)
		normalised = filename.replace('\\', '/')
		if normalised.startswith('/') or re.match(r'^[A-Za-z]:', normalised):
			return None
		parts = tuple(part for part in normalised.split('/') if part not in {'', '.'})
		if not parts or any(part == '..' for part in parts):
			return None
		assert all(part not in {'', '.', '..'} for part in parts)
		return parts

	def _extract_download_archive(self, archive_path: Path) -> _ArchiveExtractionResult:
		assert archive_path.is_file()
		assert archive_path.parent.resolve(strict=True) == self.download_dir.resolve(strict=True)
		members: list[_ExtractedArchiveMember] = []
		warnings: list[str] = []
		warning_count = 0
		skipped_count = 0

		def add_warning(message: str) -> None:
			nonlocal warning_count
			warning_count += 1
			if len(warnings) < _MAX_ARCHIVE_WARNINGS:
				warnings.append(message[:500])

		if not zipfile.is_zipfile(archive_path):
			result = _ArchiveExtractionResult(
				status='failed',
				members=[],
				skipped_count=0,
				warnings=['Downloaded file has a .zip suffix but is not a valid ZIP archive'],
			)
			assert result.status == 'failed' and not result.members
			return result

		try:
			with zipfile.ZipFile(archive_path) as archive:
				directory_aliases: dict[tuple[str, ...], Path] = {}
				total_bytes = 0
				for info in archive.infolist():
					mode = info.external_attr >> 16
					file_type = stat.S_IFMT(mode)
					if info.create_system == 3 and file_type not in {0, stat.S_IFREG, stat.S_IFDIR}:
						skipped_count += 1
						add_warning(f'{info.filename}: non-regular archive member skipped')
						continue
					parts = self._archive_member_parts(info.filename)
					if parts is None:
						skipped_count += 1
						add_warning(f'{info.filename}: unsafe archive path skipped')
						continue
					target: Path | None = None
					temporary_path: Path | None = None
					try:
						is_directory = info.is_dir() or (
							info.create_system == 3 and file_type == stat.S_IFDIR
						)
						if is_directory:
							self._archive_member_directory(parts, directory_aliases)
							continue
						if len(members) >= _MAX_ARCHIVE_FILES:
							skipped_count += 1
							add_warning(f'{info.filename}: archive exceeds 1,000 file extraction limit')
							continue
						parent = self._archive_member_directory(parts[:-1], directory_aliases)
						if not parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True)):
							raise ValueError('archive member parent escaped the downloads directory')
						target = self._reserve_unique_archive_path(parent / parts[-1], kind='file')
						with tempfile.NamedTemporaryFile(dir=parent, prefix=f'.{target.name}.', suffix='.part', delete=False) as output:
							temporary_path = Path(output.name)
							with archive.open(info) as source:
								member_bytes = 0
								while chunk := source.read(1024 * 1024):
									projected_member_bytes = member_bytes + len(chunk)
									if projected_member_bytes > _MAX_ARCHIVE_MEMBER_BYTES:
										raise _ArchiveResourceLimitError('member exceeds 100 MiB extraction limit')
									projected_total_bytes = total_bytes + len(chunk)
									if projected_total_bytes > _MAX_ARCHIVE_TOTAL_BYTES:
										raise _ArchiveResourceLimitError('archive exceeds 250 MiB total extraction limit')
									written_bytes = output.write(chunk)
									member_bytes += written_bytes
									total_bytes += written_bytes
									if written_bytes != len(chunk):
										raise OSError(f'partial archive member write: {written_bytes} of {len(chunk)} bytes')
						os.link(temporary_path, target)
						temporary_path.unlink()
						temporary_path = None
						members.append(
							_ExtractedArchiveMember(
								archive_name=info.filename,
								relative_name=target.relative_to(self.download_dir).as_posix(),
								path=target,
								size_bytes=target.stat().st_size,
							)
						)
					except Exception as exc:
						if target is not None:
							self._reserved_download_paths.discard(target)
						skipped_count += 1
						add_warning(f'{info.filename}: extraction failed: {type(exc).__name__}: {exc}')
					finally:
						if temporary_path is not None:
							temporary_path.unlink(missing_ok=True)
		except Exception as exc:
			add_warning(f'ZIP extraction failed: {type(exc).__name__}: {exc}')

		if warning_count > _MAX_ARCHIVE_WARNINGS:
			omitted_count = warning_count - (_MAX_ARCHIVE_WARNINGS - 1)
			warnings[-1] = f'{omitted_count} additional warning(s) omitted'
		status: Literal['success', 'partial', 'failed']
		if warnings and not members:
			status = 'failed'
		elif warnings:
			status = 'partial'
		else:
			status = 'success'
		result = _ArchiveExtractionResult(status=status, members=members, skipped_count=skipped_count, warnings=warnings)
		assert len(result.members) <= _MAX_ARCHIVE_FILES
		assert all(
			member.path.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
			for member in result.members
		)
		return result

	def _archive_member_directory(
		self,
		parts: tuple[str, ...],
		aliases: dict[tuple[str, ...], Path],
	) -> Path:
		assert all(part not in {'', '.', '..'} for part in parts)
		assert self.download_dir.is_dir()
		parent = self.download_dir
		prefix: tuple[str, ...] = ()
		for part in parts:
			prefix += (part,)
			if prefix in aliases:
				parent = aliases[prefix]
				continue
			candidate = parent / part
			if candidate.is_symlink() or (candidate.exists() and not candidate.is_dir()):
				candidate = self._reserve_unique_archive_path(candidate, kind='directory')
				candidate.mkdir()
			else:
				candidate.mkdir(exist_ok=True)
			aliases[prefix] = candidate
			parent = candidate
		assert parent.is_dir()
		assert parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		return parent

	def _reserve_unique_archive_path(self, candidate: Path, *, kind: Literal['directory', 'file']) -> Path:
		assert candidate.parent.is_dir()
		assert candidate.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		counter = 1
		original = candidate
		while candidate.exists() or candidate.is_symlink() or candidate in self._reserved_download_paths:
			name = (
				f'{original.name}_{counter}'
				if kind == 'directory'
				else f'{original.stem}_{counter}{original.suffix}'
			)
			candidate = original.with_name(name)
			counter += 1
		self._reserved_download_paths.add(candidate)
		assert candidate in self._reserved_download_paths
		assert candidate.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		assert not candidate.exists() and not candidate.is_symlink()
		return candidate

	@staticmethod
	def _archive_extraction_summary(result: _ArchiveExtractionResult) -> str:
		assert result.skipped_count >= 0
		summary = (
			f'ZIP extraction {result.status}: {len(result.members)} file(s) extracted, '
			f'{result.skipped_count} member(s) skipped.'
		)
		if result.warnings:
			summary += '\nWarnings:\n' + '\n'.join(f'- {warning}' for warning in result.warnings)
		assert summary
		return summary

	@staticmethod
	def _document_response_extension(response: Response) -> str | None:
		try:
			if response.request.resource_type != 'document':
				return None
			headers = response.headers
		except Exception:
			return None
		content_disposition = headers.get('content-disposition', '').lower()
		if 'attachment' in content_disposition:
			# Playwright emits a Download for attachments; using that path avoids
			# saving the same response twice.
			return None
		filename_match = re.search(r'filename\*?\s*=\s*(?:UTF-8\'\')?"?([^";]+)', content_disposition, re.IGNORECASE)
		filename = unquote(filename_match.group(1).strip()) if filename_match else ''
		for candidate in (filename, unquote(urlsplit(response.url).path)):
			suffix = Path(candidate).suffix.lower()
			if suffix in _DOCUMENT_EXTENSIONS:
				return suffix
		content_type = headers.get('content-type', '').split(';', 1)[0].strip().lower()
		return _DOCUMENT_MIME_EXTENSIONS.get(content_type)

	@classmethod
	def _response_filename(cls, response: Response, extension: str) -> str:
		disposition = response.headers.get('content-disposition', '')
		match = re.search(r"filename\*\s*=\s*(?:UTF-8'')?([^;]+)", disposition, re.IGNORECASE)
		if match is None:
			match = re.search(r'filename\s*=\s*"?([^";]+)', disposition, re.IGNORECASE)
		if match is not None:
			name = unquote(match.group(1).strip().strip('"\''))
		else:
			name = unquote(Path(urlsplit(response.url).path).name)
		name = cls._safe_download_name(name or f'document{extension}')
		if Path(name).suffix.lower() != extension:
			name = f'{name}{extension}'
		return name

	@staticmethod
	def _extract_download_text(path: Path) -> tuple[str, bool]:
		if path.stat().st_size > _MAX_DOWNLOAD_BYTES:
			return f'[text extraction omitted: file exceeds {_MAX_DOWNLOAD_BYTES} bytes]', True
		extension = path.suffix.lower()
		text = ''
		if extension == '.pdf':
			from pypdf import PdfReader

			reader = PdfReader(str(path))
			text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
		elif extension == '.docx':
			from docx import Document

			document = Document(str(path))
			parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
			for table in document.tables:
				parts.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
			text = '\n'.join(parts)
		elif extension == '.xlsx':
			text = BrowserRuntime._extract_xlsx_text(path)
		elif extension == '.xls':
			text = BrowserRuntime._extract_xls_text(path)
		elif extension == '.zip':
			text = BrowserRuntime._extract_zip_text(path)
		elif extension in {'.txt', '.text', '.md', '.csv', '.tsv', '.json', '.jsonl', '.xml', '.html', '.htm'}:
			raw = path.read_bytes()
			text = raw.decode('utf-8-sig', errors='replace')
			if extension in {'.html', '.htm'}:
				parser = _VisibleTextParser()
				parser.feed(text)
				text = '\n'.join(parser.parts)
		else:
			return '', False
		truncated = len(text) > _MAX_DOWNLOAD_TEXT
		return text[:_MAX_DOWNLOAD_TEXT], truncated

	@staticmethod
	def _extract_zip_text(path: Path) -> str:
		"""Read bounded text/CSV members from an official data ZIP in memory."""
		parts: list[str] = []
		used = 0
		supported = {'.txt', '.text', '.md', '.csv', '.tsv', '.json', '.jsonl', '.xml'}
		with zipfile.ZipFile(path) as archive:
			members = sorted((item for item in archive.infolist() if not item.is_dir()), key=lambda item: item.filename)
			for item in members[:100]:
				if Path(item.filename).suffix.lower() not in supported or item.file_size > _MAX_DOWNLOAD_BYTES:
					continue
				remaining = _MAX_DOWNLOAD_TEXT - used
				if remaining <= 0:
					break
				with archive.open(item) as member:
					raw = member.read(remaining + 1)
				text = raw.decode('utf-8-sig', errors='replace')
				chunk = f'[{item.filename}]\n{text}'
				parts.append(chunk[:remaining])
				used += min(len(chunk), remaining)
		return '\n'.join(parts)

	@staticmethod
	def _extract_xls_text(path: Path) -> str:
		"""Extract legacy BIFF Excel sheets while preserving row columns."""
		import xlrd

		workbook = xlrd.open_workbook(str(path), on_demand=True)
		lines: list[str] = []
		try:
			for sheet in workbook.sheets():
				lines.append(f'[{sheet.name}]')
				for row_index in range(sheet.nrows):
					values: list[str] = []
					for column_index in range(sheet.ncols):
						cell = sheet.cell(row_index, column_index)
						if cell.ctype == xlrd.XL_CELL_DATE:
							value = xlrd.xldate_as_datetime(float(cell.value), workbook.datemode).isoformat(sep=' ')
						elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
							value = str(int(float(cell.value)))
						else:
							value = str(cell.value)
						values.append(value)
					while values and not values[-1]:
						values.pop()
					lines.append('\t'.join(values))
		finally:
			workbook.release_resources()
		return '\n'.join(lines)

	@staticmethod
	def _extract_xlsx_text(path: Path) -> str:
		namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
		with zipfile.ZipFile(path) as archive:
			shared: list[str] = []
			if 'xl/sharedStrings.xml' in archive.namelist():
				root = ElementTree.fromstring(archive.read('xl/sharedStrings.xml'))
				for item in root.findall(f'{namespace}si'):
					shared.append(''.join(node.text or '' for node in item.iter(f'{namespace}t')))
			lines: list[str] = []
			worksheets = sorted(name for name in archive.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name))
			for worksheet in worksheets:
				lines.append(f'[{Path(worksheet).stem}]')
				root = ElementTree.fromstring(archive.read(worksheet))
				for row in root.iter(f'{namespace}row'):
					values: list[str] = []
					for cell in row.findall(f'{namespace}c'):
						cell_type = cell.get('t')
						value_node = cell.find(f'{namespace}v')
						if cell_type == 'inlineStr':
							inline = cell.find(f'{namespace}is')
							value = (
								''.join(node.text or '' for node in inline.iter(f'{namespace}t')) if inline is not None else ''
							)
						else:
							value = value_node.text if value_node is not None and value_node.text else ''
							if cell_type == 's' and value.isdigit() and int(value) < len(shared):
								value = shared[int(value)]
						values.append(value)
					lines.append('\t'.join(values))
			return '\n'.join(lines)
