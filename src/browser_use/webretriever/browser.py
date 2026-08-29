"""Playwright action runtime with read-only CDP DOM inspection.

The competition supplies an already-running browser over CDP. Connection
ownership deliberately stays with the runner; :class:`BrowserRuntime` owns
only the pages it creates inside the supplied ``BrowserContext``. CDP is used
only to inspect DOM identity, geometry, snapshots, accessibility, and listener
metadata; state-changing input remains Playwright mouse and keyboard actions.
"""

from __future__ import annotations

import asyncio
import base64
import contextlib
import copy
import hashlib
import io
import json
import logging
import math
import os
import re
import secrets
import stat
import tempfile
import time
import unicodedata
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable, Coroutine, Literal, Mapping, Sequence, cast
from urllib.parse import parse_qsl, unquote, urlsplit
from xml.etree import ElementTree

from playwright.async_api import BrowserContext, CDPSession, Dialog, Download, Frame, Locator, Page, Request, Response
from playwright.async_api import Error as PlaywrightError
from playwright.async_api import TimeoutError as PlaywrightTimeoutError
from pydantic import BaseModel, ConfigDict

from browser_use.webretriever.dom_collector import CdpCollectionError, collect_interactive_elements
from browser_use.webretriever.models import WebRetrieverActionResult

if TYPE_CHECKING:
	from browser_use.webretriever.models import AgentDecision


__all__ = [
	'BrowserObservation',
	'BrowserRuntime',
	'ElementRef',
	'cdp_headers_for_url',
	'is_sec_url',
	'is_forbidden_search_url',
	'redact_cdp_url',
]


_CDP_TOKEN_KEY = 'access_token'
_REDACTED = '<redacted>'
_INTERACTIVE_ATTRIBUTE = 'data-webretriever-index'
_OVERLAY_CLASS = '__webretriever_overlay'
_MAX_ELEMENT_TEXT = 240
_MAX_PAGE_TEXT = 80_000
_MAX_RENDERED_TEXT = 100_000
_MAX_RESPONSE_BODY_BYTES = 128 * 1024
_NETWORK_SEARCH_TIMEOUT_SECONDS = 30.0
_NETWORK_SEARCH_NODE_HEAP_MIB = 512
_TEXT_SEARCH_TIMEOUT_SECONDS = 30.0
_NETWORK_BODY_PAGE_CHARACTERS = 60_000
_MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
_MAX_DOWNLOAD_TEXT = 250_000
_DOWNLOAD_PREVIEW_HEAD_CHARS = 600
_DOWNLOAD_PREVIEW_TAIL_CHARS = 600
_DEFAULT_RUNTIME_CLEANUP_TIMEOUT_SECONDS = 60.0
_NAVIGATION_FIRST_OBSERVATION_TIMEOUT_MS = 10_000
_NAVIGATION_HARD_TIMEOUT_MS = 120_000
_INITIAL_NAVIGATION_RETRY_DELAY_SECONDS = 0.25
_TRANSIENT_INITIAL_NAVIGATION_ERROR_CODES = frozenset(
	{
		'net::err_connection_closed',
		'net::err_connection_reset',
		'net::err_connection_timed_out',
		'net::err_network_changed',
	}
)
_DOWNLOAD_HARD_TIMEOUT_MS = 10 * 60 * 1000
_MAX_ARCHIVE_FILES = 1_000
_DOWNLOAD_PLACEHOLDER_URL = ':'
_MAX_ARCHIVE_MEMBER_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_TOTAL_BYTES = 250 * 1024 * 1024
_MAX_ARCHIVE_WARNINGS = 100
_MAX_RECENT_DIALOGS = 8
_MAX_DIALOG_MESSAGE = 2_000
_MAX_ACTION_RESULT_OUTPUT = 20_000
_SCREENSHOT_RETRY_DELAY_SECONDS = 3 * 60
_STATE_CHANGING_ACTIONS = frozenset(
	{
		'click',
		'double_click',
		'drag',
		'hover',
		'hover_xy',
		'back',
		'navigate',
		'press',
		'scroll',
		'select',
		'tab',
		'type',
		'xy',
	}
)
_CONTENT_ACTIONS = frozenset({'find', 'read', 'inspect_network', 'calculate'})
_PNG_SIGNATURE = b'\x89PNG\r\n\x1a\n'
_ANNOTATION_LABEL_GAP = 2
_DOCUMENT_MIME_EXTENSIONS = {
	'application/json': '.json',
	'application/ld+json': '.json',
	'application/xml': '.xml',
	'application/pdf': '.pdf',
	'application/zip': '.zip',
	'application/x-zip-compressed': '.zip',
	'application/octet-stream': '.bin',
	'application/vnd.openxmlformats-officedocument.wordprocessingml.document': '.docx',
	'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': '.xlsx',
	'application/vnd.ms-excel': '.xls',
	'text/csv': '.csv',
	'application/csv': '.csv',
	'text/tab-separated-values': '.tsv',
	'text/tsv': '.tsv',
	'text/plain': '.txt',
	'text/xml': '.xml',
}
_DOCUMENT_EXTENSIONS = frozenset(_DOCUMENT_MIME_EXTENSIONS.values())
_STRUCTURED_DOWNLOAD_EXTENSIONS = frozenset({'.json', '.csv', '.tsv', '.xlsx', '.zip'})


class _ScreenshotFallbackError(RuntimeError):
	"""Raised when recoverable Playwright screenshot failure recovery through CDP also fails."""


class _ArchiveResourceLimitError(RuntimeError):
	"""Raised when a ZIP member exceeds an extraction resource budget."""


@dataclass(slots=True)
class _PendingNavigation:
	page: Page
	url: str
	started_at: float
	task: asyncio.Task[Any]


class _ExtractedArchiveMember(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	archive_name: str
	relative_name: str
	path: Path
	size_bytes: int


class _ArchiveExtractionResult(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	status: Literal['success', 'partial', 'failed']
	members: list[_ExtractedArchiveMember]
	skipped_count: int
	warnings: list[str]


class _ArchiveMemberDownload(BaseModel):
	model_config = ConfigDict(extra='forbid', frozen=True)

	timestamp: float
	url: str
	suggested_filename: str
	filename: str
	path: str
	size_bytes: int
	source: Literal['archive_member']
	source_archive: str
	text: str
	text_truncated: bool
	text_extraction_error: str | None = None


def _consume_detached_task_result(task: asyncio.Future[Any]) -> None:
	"""Retrieve a late cancellation-resistant result without warning."""

	if task.cancelled():
		return
	try:
		task.exception()
	except BaseException:
		pass


def cdp_headers_for_url(cdp_url: str) -> dict[str, str]:
	"""Return the cloud-sandbox authentication header encoded in a CDP URL."""

	try:
		query = parse_qsl(urlsplit(cdp_url).query, keep_blank_values=True)
	except (TypeError, ValueError):
		return {}
	for key, value in query:
		if key == _CDP_TOKEN_KEY and value:
			return {'X-Access-Token': value}
	return {}


_TOKEN_RE = re.compile(r'([?&]access_token=)[^&#]*', re.IGNORECASE)


def redact_cdp_url(cdp_url: str) -> str:
	"""Redact CDP credentials without otherwise rewriting the URL."""

	return _TOKEN_RE.sub(rf'\1{_REDACTED}', str(cdp_url))


def _host_matches(host: str, domain: str) -> bool:
	return host == domain or host.endswith(f'.{domain}')


def is_sec_url(url: str) -> bool:
	"""Whether *url* belongs to SEC.gov or one of its subdomains."""

	if not isinstance(url, str) or not url.strip():
		return False
	candidate = url.strip()
	if '://' not in candidate and not candidate.startswith('//'):
		candidate = f'//{candidate}'
	try:
		host = (urlsplit(candidate).hostname or '').rstrip('.').lower()
	except ValueError:
		return False
	return _host_matches(host, 'sec.gov')


def is_forbidden_search_url(url: str) -> bool:
	"""Whether *url* targets a known public web-search/answer engine.

	The match is hostname based (never substring based), so a site URL such as
	``example.test/?next=google.com`` is not incorrectly rejected.
	"""

	if not isinstance(url, str) or not url.strip():
		return False
	candidate = url.strip()
	if '://' not in candidate and not candidate.startswith('//'):
		candidate = f'//{candidate}'
	try:
		parsed = urlsplit(candidate)
		host = (parsed.hostname or '').rstrip('.').lower()
	except ValueError:
		return False
	if not host:
		return False

	# Google properties such as Drive and Maps are not search engines.  Match
	# Google's country-domain search frontends only when their URL is search-like.
	if re.fullmatch(r'(?:www\.)?google\.[a-z.]{2,}', host):
		path = parsed.path.rstrip('/').lower() or '/'
		if path in {'/', '/search', '/webhp', '/custom', '/m'} or 'q=' in parsed.query.lower():
			return True

	blocked_domains = (
		'bing.com',
		'duckduckgo.com',
		'baidu.com',
		'yandex.com',
		'yandex.ru',
		'ask.com',
		'ecosia.org',
		'qwant.com',
		'startpage.com',
		'dogpile.com',
		'mojeek.com',
		'search.brave.com',
		'perplexity.ai',
		'you.com',
		'so.com',
		'sm.cn',
	)
	if any(_host_matches(host, domain) for domain in blocked_domains):
		return True
	if re.fullmatch(r'(?:search|www)\.yahoo\.[a-z.]{2,}', host):
		if host.startswith('search.') or parsed.path.lower().startswith(('/search', '/s')) or parsed.query:
			return True
	if _host_matches(host, 'search.aol.com') or (_host_matches(host, 'aol.com') and parsed.path.lower().startswith('/search')):
		return True
	if _host_matches(host, 'search.naver.com') or _host_matches(host, 'search.sogou.com'):
		return True
	if _host_matches(host, 'sogou.com') and parsed.path.lower().startswith(('/web', '/search')):
		return True
	return False


@dataclass(slots=True)
class ElementRef:
	"""A visible interactive element indexed in the latest observation."""

	index: int
	tag: str
	text: str = ''
	role: str = ''
	name: str = ''
	placeholder: str = ''
	href: str = ''
	input_type: str = ''
	checked: str = ''
	frame_index: int = 0
	frame_url: str = ''
	x: float = 0.0
	y: float = 0.0
	width: float = 0.0
	height: float = 0.0
	selector: str = ''
	backend_node_id: int = 0
	frame_id: str = ''
	signals: tuple[str, ...] = ()

	def render_text(self) -> str:
		parts = [f'[{self.index}]', self.tag]
		if self.role:
			parts.append(f'role={self.role}')
		if self.input_type:
			parts.append(f'type={self.input_type}')
		if self.checked:
			parts.append(f'checked={self.checked}')
		if self.name:
			parts.append(f'name={json.dumps(self.name, ensure_ascii=False)}')
		if self.text and self.text != self.name:
			parts.append(f'text={json.dumps(self.text, ensure_ascii=False)}')
		if self.placeholder:
			parts.append(f'placeholder={json.dumps(self.placeholder, ensure_ascii=False)}')
		if self.href:
			parts.append(f'href={json.dumps(self.href, ensure_ascii=False)}')
		if self.frame_index:
			parts.append(f'frame={self.frame_index}')
		parts.append(f'box=({self.x:.0f},{self.y:.0f},{self.width:.0f},{self.height:.0f})')
		return ' '.join(parts)

	def to_dict(self) -> dict[str, Any]:
		return asdict(self)


@dataclass(slots=True)
class BrowserObservation:
	"""Multimodal state made available to the decision model at one step."""

	screenshot: bytes
	url: str
	title: str
	tabs: list[dict[str, Any]]
	viewport_width: int
	viewport_height: int
	elements: list[ElementRef]
	page_text: str
	recent_network: list[dict[str, Any]]
	downloads: list[dict[str, Any]]
	step: int = 0
	screenshot_path: str = ''
	visual_screenshot_path: str = ''

	@staticmethod
	def download_prompt_record(download: Mapping[str, Any]) -> dict[str, Any]:
		"""Return complete download metadata plus a bounded content preview.

		The raw ``text`` field can contain hundreds of thousands of characters and
		must not be serialized wholesale into a model prompt.  Every other field is
		metadata and is retained verbatim (with non-JSON values stringified); the
		preview is deliberately split into head and tail so both document identity
		and late-page evidence remain visible.
		"""

		def json_safe(value: Any) -> Any:
			if isinstance(value, Mapping):
				return {str(key): json_safe(child) for key, child in value.items()}
			if isinstance(value, (list, tuple)):
				return [json_safe(child) for child in value]
			try:
				json.dumps(value, ensure_ascii=False)
			except (TypeError, ValueError):
				return str(value)
			return value

		text = str(download.get('text', ''))
		result = {str(key): json_safe(value) for key, value in download.items() if key != 'text'}
		result.update(
			{
				'content_characters': len(text),
				'content_preview_truncated': len(text) > _DOWNLOAD_PREVIEW_HEAD_CHARS + _DOWNLOAD_PREVIEW_TAIL_CHARS,
				'content_preview_head': text[:_DOWNLOAD_PREVIEW_HEAD_CHARS],
				'content_preview_tail': (
					text[-_DOWNLOAD_PREVIEW_TAIL_CHARS:]
					if len(text) > _DOWNLOAD_PREVIEW_HEAD_CHARS + _DOWNLOAD_PREVIEW_TAIL_CHARS
					else ''
				),
			}
		)
		return result

	def render_text(self) -> str:
		"""Render bounded, model-friendly browser state (image bytes excluded)."""

		def clip(value: Any, limit: int) -> str:
			text = str(value)
			return text if len(text) <= limit else text[: limit - 1] + '…'

		def section(name: str, lines: list[str], budget: int) -> str:
			body = '\n'.join(lines) or '  (none)'
			return f'{name}:\n{clip(body, budget)}'

		tab_lines = [
			f'  {tab["index"]}: {clip(tab.get("title", ""), 300)!r} {clip(tab.get("url", ""), 1_500)}'
			+ (' [active]' if tab.get('active') else '')
			for tab in self.tabs
		]
		element_lines = [f'  {clip(element.render_text(), 1_200)}' for element in self.elements]
		network_lines = []
		for item in self.recent_network[-12:]:
			status = item.get('status', 'pending')
			network_lines.append(f'  {clip(item.get("method", ""), 20)} {clip(item.get("url", ""), 1_500)} [{status}]')
		download_items = sorted(self.downloads, key=lambda item: float(item.get('timestamp', 0)))
		download_lines = [
			f'  {json.dumps(self.download_prompt_record(item), ensure_ascii=False, separators=(",", ":"))}'
			for item in download_items
		]
		# Put compact, high-value evidence before potentially very long page text,
		# then budget the page section independently so downloads/network cannot be
		# cut off merely because a page has a huge body.
		sections = [
			f'URL: {clip(self.url, 4_000)}',
			f'Title: {clip(self.title, 1_000)}',
			f'Viewport: {self.viewport_width}x{self.viewport_height}',
			section('Tabs', tab_lines, 8_000),
			section('Interactive elements', element_lines, 40_000),
			section('Recent XHR/Fetch', network_lines, 14_000),
			section('Downloads', download_lines, 36_000),
		]
		prefix = '\n\n'.join(sections)
		page_budget = max(0, _MAX_RENDERED_TEXT - len(prefix) - len('\n\nPage text:\n'))
		return (prefix + '\n\nPage text:\n' + self.page_text[:page_budget])[:_MAX_RENDERED_TEXT]


@dataclass(slots=True)
class _ElementBinding:
	frame: Frame
	selector: str = ''
	backend_node_id: int = 0
	cdp_target: Page | Frame | None = None
	coordinate_frame: Frame | None = None
	frame_id: str = ''
	read_text: str = ''
	options: tuple[tuple[str, str], ...] = ()


class _VisibleTextParser(HTMLParser):
	def __init__(self) -> None:
		super().__init__()
		self.parts: list[str] = []
		self._ignored_depth = 0

	def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
		if tag in {'script', 'style', 'noscript'}:
			self._ignored_depth += 1

	def handle_endtag(self, tag: str) -> None:
		if tag in {'script', 'style', 'noscript'} and self._ignored_depth:
			self._ignored_depth -= 1

	def handle_data(self, data: str) -> None:
		if not self._ignored_depth and data.strip():
			self.parts.append(data.strip())


_MARK_ELEMENTS_JS = r"""
({start, markerAttribute, overlayClass, frameIndex, maxText}) => {
  document.querySelectorAll('.' + overlayClass).forEach((node) => node.remove());
  document.querySelectorAll('[' + markerAttribute + ']').forEach((node) => node.removeAttribute(markerAttribute));
  const selector = [
    'a[href]', 'button', 'input:not([type="hidden"])', 'textarea', 'select',
    '[contenteditable="true"]', 'summary', 'label[for]',
    '[role="button"]', '[role="link"]', '[role="checkbox"]', '[role="radio"]',
    '[role="menuitem"]', '[role="option"]', '[role="tab"]', '[role="switch"]',
    '[role="textbox"]', '[role="combobox"]', '[tabindex]:not([tabindex="-1"])'
  ].join(',');
  const candidates = Array.from(document.querySelectorAll(selector));
  const results = [];
  let next = start;
  for (const element of candidates) {
    if (!(element instanceof HTMLElement) && !(element instanceof SVGElement)) continue;
    if (element.matches(':disabled,[aria-disabled="true"],[inert]')) continue;
    const style = getComputedStyle(element);
    const rect = element.getBoundingClientRect();
    if (style.display === 'none' || style.visibility === 'hidden' || Number(style.opacity) === 0) continue;
    if (rect.width < 2 || rect.height < 2) continue;
    if (rect.bottom <= 0 || rect.right <= 0 || rect.top >= innerHeight || rect.left >= innerWidth) continue;
    if (typeof element.checkVisibility === 'function' && !element.checkVisibility({checkOpacity: true, checkVisibilityCSS: true})) continue;
    const index = next++;
    element.setAttribute(markerAttribute, String(index));
    const isSelection = element.matches('input[type="checkbox"], input[type="radio"], [role="checkbox"], [role="radio"], [role="switch"]');
    const labelledBy = (element.getAttribute('aria-labelledby') || '').split(/\s+/).filter(Boolean)
      .map(id => document.getElementById(id)?.innerText || document.getElementById(id)?.textContent || '')
      .join(' ');
    const explicitLabel = element.id
      ? Array.from(document.querySelectorAll('label[for]')).find(label => label.getAttribute('for') === element.id)?.innerText || ''
      : '';
    const wrappingLabel = element.closest('label')?.innerText || '';
    const text = (element.innerText || element.value || element.getAttribute('aria-label') || '').replace(/\s+/g, ' ').trim();
    const name = (isSelection
      ? (element.getAttribute('aria-label') || labelledBy || explicitLabel || wrappingLabel || text || element.value || '')
      : (element.getAttribute('aria-label') || element.getAttribute('title') || text)).replace(/\s+/g, ' ').trim();
    const checked = isSelection
      ? ('checked' in element ? (element.checked ? 'true' : 'false') : (element.getAttribute('aria-checked') || ''))
      : '';
    const overlay = document.createElement('div');
    overlay.className = overlayClass;
    Object.assign(overlay.style, {
      position: 'fixed', left: `${Math.max(0, rect.left)}px`, top: `${Math.max(0, rect.top)}px`,
      width: `${Math.max(1, Math.min(rect.width, innerWidth - Math.max(0, rect.left)))}px`,
      height: `${Math.max(1, Math.min(rect.height, innerHeight - Math.max(0, rect.top)))}px`,
      border: '2px solid #ff2d55', background: 'rgba(255,45,85,.04)', boxSizing: 'border-box',
      pointerEvents: 'none', zIndex: '2147483646'
    });
    const badge = document.createElement('span');
    badge.textContent = String(index);
    Object.assign(badge.style, {
      position: 'absolute', left: '-2px', top: '-18px', background: '#ff2d55', color: '#fff',
      font: 'bold 12px/16px sans-serif', padding: '0 3px', minWidth: '16px', textAlign: 'center'
    });
    overlay.appendChild(badge);
    (document.documentElement || document.body).appendChild(overlay);
    results.push({
      index, tag: element.tagName.toLowerCase(), text: text.slice(0, maxText),
      role: element.getAttribute('role') || (isSelection ? (element.getAttribute('type') || '') : ''), name: name.slice(0, maxText),
      placeholder: (element.getAttribute('placeholder') || '').slice(0, maxText),
      href: String(element.href || '').slice(0, 2000), input_type: element.getAttribute('type') || '', checked, frame_index: frameIndex,
      x: rect.x, y: rect.y, width: rect.width, height: rect.height
    });
  }
  return results;
}
"""


def _place_annotation_label(
	*,
	anchor_left: float,
	anchor_top: float,
	label_width: int,
	label_height: int,
	image_width: int,
	image_height: int,
	occupied: Sequence[tuple[int, int, int, int]],
) -> tuple[int, int]:
	"""Place one numbered screenshot label without covering an earlier label.

	Labels start at an element's top-left corner.  When that slot is occupied,
	they are packed horizontally from left to right (the model still sees the
	original element index).  If a row reaches the image edge, the next row is
	used.  The bounded grid search keeps labels inside the screenshot even on a
	dense page; the final clamped slot is a safe fallback if the image cannot fit
	all labels without overlap.
	"""
	if label_width <= 0 or label_height <= 0 or image_width <= 0 or image_height <= 0:
		return 0, 0
	max_left = max(0, image_width - label_width)
	max_top = max(0, image_height - label_height)
	base_left = min(max(0, int(round(anchor_left))), max_left)
	base_top = min(max(0, int(round(anchor_top))), max_top)
	row_step = max(1, label_height + _ANNOTATION_LABEL_GAP)
	row_count = max(1, image_height // row_step + 2)

	def overlaps(left: int, top: int) -> list[tuple[int, int, int, int]]:
		right = left + label_width
		bottom = top + label_height
		return [rect for rect in occupied if left < rect[2] and right > rect[0] and top < rect[3] and bottom > rect[1]]

	for row in range(row_count):
		top = min(max_top, base_top + row * row_step)
		left = base_left
		# At most one rightward jump per existing label is needed.  This is
		# deliberately bounded because the screenshot is a diagnostic artifact,
		# not a reason to delay the browser observation.
		for _ in range(len(occupied) + 1):
			conflicts = overlaps(left, top)
			if not conflicts:
				return left, top
			next_left = max(rect[2] + _ANNOTATION_LABEL_GAP for rect in conflicts)
			if next_left > max_left:
				break
			left = next_left

	# There is no free slot in the bounded grid.  Keep the label visible and
	# deterministic rather than allowing it to be clipped by the image edge.
	return base_left, max_top


_CLEAR_MARKERS_JS = r"""
({markerAttribute, overlayClass, removeAttributes}) => {
  document.querySelectorAll('.' + overlayClass).forEach((node) => node.remove());
  if (removeAttributes) {
    document.querySelectorAll('[' + markerAttribute + ']').forEach((node) => node.removeAttribute(markerAttribute));
  }
}
"""

_TARGET_STATE_JS = r"""
function(element) {
  const target = element || this;
  const rect = target.getBoundingClientRect();
  const text = (target.innerText || target.textContent || '').replace(/\s+/g, ' ').trim();
  const ariaChecked = target.getAttribute('aria-checked');
  return {
    tag: target.tagName.toLowerCase(),
    value: 'value' in target ? String(target.value ?? '') : null,
    checked: 'checked' in target ? Boolean(target.checked) : ariaChecked,
    selected: target.getAttribute('aria-selected'),
    expanded: target.getAttribute('aria-expanded'),
    disabled: 'disabled' in target ? Boolean(target.disabled) : null,
    text: text.slice(0, 500),
    visible: rect.bottom > 0 && rect.right > 0
      && rect.top < window.innerHeight && rect.left < window.innerWidth
  };
}
"""

_SELECT_OPTIONS_JS = r"""
function(element) {
  const target = element || this;
  if (!target || target.tagName.toLowerCase() !== 'select') return null;
  return Array.from(target.options || []).map(option => ({
    value: String(option.value ?? ''),
    label: String(option.label ?? option.textContent ?? '').replace(/\s+/g, ' ').trim(),
  }));
}
"""


class BrowserRuntime:
	"""A small async runtime that performs every browser operation with Playwright."""

	def __init__(
		self,
		context: BrowserContext,
		task_dir: Path,
		logger: logging.Logger,
		*,
		navigation_timeout_ms: int = _NAVIGATION_HARD_TIMEOUT_MS,
		navigation_first_observation_timeout_ms: int = _NAVIGATION_FIRST_OBSERVATION_TIMEOUT_MS,
		download_timeout_ms: int = _DOWNLOAD_HARD_TIMEOUT_MS,
		action_timeout_ms: int = 30_000,
		screenshot_timeout_ms: int = 20_000,
		cdp_screenshot_timeout_ms: int = 20_000,
		screenshot_retry_delay_seconds: float = _SCREENSHOT_RETRY_DELAY_SECONDS,
		max_response_body_bytes: int = _MAX_RESPONSE_BODY_BYTES,
		declared_user_agent: str | None = None,
		task_identity: Mapping[str, Any] | None = None,
		before_close: Callable[[], Coroutine[Any, Any, Any]] | None = None,
	) -> None:
		if not 0 < navigation_first_observation_timeout_ms <= navigation_timeout_ms:
			raise ValueError('navigation_first_observation_timeout_ms must be in (0, navigation_timeout_ms]')
		if download_timeout_ms <= 0:
			raise ValueError('download_timeout_ms must be greater than 0')
		if screenshot_retry_delay_seconds < 0:
			raise ValueError('screenshot_retry_delay_seconds must be non-negative')
		self.context = context
		self.task_dir = Path(task_dir)
		self.logger = logger
		self.navigation_timeout_ms = navigation_timeout_ms
		self.navigation_first_observation_timeout_ms = navigation_first_observation_timeout_ms
		self.download_timeout_ms = download_timeout_ms
		self.action_timeout_ms = action_timeout_ms
		self.screenshot_timeout_ms = screenshot_timeout_ms
		self.cdp_screenshot_timeout_ms = cdp_screenshot_timeout_ms
		self.screenshot_retry_delay_seconds = screenshot_retry_delay_seconds
		self.max_response_body_bytes = max(0, max_response_body_bytes)
		self.declared_user_agent = declared_user_agent
		self.task_identity = dict(task_identity) if isinstance(task_identity, Mapping) else None
		self._before_close = before_close

		self.page: Page | None = None
		self.website = ''
		self.visited_urls: list[str] = []
		self.all_requests: list[dict[str, Any]] = []
		# ``all_requests`` is the competition-compatible XHR/Fetch capture.  Keep a
		# separate all-resource cache for chart-request discovery so broad capture
		# does not change the evaluator-facing capture.json contract.
		self.network_requests: list[dict[str, Any]] = []
		self.downloads: list[dict[str, Any]] = []

		self._owned_pages: list[Page] = []
		self._page_handlers: dict[int, list[tuple[str, Callable[..., Any]]]] = {}
		self._context_handlers: list[tuple[str, Callable[..., Any]]] = []
		self._request_entries: dict[int, dict[str, Any]] = {}
		self._request_entries_by_id: dict[int, dict[str, Any]] = {}
		self._request_ids_by_entry: dict[int, int] = {}
		self._network_request_entries: dict[int, dict[str, Any]] = {}
		self._network_entries_by_id: dict[int, dict[str, Any]] = {}
		self._network_responses: dict[int, Response] = {}
		self._network_page_cursors: dict[str, tuple[int, int, str, int]] = {}
		self._next_network_request_id = 0
		self._page_ids: dict[int, int] = {}
		self._page_document_generations: dict[int, int] = {}
		self._element_bindings: dict[int, _ElementBinding] = {}
		self._legacy_marker_bindings_active = False
		self._last_safe_urls: dict[int, str] = {}
		self._background_tasks: set[asyncio.Task[Any]] = set()
		self._download_tasks: set[asyncio.Task[Any]] = set()
		self._pending_navigations: dict[int, _PendingNavigation] = {}
		self._navigation_notices: list[dict[str, Any]] = []
		self._policy_tasks: set[asyncio.Task[Any]] = set()
		self._dialog_tasks: set[asyncio.Task[Any]] = set()
		self._screenshot_recovery_tasks: set[asyncio.Task[Any]] = set()
		self._rollback_pages: set[int] = set()
		self._captured_document_responses: set[tuple[str, int]] = set()
		self._download_urls_seen: set[str] = set()
		self._reserved_download_paths: set[Path] = set()
		# Native JavaScript dialogs are not part of DOM text or page screenshots.
		# Retain a bounded, task-local transcript so form-validation alerts reach
		# the decision model through both the action result and the next observation.
		self._recent_dialogs: list[dict[str, Any]] = []
		self._next_dialog_id = 0
		self._started = False
		self._closed = False
		self._cleanup_diagnostics: dict[str, Any] = {'status': 'not_started', 'residual_tasks': {}}
		self._execute_lock = asyncio.Lock()
		self._archive_extraction_lock = asyncio.Lock()

		self.trajectory_dir = self.task_dir / 'trajectory'
		self.trajectory_visual_dir = self.task_dir / 'trajectory_visual'
		self.download_dir = self.task_dir / 'downloads'
		self.trajectory_dir.mkdir(parents=True, exist_ok=True)
		self.trajectory_visual_dir.mkdir(parents=True, exist_ok=True)
		self.download_dir.mkdir(parents=True, exist_ok=True)

	async def __aenter__(self) -> BrowserRuntime:
		return self

	async def __aexit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
		await self.close()

	async def start(self, website: str) -> Page:
		"""Create an owned page and navigate to the exact task start URL."""

		if self._closed:
			raise RuntimeError('BrowserRuntime is closed')
		if self._started:
			raise RuntimeError('BrowserRuntime.start() may only be called once')
		self._validate_navigation_url(website)
		if is_forbidden_search_url(website):
			raise ValueError('The designated start URL is a prohibited external search engine')

		self.website = website
		self._install_context_listeners()
		page = await self.context.new_page()
		self._register_page(page, make_active=True)
		await self._configure_owned_page(page)
		self._started = True
		self.logger.info('Opening exact task start URL: %s', redact_cdp_url(website))
		try:
			download_started = await self._goto_exact(page, website)
		except PlaywrightError as exc:
			if page.is_closed() or not self._is_transient_initial_navigation_error(exc):
				raise
			self._record_navigation_notice(page, website, 'initial_navigation_retry', str(exc))
			self.logger.warning(
				'Initial navigation failed with a transient transport error; retrying once: %s',
				redact_cdp_url(str(exc)),
			)
			await asyncio.sleep(_INITIAL_NAVIGATION_RETRY_DELAY_SECONDS)
			download_started = await self._goto_exact(page, website)
		self._record_url(website if download_started or id(page) in self._pending_navigations else page.url, unless_last=True)
		if not is_forbidden_search_url(page.url):
			self._last_safe_urls[id(page)] = page.url
		await self._enforce_search_policy()
		return page

	async def observe(self, step: int) -> BrowserObservation:
		"""Capture raw/annotated screenshots plus a cross-frame textual state."""

		self._ensure_started()
		await self._drain_dialog_tasks()
		await self._enforce_search_policy()
		await self._drain_dialog_tasks()
		await self._settle_pending_navigations()
		await self._drain_download_tasks()
		page = await self._active_page_for_observation()
		await self._clear_markers(remove_attributes=True)

		step_name = self._safe_step_name(step)
		raw_path = self.trajectory_dir / f'{step_name}.png'
		visual_path = self.trajectory_visual_dir / f'{step_name}.png'
		raw_screenshot: bytes | None = None
		try:
			page, raw_screenshot = await self._capture_observation_screenshot(page, raw_path)
		except asyncio.CancelledError:
			raise
		except Exception as first_capture_error:
			self.logger.warning(
				'Observation screenshot failed for step %s; retrying the full capture chain after %gs: %s',
				step_name,
				self.screenshot_retry_delay_seconds,
				first_capture_error,
			)
			if self.screenshot_retry_delay_seconds:
				await asyncio.sleep(self.screenshot_retry_delay_seconds)
			try:
				page, raw_screenshot = await self._capture_observation_screenshot(page, raw_path)
			except asyncio.CancelledError:
				raise
			except Exception as retry_capture_error:
				raw_path.unlink(missing_ok=True)
				visual_path.unlink(missing_ok=True)
				self.logger.warning(
					'Observation screenshot retry failed for step %s; continuing this step without a screenshot: %s',
					step_name,
					retry_capture_error,
				)

		page_text = (
			self._dialog_observation_text() + self._navigation_observation_text(page) + await self._collect_page_text(page)
		)
		elements = await self._collect_elements(page)
		viewport = await self._viewport(page)
		if raw_screenshot is None:
			screenshot = b''
		else:
			try:
				visual_screenshot = await self._capture_screenshot(page, visual_path)
				screenshot = visual_screenshot
				if not self._legacy_marker_bindings_active:
					screenshot = self._annotate_element_screenshot(visual_screenshot, elements, viewport)
					if screenshot != visual_screenshot:
						visual_path.write_bytes(screenshot)
			except Exception:
				self.logger.exception('Could not capture annotated screenshot for step %s', step_name)
				screenshot = raw_screenshot

		tabs = await self._tabs()
		title = await self._page_title(page)
		return BrowserObservation(
			screenshot=screenshot,
			url=page.url,
			title=title,
			tabs=tabs,
			viewport_width=viewport['width'],
			viewport_height=viewport['height'],
			elements=elements,
			page_text=page_text,
			recent_network=copy.deepcopy(self.all_requests[-20:]),
			downloads=copy.deepcopy(self.downloads),
			step=int(step),
			screenshot_path=str(raw_path),
			visual_screenshot_path=str(visual_path),
		)

	async def _capture_screenshot(self, page: Page, path: Path) -> bytes:
		try:
			return await page.screenshot(
				path=str(path),
				type='png',
				animations='disabled',
				timeout=self.screenshot_timeout_ms,
			)
		except PlaywrightTimeoutError as screenshot_timeout:
			return await self._recover_screenshot_through_cdp(
				page,
				path,
				failure=screenshot_timeout,
				failure_description=f'timed out after {self.screenshot_timeout_ms}ms',
			)
		except PlaywrightError as screenshot_error:
			if not self._is_capture_screenshot_protocol_error(screenshot_error):
				raise
			return await self._recover_screenshot_through_cdp(
				page,
				path,
				failure=screenshot_error,
				failure_description='failed with Page.captureScreenshot protocol error',
			)

	async def _capture_observation_screenshot(self, page: Page, path: Path) -> tuple[Page, bytes]:
		"""Capture an observation screenshot, switching only to a live owned recovery page on failure."""

		try:
			return page, await self._capture_screenshot(page, path)
		except asyncio.CancelledError:
			raise
		except Exception as capture_error:
			for candidate in await self._screenshot_recovery_pages(page):
				path.unlink(missing_ok=True)
				try:
					screenshot = await self._capture_screenshot(candidate, path)
				except asyncio.CancelledError:
					raise
				except Exception as recovery_error:
					self.logger.warning(
						'Screenshot recovery candidate failed (%s): %s',
						redact_cdp_url(candidate.url),
						recovery_error,
					)
					continue
				self.logger.warning(
					'Screenshot failed for %s; switching observation to live page %s',
					redact_cdp_url(page.url),
					redact_cdp_url(candidate.url),
				)
				self.page = candidate
				self._element_bindings.clear()
				self._legacy_marker_bindings_active = False
				with contextlib.suppress(Exception):
					await candidate.bring_to_front()
				return candidate, screenshot
			raise capture_error

	async def _screenshot_recovery_pages(self, failed_page: Page) -> list[Page]:
		"""Return live owned recovery pages, preferring the failed page's opener."""

		candidates: list[Page] = []

		def add_if_safe(candidate: Page | None) -> None:
			if (
				candidate is None
				or candidate is failed_page
				or candidate.is_closed()
				or candidate.url == _DOWNLOAD_PLACEHOLDER_URL
				or is_forbidden_search_url(candidate.url)
				or all(id(owned) != id(candidate) for owned in self._owned_pages)
			):
				return
			if all(id(existing) != id(candidate) for existing in candidates):
				candidates.append(candidate)

		with contextlib.suppress(Exception):
			add_if_safe(await failed_page.opener())
		for candidate in reversed(self._live_owned_pages()):
			add_if_safe(candidate)
		return candidates

	@staticmethod
	def _annotate_element_screenshot(
		screenshot: bytes,
		elements: Sequence[ElementRef],
		viewport: Mapping[str, int],
	) -> bytes:
		"""Draw element IDs on screenshot pixels instead of injecting page overlays."""

		width = int(viewport.get('width', 0))
		height = int(viewport.get('height', 0))
		if not screenshot or width <= 0 or height <= 0 or not elements:
			return screenshot
		try:
			from PIL import Image, ImageDraw, ImageFont

			with Image.open(io.BytesIO(screenshot)).convert('RGBA') as image:
				draw = ImageDraw.Draw(image)
				font = ImageFont.load_default()
				scale_x = image.width / width
				scale_y = image.height / height
				occupied_labels: list[tuple[int, int, int, int]] = []
				for element in elements:
					left = max(0.0, element.x)
					top = max(0.0, element.y)
					right = min(float(width), element.x + element.width)
					bottom = min(float(height), element.y + element.height)
					if right <= left or bottom <= top:
						continue
					box = (left * scale_x, top * scale_y, right * scale_x, bottom * scale_y)
					draw.rectangle(box, outline=(255, 45, 85, 255), width=max(1, round(2 * min(scale_x, scale_y))))
					label = str(element.index)
					label_box = draw.textbbox((0, 0), label, font=font)
					label_width = label_box[2] - label_box[0] + 4
					label_height = label_box[3] - label_box[1] + 2
					label_left, label_top = _place_annotation_label(
						anchor_left=box[0],
						anchor_top=box[1] - label_height,
						label_width=label_width,
						label_height=label_height,
						image_width=image.width,
						image_height=image.height,
						occupied=occupied_labels,
					)
					label_rect = (label_left, label_top, label_left + label_width, label_top + label_height)
					occupied_labels.append(label_rect)
					draw.rectangle(label_rect, fill=(255, 45, 85, 255))
					draw.text((label_left + 2, label_top + 1), label, fill=(255, 255, 255, 255), font=font)
				buffer = io.BytesIO()
				image.save(buffer, format='PNG')
				return buffer.getvalue()
		except Exception:
			# Screenshot annotation is a debugging aid.  A stripped-down evaluator
			# image or absent Pillow must not prevent the textual observation.
			return screenshot

	@staticmethod
	def _is_capture_screenshot_protocol_error(error: PlaywrightError) -> bool:
		"""Match only the transient capture failure seen from Playwright's screenshot wrapper."""

		return 'Protocol error (Page.captureScreenshot): Unable to capture screenshot' in str(error)

	async def _recover_screenshot_through_cdp(
		self,
		page: Page,
		path: Path,
		*,
		failure: Exception,
		failure_description: str,
	) -> bytes:
		self.logger.warning('Playwright screenshot %s for %s; falling back to CDP capture', failure_description, path.name)
		try:
			screenshot = await self._capture_cdp_screenshot_before_deadline(page, path)
		except Exception as fallback_error:
			path.unlink(missing_ok=True)
			raise _ScreenshotFallbackError(
				f'Playwright screenshot {failure_description} ({failure}) and CDP fallback failed: '
				f'{type(fallback_error).__name__}: {fallback_error}'
			) from fallback_error
		self.logger.info('Recovered screenshot %s through CDP capture', path.name)
		return screenshot

	async def _capture_cdp_screenshot_before_deadline(self, page: Page, path: Path) -> bytes:
		expired = asyncio.Event()
		task = asyncio.create_task(self._capture_cdp_screenshot(page, path, expired))
		timeout_seconds = self.cdp_screenshot_timeout_ms / 1000
		try:
			done, _ = await asyncio.wait({task}, timeout=timeout_seconds)
		except BaseException:
			expired.set()
			if not task.done():
				self._track_screenshot_recovery_task(task)
				task.cancel()
			raise
		if task in done or task.done():
			return task.result()
		expired.set()
		self._track_screenshot_recovery_task(task)
		task.cancel()
		raise TimeoutError(f'CDP screenshot recovery exceeded the {timeout_seconds:g}-second hard deadline')

	def _track_screenshot_recovery_task(self, task: asyncio.Task[Any]) -> None:
		self._screenshot_recovery_tasks.add(task)

		def consume_and_discard(finished: asyncio.Task[Any]) -> None:
			self._screenshot_recovery_tasks.discard(finished)
			_consume_detached_task_result(finished)

		task.add_done_callback(consume_and_discard)

	async def _capture_cdp_screenshot(self, page: Page, path: Path, expired: asyncio.Event) -> bytes:
		cdp_session = None
		try:
			cdp_session = await self.context.new_cdp_session(page)
			result = await cdp_session.send(
				'Page.captureScreenshot',
				{'format': 'png', 'captureBeyondViewport': False},
			)
			encoded = result.get('data')
			if not isinstance(encoded, str) or not encoded:
				raise ValueError('CDP screenshot response did not include image data')
			try:
				screenshot = base64.b64decode(encoded, validate=True)
			except (ValueError, TypeError) as exc:
				raise ValueError('CDP screenshot response was not valid base64') from exc
			if not screenshot.startswith(_PNG_SIGNATURE):
				raise ValueError('CDP screenshot response was not a PNG image')
			if expired.is_set():
				raise TimeoutError('CDP screenshot recovery expired before the image could be persisted')
			path.write_bytes(screenshot)
			return screenshot
		finally:
			if cdp_session is not None:
				await cdp_session.detach()

	async def execute(self, decision: AgentDecision | Mapping[str, Any]) -> WebRetrieverActionResult:
		"""Execute one model decision and return its bounded structured result."""

		self._ensure_started()
		async with self._execute_lock:
			data = self._decision_dict(decision)
			action, params = self._normalise_action(data)
			page = self._active_page()
			previous_page = page
			previous_url = page.url
			known_page_ids = {id(item) for item in self._live_owned_pages()}
			dialog_cursor = self._next_dialog_id
			await self._clear_markers(remove_attributes=False)
			before = await self._capture_action_state(page, action, params)

			try:
				raw_result = await self._perform_action(action, params)
				if action not in {'wait'} and self.page is not None and not self.page.is_closed():
					settle_page = self.page
					try:
						await settle_page.wait_for_timeout(200)
					except PlaywrightError:
						if settle_page.url != _DOWNLOAD_PLACEHOLDER_URL:
							raise
						self.logger.info('Download placeholder page closed while settling action; restoring a live page')
					if action in {'click', 'double_click', 'press', 'xy'} and not settle_page.is_closed():
						# Link/form navigation is normally triggered by an interaction rather
						# than a navigate action. If it is still loading after the short
						# settle, give it the same 10s/120s observation contract.
						with contextlib.suppress(PlaywrightError, AttributeError):
							await self._observe_interaction_navigation(settle_page)
				await self._drain_dialog_tasks()
				await self._enforce_search_policy(previous_page, previous_url, known_page_ids)
				await self._drain_dialog_tasks()
			except Exception as exc:
				return self._action_error_result(action, exc, before)

			try:
				after_page = self._active_page()
				after = await self._capture_action_state(after_page, action, params)
			except Exception as exc:
				after = {'probe_error': f'{type(exc).__name__}: {exc}'[:500]}
			result_text = self._append_dialog_outcome(raw_result, since_id=dialog_cursor)
			return self._build_action_result(action, result_text, before=before, after=after)

	def capture_payload(self) -> dict[str, Any]:
		"""Return the official ``capture.json`` envelope.

		The request keys from the reference implementation are retained verbatim;
		bounded response fields are additive and ignored by older evaluators.  Some
		legitimate sites render entirely from their initial HTML document and never
		issue an XHR or fetch.  In that case, retain the real top-level document
		request as the capture fallback rather than emitting an empty artifact.
		"""

		requests = copy.deepcopy(self.all_requests)
		if not requests:
			# ``all_requests`` historically held XHR/fetch calls only.  The smoke
			# evaluator, however, needs evidence that the supplied browser actually
			# reached the task website.  A document navigation is that evidence for
			# static pages, and is recorded by the same Playwright request listener.
			for network_entry in self.network_requests:
				if network_entry.get('resource_type') != 'document':
					continue
				entry = {
					key: copy.deepcopy(network_entry[key])
					for key in (
						'timestamp',
						'url',
						'method',
						'headers',
						'resource_type',
						'post_data',
						'post_text',
						'json_data',
						'status',
						'response_headers',
						'failure',
					)
					if key in network_entry
				}
				requests.append(entry)
				break

		return {
			'capture_time': datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S'),
			'total_requests': len(requests),
			'all_requests': requests,
		}

	async def save_capture(self, path: str | Path | None = None) -> Path:
		await self._drain_background_tasks()
		await self._drain_download_tasks()
		destination = Path(path) if path is not None else self.task_dir / 'capture.json'
		destination.parent.mkdir(parents=True, exist_ok=True)
		destination.write_text(json.dumps(self.capture_payload(), ensure_ascii=False, indent=2), encoding='utf-8')
		return destination

	def cleanup_diagnostics(self) -> dict[str, Any]:
		"""Return the latest bounded-cleanup report for the task artifact."""

		return copy.deepcopy(self._cleanup_diagnostics)

	@staticmethod
	def _cleanup_remaining_seconds(deadline: float) -> float:
		return max(0.0, deadline - time.monotonic())

	@staticmethod
	def _detach_cleanup_task(task: asyncio.Task[Any]) -> None:
		task.add_done_callback(_consume_detached_task_result)
		task.cancel()

	async def _await_cleanup_operation(self, awaitable: Coroutine[Any, Any, Any], deadline: float) -> bool:
		"""Await one cleanup operation only until the shared cleanup deadline."""

		task = asyncio.create_task(awaitable)
		try:
			remaining = self._cleanup_remaining_seconds(deadline)
			if remaining <= 0:
				self._detach_cleanup_task(task)
				return False
			done, _ = await asyncio.wait({task}, timeout=remaining)
		except BaseException:
			if not task.done():
				self._detach_cleanup_task(task)
			raise
		if task in done or task.done():
			with contextlib.suppress(asyncio.CancelledError, Exception):
				task.result()
			return True
		self._detach_cleanup_task(task)
		return False

	async def _cancel_and_drain_task_groups(self, deadline: float) -> tuple[int, dict[str, int]]:
		"""Cancel non-essential runtime work and wait only for the shared deadline."""

		groups = {
			'background': self._background_tasks,
			'download': self._download_tasks,
			'policy': self._policy_tasks,
			'dialog': self._dialog_tasks,
		}
		task_groups: dict[asyncio.Task[Any], str] = {}
		cancelled = 0
		for name, group in groups.items():
			tasks = tuple(group)
			group.clear()
			for task in tasks:
				task_groups[task] = name
				if not task.done():
					cancelled += 1
					task.cancel()
		if not task_groups:
			return cancelled, {}

		try:
			remaining = self._cleanup_remaining_seconds(deadline)
			if remaining > 0:
				done, pending = await asyncio.wait(tuple(task_groups), timeout=remaining)
			else:
				done, pending = set(), set(task_groups)
		except BaseException:
			for task in task_groups:
				if not task.done():
					self._detach_cleanup_task(task)
			raise

		for task in done:
			_consume_detached_task_result(task)
		residual: dict[str, int] = {}
		for task in pending:
			name = task_groups[task]
			residual[name] = residual.get(name, 0) + 1
			self._detach_cleanup_task(task)
		return cancelled, residual

	async def _close_owned_pages_before_deadline(self, deadline: float) -> int:
		tasks: list[asyncio.Task[Any]] = []
		for page in reversed(self._owned_pages):
			with contextlib.suppress(Exception):
				if not page.is_closed():
					tasks.append(asyncio.create_task(page.close(run_before_unload=False)))
		if not tasks:
			return 0
		try:
			remaining = self._cleanup_remaining_seconds(deadline)
			if remaining > 0:
				done, pending = await asyncio.wait(tasks, timeout=remaining)
			else:
				done, pending = set(), set(tasks)
		except BaseException:
			for task in tasks:
				if not task.done():
					self._detach_cleanup_task(task)
			raise
		for task in done:
			_consume_detached_task_result(task)
		for task in pending:
			self._detach_cleanup_task(task)
		return len(pending)

	async def close(self, *, timeout_seconds: float = _DEFAULT_RUNTIME_CLEANUP_TIMEOUT_SECONDS) -> dict[str, Any]:
		"""Boundedly cancel runtime work and close only runtime-owned pages.

		A task result must outlive an unresponsive CDP operation.  Closing therefore
		cancels collector work first, then waits only within ``timeout_seconds``.
		"""

		if timeout_seconds <= 0:
			raise ValueError('timeout_seconds must be greater than 0')
		if self._closed:
			return self.cleanup_diagnostics()
		self._closed = True
		started_at = time.monotonic()
		deadline = started_at + timeout_seconds
		report: dict[str, Any] = {
			'status': 'in_progress',
			'grace_seconds': timeout_seconds,
			'elapsed_seconds': 0.0,
			'cancelled_tasks': 0,
			'residual_tasks': {},
		}
		self._cleanup_diagnostics = report
		for event, handler in self._context_handlers:
			with contextlib.suppress(Exception):
				self.context.remove_listener(event, handler)
		self._context_handlers.clear()
		if self._before_close is not None:
			anchor_restored = await self._await_cleanup_operation(self._before_close(), deadline)
			if not anchor_restored:
				report['residual_tasks']['anchor_restore'] = 1
		residual_navigation = await self._cancel_pending_navigations(timeout_seconds=self._cleanup_remaining_seconds(deadline))
		if residual_navigation:
			report['residual_tasks']['navigation'] = residual_navigation

		for page in list(self._owned_pages):
			for event, handler in self._page_handlers.pop(id(page), []):
				with contextlib.suppress(Exception):
					page.remove_listener(event, handler)
		residual_screenshots = await self._cancel_screenshot_recovery_tasks(
			timeout_seconds=self._cleanup_remaining_seconds(deadline)
		)
		if residual_screenshots:
			report['residual_tasks']['screenshot_recovery'] = residual_screenshots
		cancelled_tasks, residual_tasks = await self._cancel_and_drain_task_groups(deadline)
		report['cancelled_tasks'] = cancelled_tasks
		report['residual_tasks'].update(residual_tasks)
		markers_cleared = await self._await_cleanup_operation(self._clear_markers(remove_attributes=True), deadline)
		if not markers_cleared:
			report['residual_tasks']['marker_cleanup'] = 1
		residual_page_closes = await self._close_owned_pages_before_deadline(deadline)
		if residual_page_closes:
			report['residual_tasks']['page_close'] = residual_page_closes
		self._owned_pages.clear()
		self._element_bindings.clear()
		self._request_entries.clear()
		self._request_entries_by_id.clear()
		self._request_ids_by_entry.clear()
		self._network_request_entries.clear()
		self._network_entries_by_id.clear()
		self._network_responses.clear()
		self._network_page_cursors.clear()
		self._page_ids.clear()
		self._page_document_generations.clear()
		self._reserved_download_paths.clear()
		self.page = None
		report['elapsed_seconds'] = round(time.monotonic() - started_at, 3)
		report['status'] = 'timed_out' if report['residual_tasks'] else 'completed'
		self._cleanup_diagnostics = report
		return self.cleanup_diagnostics()

	async def _cancel_screenshot_recovery_tasks(self, *, timeout_seconds: float | None = None) -> int:
		tasks = tuple(self._screenshot_recovery_tasks)
		if not tasks:
			return 0
		for task in tasks:
			task.cancel()
		cleanup_timeout = min(1.0, max(0.0, self.cdp_screenshot_timeout_ms / 1000))
		if timeout_seconds is not None:
			cleanup_timeout = min(cleanup_timeout, max(0.0, timeout_seconds))
		_, pending = await asyncio.wait(tasks, timeout=cleanup_timeout)
		if pending:
			self.logger.warning('%d timed-out CDP screenshot recovery task(s) resisted bounded cleanup', len(pending))
			for task in pending:
				self._detach_cleanup_task(task)
		return len(pending)

	def _install_context_listeners(self) -> None:
		handlers: list[tuple[str, Callable[..., Any]]] = [
			('page', self._on_context_page),
			('request', self._on_request),
			('response', self._on_response),
			('requestfailed', self._on_request_failed),
		]
		for event, handler in handlers:
			cast(Any, self.context).on(event, handler)
		self._context_handlers.extend(handlers)

	def _register_page(self, page: Page, *, make_active: bool) -> None:
		if all(id(existing) != id(page) for existing in self._owned_pages):
			self._owned_pages.append(page)
			self._page_ids[id(page)] = len(self._page_ids)
			self._page_document_generations[id(page)] = 0
		page.set_default_timeout(self.action_timeout_ms)
		page.set_default_navigation_timeout(self.navigation_timeout_ms)
		if id(page) not in self._page_handlers:
			frame_handler = lambda frame, owner=page: self._on_frame_navigated(owner, frame)
			download_handler = self._on_download
			close_handler = lambda owner=page: self._on_page_closed(owner)
			dialog_handler = lambda dialog, owner=page: self._on_dialog(owner, dialog)
			page.on('framenavigated', frame_handler)
			page.on('download', download_handler)
			page.on('close', close_handler)
			page.on('dialog', dialog_handler)
			self._page_handlers[id(page)] = [
				('framenavigated', frame_handler),
				('download', download_handler),
				('close', close_handler),
				('dialog', dialog_handler),
			]
		if make_active:
			self.page = page
		self._record_url(page.url, unless_last=True)

	async def _configure_owned_page(self, page: Page) -> None:
		"""Configure the User-Agent consistently before an owned page navigates."""

		if self.declared_user_agent and is_sec_url(self.website):
			await page.set_extra_http_headers({'User-Agent': self.declared_user_agent})
			return

		user_agent = await page.evaluate('navigator.userAgent')
		if not isinstance(user_agent, str):
			return
		normalized_user_agent = user_agent.replace('HeadlessChrome/', 'Chrome/')
		if normalized_user_agent == user_agent:
			return

		# Keep the HTTP header and JavaScript-visible value in sync for the page's
		# future documents. This must run before its first navigation.
		await page.set_extra_http_headers({'User-Agent': normalized_user_agent})
		await page.add_init_script(
			f"""
			(() => {{
				const userAgent = {json.dumps(normalized_user_agent)};
				Object.defineProperty(Navigator.prototype, 'userAgent', {{
					configurable: true,
					get: () => userAgent,
				}});
			}})();
			"""
		)

	def _capture_request_headers(self, headers: Mapping[str, str]) -> dict[str, str]:
		"""Copy request headers while keeping configured contact details private."""

		captured = dict(headers)
		if not self.declared_user_agent:
			return captured
		for name, value in tuple(captured.items()):
			if name.lower() == 'user-agent' and value == self.declared_user_agent:
				captured[name] = _REDACTED
		return captured

	def _on_context_page(self, page: Page) -> None:
		if self._closed:
			return
		self._register_page(page, make_active=True)
		# A target=_blank link or script-created popup is registered by the
		# browser, not by one of our explicit new-page actions. Apply the same
		# page-scoped User-Agent configuration for its follow-up requests.
		self._spawn_background(self._configure_owned_page(page))

	def _on_page_closed(self, page: Page) -> None:
		self._element_bindings.clear()
		self._legacy_marker_bindings_active = False
		if self.page is page:
			live = [item for item in self._owned_pages if item is not page and not item.is_closed()]
			self.page = live[-1] if live else None

	def _on_dialog(self, page: Page, dialog: Dialog) -> None:
		"""Record and dismiss a native dialog without letting it block the browser.

		Playwright auto-dismisses dialogs only when no listener exists.  Once this
		listener is registered, dismissal becomes our responsibility; doing it in a
		task keeps the event callback synchronous while still allowing the action
		that triggered the dialog to settle normally.
		"""

		message = str(dialog.message).strip()
		record = {
			'id': self._next_dialog_id,
			'type': str(dialog.type),
			'message': message[:_MAX_DIALOG_MESSAGE],
			'url': str(page.url),
		}
		self._next_dialog_id += 1
		self._recent_dialogs.append(record)
		if len(self._recent_dialogs) > _MAX_RECENT_DIALOGS:
			del self._recent_dialogs[:-_MAX_RECENT_DIALOGS]
		self.logger.info('Browser dialog on %s: [%s] %s', redact_cdp_url(record['url']), record['type'], record['message'])
		self._spawn_dialog(self._dismiss_dialog(dialog))

	async def _dismiss_dialog(self, dialog: Dialog) -> None:
		"""Preserve the old no-listener behavior after retaining the diagnostic."""

		with contextlib.suppress(PlaywrightError):
			await dialog.dismiss()

	def _append_dialog_outcome(self, result: str, *, since_id: int) -> str:
		new_dialogs = [record for record in self._recent_dialogs if int(record['id']) >= since_id]
		if not new_dialogs:
			return result
		return result + '\nBrowser dialogs observed (untrusted):\n' + self._render_dialog_lines(new_dialogs)

	def _dialog_observation_text(self) -> str:
		if not self._recent_dialogs:
			return ''
		return 'Recent browser dialogs (untrusted):\n' + self._render_dialog_lines(self._recent_dialogs) + '\n\n'

	@staticmethod
	def _render_dialog_lines(records: Sequence[Mapping[str, Any]]) -> str:
		return '\n'.join(f'  [{str(record.get("type", "dialog"))}] {str(record.get("message", ""))}' for record in records)

	def _on_frame_navigated(self, page: Page, frame: Frame) -> None:
		# backendNodeId is scoped to one document and can be reused after any frame
		# navigation.  Conservatively invalidate the complete observation instead
		# of risking a stale index operating on a different control.
		self._element_bindings.clear()
		self._legacy_marker_bindings_active = False
		if frame is not page.main_frame:
			return
		self._page_document_generations[id(page)] = self._page_document_generations.get(id(page), 0) + 1
		url = frame.url
		if is_forbidden_search_url(url):
			if id(page) not in self._rollback_pages and not self._closed:
				self._rollback_pages.add(id(page))
				self._spawn_policy(self._rollback_forbidden_page(page))
		elif url and url not in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
			self._record_url(url)
			self._last_safe_urls[id(page)] = url

	def _network_request_owner(self, request: Request) -> tuple[int | None, int | None, str, str]:
		"""Resolve a request to its owning top-level page and document generation."""

		try:
			frame = request.frame
			page = frame.page
		except Exception:
			return None, None, '', ''
		page_key = id(page)
		document_generation = self._page_document_generations.get(page_key)
		# The main-document request fires before ``framenavigated`` advances the
		# generation. Attribute the entire redirect chain to the document it is
		# creating so the current-page snapshot includes its own HTML request.
		with contextlib.suppress(Exception):
			if request.is_navigation_request() and frame is page.main_frame and document_generation is not None:
				document_generation += 1
		return (
			self._page_ids.get(page_key),
			document_generation,
			str(page.url),
			str(frame.url),
		)

	def _on_request(self, request: Request) -> None:
		post_data: str | None = None
		raw_post_data: str | None = None
		json_data: Any = None
		with contextlib.suppress(Exception):
			post_data = request.post_data
		with contextlib.suppress(Exception):
			raw = request.post_data_buffer
			if raw:
				raw_post_data = raw.hex()
				if post_data is None:
					post_data = raw.decode('utf-8', errors='replace')
		if post_data:
			with contextlib.suppress(json.JSONDecodeError):
				json_data = json.loads(post_data)
		page_id, document_generation, page_url, frame_url = self._network_request_owner(request)
		request_id = self._next_network_request_id
		self._next_network_request_id += 1
		network_entry: dict[str, Any] = {
			'request_id': request_id,
			'timestamp': time.time(),
			'url': request.url,
			'method': request.method,
			'headers': self._capture_request_headers(request.headers),
			'resource_type': request.resource_type,
			'post_data': post_data,
			'post_text': post_data,
			'json_data': json_data,
			'page_id': page_id,
			'document_generation': document_generation,
			'page_url': page_url,
			'frame_url': frame_url,
			'response_body_state': 'pending',
		}
		if raw_post_data is not None:
			network_entry['raw_post_data'] = raw_post_data
		self.network_requests.append(network_entry)
		self._network_request_entries[id(request)] = network_entry
		self._network_entries_by_id[request_id] = network_entry

		if request.resource_type in {'xhr', 'fetch'}:
			# Preserve the historical capture schema exactly; page provenance and
			# request IDs belong only to the action-specific all-resource cache.
			entry = {
				'timestamp': network_entry['timestamp'],
				'url': network_entry['url'],
				'method': network_entry['method'],
				'headers': copy.deepcopy(network_entry['headers']),
				'resource_type': network_entry['resource_type'],
				'post_data': post_data,
				'post_text': post_data,
				'json_data': copy.deepcopy(json_data),
			}
			if raw_post_data is not None:
				entry['raw_post_data'] = raw_post_data
			self.all_requests.append(entry)
			self._request_entries[id(request)] = entry
			self._request_entries_by_id[request_id] = entry
			self._request_ids_by_entry[id(entry)] = request_id

	def _on_response(self, response: Response) -> None:
		network_entry = self._network_request_entries.get(id(response.request))
		if network_entry is not None:
			network_entry['status'] = response.status
			network_entry['response_headers'] = dict(response.headers)
			network_entry['response_body_state'] = 'available'
			self._network_responses[int(network_entry['request_id'])] = response
		entry = self._request_entries.get(id(response.request))
		if entry is not None:
			self._spawn_background(self._capture_response(response, entry))
		extension = self._document_response_extension(response)
		if extension is not None:
			key = (response.url, response.status)
			if key not in self._captured_document_responses:
				self._captured_document_responses.add(key)
				self._spawn_download(self._save_document_response(response, extension))

	def _on_request_failed(self, request: Request) -> None:
		network_entry = self._network_request_entries.get(id(request))
		if network_entry is not None:
			with contextlib.suppress(Exception):
				network_entry['failure'] = request.failure
			network_entry['response_body_state'] = 'failed'
		entry = self._request_entries.get(id(request))
		if entry is None:
			return
		with contextlib.suppress(Exception):
			entry['failure'] = request.failure

	async def _capture_response(self, response: Response, entry: dict[str, Any]) -> None:
		try:
			entry['status'] = response.status
			entry['response_headers'] = dict(response.headers)
			content_length = self._content_length(response.headers)
			if self.max_response_body_bytes <= 0:
				return
			if content_length is not None and content_length > self.max_response_body_bytes:
				entry['response_body_omitted'] = f'content-length {content_length} exceeds limit'
				entry['response_body_truncated'] = True
				return
			body = await asyncio.wait_for(response.body(), timeout=max(1, self.action_timeout_ms / 1000))
			truncated = len(body) > self.max_response_body_bytes
			bounded = body[: self.max_response_body_bytes]
			content_type = response.headers.get('content-type', '').lower()
			if self._is_textual_content(content_type, bounded):
				text = bounded.decode(self._charset(content_type), errors='replace')
				entry['response_body'] = text
				if 'json' in content_type:
					with contextlib.suppress(json.JSONDecodeError):
						entry['response_json'] = json.loads(text)
			else:
				entry['response_body_base64'] = base64.b64encode(bounded).decode('ascii')
				entry['response_body_encoding'] = 'base64'
			entry['response_body_bytes'] = len(body)
			entry['response_body_truncated'] = truncated
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			entry['response_error'] = f'{type(exc).__name__}: {exc}'[:500]

	async def settle_network_capture(self) -> None:
		"""Finish response collectors that were already scheduled by the browser."""

		await self._drain_background_tasks()

	def current_page_network_requests(self) -> list[dict[str, Any]]:
		"""Return all HTTP requests belonging to the active page's current document.

		Requests initiated by child frames retain the top-level page identifier, so
		Tableau and similar cross-origin embeds are included. Requests from previous
		documents and other tabs are excluded.
		"""

		page = self._active_page()
		page_key = id(page)
		page_id = self._page_ids.get(page_key)
		document_generation = self._page_document_generations.get(page_key)
		if page_id is None or document_generation is None:
			return []
		return [
			copy.deepcopy(entry)
			for entry in self.network_requests
			if entry.get('page_id') == page_id and entry.get('document_generation') == document_generation
		]

	async def materialize_network_request(
		self,
		request_id: int,
		*,
		max_body_bytes: int = _MAX_DOWNLOAD_BYTES,
	) -> dict[str, Any]:
		"""Return one all-resource request with its complete bounded response body."""

		entry = self._network_entries_by_id.get(int(request_id))
		if entry is None:
			raise ValueError(f'Unknown network request_id {request_id}')
		state = str(entry.get('response_body_state', 'pending'))
		if state in {'complete', 'body_too_large', 'failed', 'unavailable', 'error'}:
			return copy.deepcopy(entry)

		response = self._network_responses.get(int(request_id))
		if response is None:
			entry['response_body_state'] = 'unavailable' if entry.get('status') is not None else 'pending'
			return copy.deepcopy(entry)

		headers = entry.get('response_headers')
		if not isinstance(headers, Mapping):
			headers = dict(response.headers)
			entry['response_headers'] = dict(headers)
		content_length = self._content_length(headers)
		if content_length is not None and content_length > max_body_bytes:
			entry['response_body_state'] = 'body_too_large'
			entry['response_body_bytes'] = content_length
			entry['response_error'] = f'content-length {content_length} exceeds {max_body_bytes} byte limit'
			return copy.deepcopy(entry)

		try:
			body = await asyncio.wait_for(response.body(), timeout=max(1, self.action_timeout_ms / 1000))
			entry['response_body_bytes'] = len(body)
			entry['response_body_sha256'] = hashlib.sha256(body).hexdigest()
			if len(body) > max_body_bytes:
				entry['response_body_state'] = 'body_too_large'
				entry['response_error'] = f'response body {len(body)} exceeds {max_body_bytes} byte limit'
				return copy.deepcopy(entry)

			content_type = str(headers.get('content-type', '')).lower()
			if self._is_textual_content(content_type, body):
				text = body.decode(self._charset(content_type), errors='replace')
				entry['response_body'] = text
				if 'json' in content_type:
					with contextlib.suppress(json.JSONDecodeError):
						entry['response_json'] = json.loads(text)
			else:
				entry['response_body_base64'] = base64.b64encode(body).decode('ascii')
				entry['response_body_encoding'] = 'base64'
			entry['response_body_state'] = 'complete'
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			entry['response_body_state'] = 'error'
			entry['response_error'] = f'{type(exc).__name__}: {exc}'[:500]
		return copy.deepcopy(entry)

	def _on_download(self, download: Download) -> None:
		self._download_urls_seen.add(download.url)
		self._spawn_download(self._save_download(download))

	async def _run_download_until_deadline(self, item: dict[str, Any], awaitable: Coroutine[Any, Any, None]) -> None:
		"""Run one download pipeline without letting cancellation resistance block observe()."""

		task = asyncio.create_task(awaitable)
		try:
			done, _ = await asyncio.wait({task}, timeout=self.download_timeout_ms / 1000)
		except BaseException:
			if not task.done():
				task.add_done_callback(_consume_detached_task_result)
				task.cancel()
			raise
		if task in done or task.done():
			task.result()
			return
		task.add_done_callback(_consume_detached_task_result)
		task.cancel()
		item['status'] = 'timed_out'
		item['failure'] = f'download exceeded the {self.download_timeout_ms / 1000:g}-second hard deadline'
		item.setdefault('text', '')
		item.setdefault('text_truncated', True)

	async def _prepare_download_data_artifact(self, item: dict[str, Any], destination: Path) -> None:
		if destination.suffix.casefold() not in _STRUCTURED_DOWNLOAD_EXTENSIONS or self.task_identity is None:
			return
		from browser_use.webretriever.download_artifacts import prepare_download_artifact

		artifact = await asyncio.to_thread(
			prepare_download_artifact,
			task_dir=self.task_dir,
			task_identity=self.task_identity,
			source_path=destination,
			source_url=str(item.get('url', '')),
			content_type=str(item.get('mime_type', '')),
		)
		item['data_artifact'] = artifact

	async def _finish_saved_download(self, item: dict[str, Any], destination: Path) -> None:
		if item.get('status') != 'pending':
			return
		item['size_bytes'] = destination.stat().st_size
		if destination.suffix.lower() == '.zip':
			async with self._archive_extraction_lock:
				extraction = await asyncio.to_thread(self._extract_download_archive, destination)
			if item.get('status') != 'pending':
				return
			item['extraction_status'] = extraction.status
			item['extracted_count'] = len(extraction.members)
			item['skipped_count'] = extraction.skipped_count
			item['extraction_warnings'] = extraction.warnings
			item['text'] = self._archive_extraction_summary(extraction)
			item['text_truncated'] = False
			for member in extraction.members:
				text = ''
				truncated = False
				text_extraction_error: str | None = None
				if member.path.suffix.lower() != '.zip':
					try:
						text, truncated = await asyncio.to_thread(self._extract_download_text, member.path)
					except Exception as exc:
						text_extraction_error = f'{type(exc).__name__}: {exc}'[:500]
						text = f'[text extraction failed: {text_extraction_error}]'
				member_item = _ArchiveMemberDownload(
					timestamp=time.time(),
					url=str(item.get('url', '')),
					suggested_filename=member.archive_name,
					filename=member.relative_name,
					path=str(member.path),
					size_bytes=member.size_bytes,
					source='archive_member',
					source_archive=destination.name,
					text=text,
					text_truncated=truncated,
					text_extraction_error=text_extraction_error,
				)
				self.downloads.append(member_item.model_dump(exclude_none=True))
		else:
			text, truncated = await asyncio.to_thread(self._extract_download_text, destination)
			if item.get('status') != 'pending':
				return
			item['text'] = text
			item['text_truncated'] = truncated
		try:
			await self._prepare_download_data_artifact(item, destination)
		except Exception as exc:
			if item.get('status') != 'pending':
				return
			item['status'] = 'failed'
			item['failure'] = f'data artifact normalization failed: {type(exc).__name__}: {exc}'[:500]
			return
		if item.get('status') == 'pending':
			item['status'] = 'ready'

	async def _save_download(self, download: Download) -> None:
		suggested = self._safe_download_name(download.suggested_filename)
		destination = self._unique_download_path(suggested)
		item: dict[str, Any] = {
			'timestamp': time.time(),
			'url': download.url,
			'suggested_filename': download.suggested_filename,
			'filename': destination.name,
			'path': str(destination),
			'source': 'browser_download',
			'status': 'pending',
		}
		self.downloads.append(item)
		try:

			async def save() -> None:
				await download.save_as(str(destination))
				if item.get('status') != 'pending':
					return
				failure = await download.failure()
				if item.get('status') != 'pending':
					return
				if failure:
					item['status'] = 'failed'
					item['failure'] = failure
					return
				await self._finish_saved_download(item, destination)

			await self._run_download_until_deadline(item, save())
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			item['status'] = 'failed'
			item['failure'] = f'{type(exc).__name__}: {exc}'[:500]

	async def _save_document_response(self, response: Response, extension: str) -> None:
		"""Persist an inline top-level document, then normalize supported data formats."""

		# A headless Chromium build can convert even an ``inline`` PDF response
		# into a Playwright Download.  Give that higher-fidelity event one loop
		# turn to arrive and avoid writing the same document through both paths.
		await asyncio.sleep(0.1)
		if response.url in self._download_urls_seen:
			return
		content_type = response.headers.get('content-type', '').split(';', 1)[0].strip().lower()
		name = self._response_filename(response, extension)
		destination = self._unique_download_path(name)
		item: dict[str, Any] = {
			'timestamp': time.time(),
			'url': response.url,
			'suggested_filename': name,
			'filename': destination.name,
			'path': str(destination),
			'mime_type': content_type,
			'source': 'document_response',
			'status': 'pending',
		}
		self.downloads.append(item)
		try:

			async def save() -> None:
				content_length = self._content_length(response.headers)
				if content_length is not None and content_length > _MAX_DOWNLOAD_BYTES:
					item['status'] = 'failed'
					item['failure'] = f'document body exceeds {_MAX_DOWNLOAD_BYTES} byte extraction limit'
					item['text'] = ''
					item['text_truncated'] = True
					return
				body = await response.body()
				if item.get('status') != 'pending':
					return
				if len(body) > _MAX_DOWNLOAD_BYTES:
					item['status'] = 'failed'
					item['failure'] = f'document body exceeds {_MAX_DOWNLOAD_BYTES} byte extraction limit'
					item['size_bytes'] = len(body)
					item['text'] = ''
					item['text_truncated'] = True
					return
				destination.write_bytes(body)
				await self._finish_saved_download(item, destination)

			await self._run_download_until_deadline(item, save())
		except asyncio.CancelledError:
			raise
		except Exception as exc:
			item['status'] = 'failed'
			item['failure'] = f'{type(exc).__name__}: {exc}'[:500]

	async def _capture_action_state(self, page: Page, action: str, params: Mapping[str, Any]) -> dict[str, Any]:
		"""Capture a bounded, action-relevant state probe without a screenshot."""

		state: dict[str, Any] = {'url': str(page.url), 'action': action}
		live_pages = self._live_owned_pages()
		try:
			state['active_page_index'] = next(index for index, candidate in enumerate(live_pages) if candidate is page)
		except StopIteration:
			state['active_page_index'] = -1
		state['title'] = await self._page_title(page)
		try:
			page_text = await self._collect_page_text(page)
			state['page_fingerprint'] = hashlib.sha256(page_text.encode('utf-8', errors='replace')).hexdigest()[:16]
			state['page_text_length'] = len(page_text)
		except Exception as exc:
			state['probe_error'] = f'page_text: {type(exc).__name__}: {exc}'[:500]

		try:
			scroll = await page.evaluate(
				"""() => ({
					x: Number(window.scrollX || 0),
					y: Number(window.scrollY || 0),
					height: Number(document.documentElement?.scrollHeight || 0)
			})"""
			)
			if isinstance(scroll, Mapping):
				state['scroll_x'] = float(scroll.get('x', 0))
				state['scroll_y'] = float(scroll.get('y', 0))
				state['document_height'] = float(scroll.get('height', 0))
		except Exception as exc:
			state['probe_error'] = f'{state.get("probe_error", "")} scroll: {type(exc).__name__}: {exc}'[:500]

		try:
			state['network_request_count'] = len(self.current_page_network_requests())
		except Exception:
			state['network_request_count'] = len(self.network_requests)
		state['download_count'] = len(self.downloads)
		state['dialog_count'] = len(self._recent_dialogs)
		state['target'] = await self._capture_target_state(params)
		return state

	async def _capture_target_state(self, params: Mapping[str, Any]) -> dict[str, Any]:
		"""Read common form/ARIA properties from the action target when available."""

		if not self._has_explicit_selector(params) and not self._has_target(params):
			return {}
		try:
			if self._has_explicit_selector(params):
				locator = self._target_locator(params)
				value = await locator.evaluate(_TARGET_STATE_JS)
			else:
				index = self._target_index(params)
				binding = self._binding_for_index(index)
				if self._is_backend_binding(binding):
					return await self._capture_backend_target_state(binding, index)
				locator = self._locator_for_index(index)
				value = await locator.evaluate(_TARGET_STATE_JS)
			return dict(value) if isinstance(value, Mapping) else {'value': str(value)}
		except Exception as exc:
			return {'probe_error': f'{type(exc).__name__}: {exc}'[:500]}

	async def _capture_backend_target_state(self, binding: _ElementBinding, index: int) -> dict[str, Any]:
		"""Read target state from a CDP backend node without requiring a selector."""

		async with self._backend_session(binding, index) as session:
			resolved = await session.send('DOM.resolveNode', {'backendNodeId': binding.backend_node_id})
			remote_object = resolved.get('object') if isinstance(resolved, Mapping) else None
			object_id = remote_object.get('objectId') if isinstance(remote_object, Mapping) else None
			if not isinstance(object_id, str) or not object_id:
				raise ValueError(f'Could not resolve backend-bound element {index} for state probing')
			try:
				result = await session.send(
					'Runtime.callFunctionOn',
					{
						'objectId': object_id,
						'functionDeclaration': _TARGET_STATE_JS,
						'returnByValue': True,
						'awaitPromise': False,
					},
				)
				remote_result = result.get('result') if isinstance(result, Mapping) else None
				value = remote_result.get('value') if isinstance(remote_result, Mapping) else None
				if not isinstance(value, Mapping):
					raise ValueError(f'Could not read backend-bound element {index} state')
				return dict(value)
			finally:
				with contextlib.suppress(Exception):
					await session.send('Runtime.releaseObject', {'objectId': object_id})

	async def _capture_backend_select_options(
		self,
		binding: _ElementBinding,
		index: int,
	) -> tuple[tuple[str, str], ...]:
		"""Read the current native-select options without changing browser state."""

		async with self._backend_session(binding, index) as session:
			resolved = await session.send('DOM.resolveNode', {'backendNodeId': binding.backend_node_id})
			remote_object = resolved.get('object') if isinstance(resolved, Mapping) else None
			object_id = remote_object.get('objectId') if isinstance(remote_object, Mapping) else None
			if not isinstance(object_id, str) or not object_id:
				raise ValueError(f'Could not resolve backend-bound select {index}')
			try:
				result = await session.send(
					'Runtime.callFunctionOn',
					{
						'objectId': object_id,
						'functionDeclaration': _SELECT_OPTIONS_JS,
						'returnByValue': True,
						'awaitPromise': False,
					},
				)
				remote_result = result.get('result') if isinstance(result, Mapping) else None
				options = remote_result.get('value') if isinstance(remote_result, Mapping) else None
				if not isinstance(options, list):
					raise ValueError(f'Element {index} is not a native select control')
				return tuple(
					(str(option.get('value', '')), str(option.get('label', '')))
					for option in options
					if isinstance(option, Mapping)
				)
			finally:
				with contextlib.suppress(Exception):
					await session.send('Runtime.releaseObject', {'objectId': object_id})

	@staticmethod
	def _action_state_diff(before: Mapping[str, Any], after: Mapping[str, Any]) -> list[str]:
		"""Return bounded human-readable differences between two probes."""

		if before.get('probe_error') or after.get('probe_error'):
			return []
		diff: list[str] = []
		for key in (
			'url',
			'active_page_index',
			'title',
			'page_fingerprint',
			'page_text_length',
			'scroll_x',
			'scroll_y',
			'document_height',
			'network_request_count',
			'download_count',
			'dialog_count',
		):
			if before.get(key) != after.get(key):
				diff.append(f'{key}: {before.get(key)!r} -> {after.get(key)!r}')

		before_target = before.get('target') if isinstance(before.get('target'), Mapping) else {}
		after_target = after.get('target') if isinstance(after.get('target'), Mapping) else {}
		for key in ('value', 'checked', 'selected', 'expanded', 'disabled', 'text', 'visible'):
			if before_target.get(key) != after_target.get(key):
				diff.append(f'target.{key}: {before_target.get(key)!r} -> {after_target.get(key)!r}')
		return diff[:32]

	@staticmethod
	def _bounded_action_output(value: str) -> str:
		if len(value) <= _MAX_ACTION_RESULT_OUTPUT:
			return value
		marker = '\n...[action result truncated]...\n'
		available = _MAX_ACTION_RESULT_OUTPUT - len(marker)
		if available <= 0:
			return marker[:_MAX_ACTION_RESULT_OUTPUT]
		head = min(14_000, available)
		tail = available - head
		return value[:head] + marker + (value[-tail:] if tail else '')

	@staticmethod
	def _error_metadata(exc: Exception) -> tuple[str, str]:
		message = str(exc)
		folded = message.casefold()
		if 'outside the current viewport' in folded:
			return 'ElementOutsideViewport', 'observe'
		if 'unknown element index' in folded or 'stale' in folded or 'page changed' in folded:
			return 'StaleElement', 'observe'
		if isinstance(exc, (PlaywrightTimeoutError, TimeoutError)) or 'timeout' in folded:
			return 'ActionTimeout', 'observe'
		return type(exc).__name__, 'replan'

	def _action_error_result(
		self,
		action: str,
		exc: Exception,
		before: Mapping[str, Any],
	) -> WebRetrieverActionResult:
		error_type, recovery = self._error_metadata(exc)
		return WebRetrieverActionResult(
			action=action,
			status='error',
			executed=False,
			state_changed=False,
			error_type=error_type,
			error=str(exc)[:4_000],
			recovery=recovery,
			before=dict(before),
			summary=f'{action} was not completed: {error_type}',
		)

	def _build_action_result(
		self,
		action: str,
		raw_result: str,
		*,
		before: Mapping[str, Any],
		after: Mapping[str, Any],
	) -> WebRetrieverActionResult:
		diff = self._action_state_diff(before, after)
		probe_uncertain = bool(before.get('probe_error') or after.get('probe_error'))
		state_changed: bool | None = None if probe_uncertain else bool(diff)
		if action in _STATE_CHANGING_ACTIONS and probe_uncertain:
			status: Literal['ok', 'error', 'no_change', 'uncertain'] = 'uncertain'
			recovery: Literal['none', 'observe', 're_ground', 'replan'] = 'observe'
		elif action in _STATE_CHANGING_ACTIONS and state_changed is False:
			status = 'no_change'
			recovery = 're_ground'
		else:
			status = 'ok'
			recovery = 'none'

		bounded_output = self._bounded_action_output(raw_result)
		is_content_action = action in _CONTENT_ACTIONS
		details: dict[str, Any] = {
			'mutation_expected': action in _STATE_CHANGING_ACTIONS,
			'evidence_changed': bool(raw_result.strip()) if is_content_action else False,
		}
		if probe_uncertain:
			details['probe_error'] = str(before.get('probe_error') or after.get('probe_error'))[:500]
		return WebRetrieverActionResult(
			action=action,
			status=status,
			executed=True,
			state_changed=state_changed,
			summary='Action completed; extracted content is attached.' if is_content_action else bounded_output,
			extracted_content=bounded_output if is_content_action else None,
			recovery=recovery,
			before=dict(before),
			after=dict(after),
			diff=diff,
			details=details,
		)

	async def _perform_action(self, action: str, params: dict[str, Any]) -> str:
		page = self._active_page()
		if action == 'click':
			if self._has_explicit_selector(params):
				await self._target_locator(params).click(timeout=self.action_timeout_ms)
				return 'clicked selected element'
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				await self._backend_click(binding, index)
			else:
				await self._locator_for_index(index).click(timeout=self.action_timeout_ms)
			return f'clicked element {index}'
		if action == 'double_click':
			if self._has_explicit_selector(params):
				await self._target_locator(params).dblclick(timeout=self.action_timeout_ms)
				return 'double-clicked selected element'
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				await self._backend_click(binding, index, click_count=2)
			else:
				await self._locator_for_index(index).dblclick(timeout=self.action_timeout_ms)
			return f'double-clicked element {index}'
		if action == 'type':
			text = str(self._first(params, 'text', 'value', 'input_text', default=''))
			submit = self._as_bool(self._first(params, 'submit', default=False))
			if self._has_explicit_selector(params):
				locator = self._target_locator(params)
				try:
					await locator.fill(text, timeout=self.action_timeout_ms)
				except Exception:
					await locator.click(timeout=self.action_timeout_ms)
					await locator.press('ControlOrMeta+A', timeout=self.action_timeout_ms)
					await locator.type(text, timeout=self.action_timeout_ms)
				if submit:
					await locator.press('Enter', timeout=self.action_timeout_ms)
				return 'typed into selected element'
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				await self._backend_type(binding, index, text, submit=submit)
			else:
				locator = self._locator_for_index(index)
				try:
					await locator.fill(text, timeout=self.action_timeout_ms)
				except Exception:
					await locator.click(timeout=self.action_timeout_ms)
					await locator.press('ControlOrMeta+A', timeout=self.action_timeout_ms)
					await locator.type(text, timeout=self.action_timeout_ms)
				if submit:
					await locator.press('Enter', timeout=self.action_timeout_ms)
			return f'typed into element {index}'
		if action == 'select':
			value = self._first(params, 'value', 'text', 'option')
			if value is None:
				raise ValueError('select requires value/text/option')
			if self._has_explicit_selector(params):
				locator = self._target_locator(params)
				try:
					selected = await locator.select_option(value=str(value), timeout=self.action_timeout_ms)
				except Exception:
					selected = await locator.select_option(label=str(value), timeout=self.action_timeout_ms)
				return f'selected {selected!r} on selected element'
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				selected = await self._backend_select(binding, index, str(value))
			else:
				locator = self._locator_for_index(index)
				try:
					selected = await locator.select_option(value=str(value), timeout=self.action_timeout_ms)
				except Exception:
					selected = await locator.select_option(label=str(value), timeout=self.action_timeout_ms)
			return f'selected {selected!r} on element {index}'
		if action == 'press':
			key = str(self._first(params, 'key', 'text', 'value', default='Enter'))
			if self._has_target(params):
				if self._has_explicit_selector(params):
					await self._target_locator(params).press(key, timeout=self.action_timeout_ms)
				else:
					index = self._target_index(params)
					binding = self._binding_for_index(index)
					if self._is_backend_binding(binding):
						await self._backend_focus(binding, index)
						await page.keyboard.press(key)
					else:
						await self._locator_for_index(index).press(key, timeout=self.action_timeout_ms)
			else:
				await page.keyboard.press(key)
			return f'pressed {key}'
		if action == 'scroll':
			return await self._scroll(params)
		if action == 'hover':
			if self._has_explicit_selector(params):
				await self._target_locator(params).hover(timeout=self.action_timeout_ms)
				return 'hovered selected element'
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				await self._backend_hover(binding, index)
			else:
				await self._locator_for_index(index).hover(timeout=self.action_timeout_ms)
			return f'hovered element {index}'
		if action == 'xy':
			x = float(self._required(params, 'x'))
			y = float(self._required(params, 'y'))
			button = cast(Literal['left', 'middle', 'right'], str(self._first(params, 'button', default='left')))
			if button not in {'left', 'middle', 'right'}:
				raise ValueError(f'Unsupported mouse button: {button!r}')
			await page.mouse.click(x, y, button=button)
			return f'clicked coordinates ({x:g}, {y:g})'
		if action == 'hover_xy':
			x = float(self._required(params, 'x'))
			y = float(self._required(params, 'y'))
			await page.mouse.move(x, y)
			return f'hovered coordinates ({x:g}, {y:g})'
		if action == 'drag':
			return await self._drag(params)
		if action == 'back':
			status = await self._start_navigation(
				page,
				page.url,
				page.go_back(wait_until='domcontentloaded', timeout=self.navigation_timeout_ms),
			)
			if status == 'pending':
				return 'first observation is available while back navigation continues'
			if status == 'download':
				return 'downloaded document while navigating back'
			return f'navigated back to {page.url}'
		if action == 'navigate':
			url = str(self._required(params, 'url'))
			self._validate_navigation_url(url)
			if is_forbidden_search_url(url):
				raise ValueError('External search engines are prohibited by the challenge rules')
			new_tab = self._as_bool(self._first(params, 'new_tab', default=False))
			if new_tab:
				page = await self.context.new_page()
				self._register_page(page, make_active=True)
				await self._configure_owned_page(page)
			download_started = await self._goto_exact(page, url)
			self._record_url(url if download_started else page.url, unless_last=True)
			if id(page) in self._pending_navigations:
				return f'first observation is available while navigation continues to {url}'
			return f'downloaded document from {url}' if download_started else f'navigated to {page.url}'
		if action == 'wait':
			milliseconds = self._wait_milliseconds(params)
			await page.wait_for_timeout(milliseconds)
			return f'waited {milliseconds}ms'
		if action == 'tab':
			return await self._tab_action(params)
		if action == 'read':
			return await self._read(params)
		if action == 'find':
			return await self._find(params)
		if action == 'inspect_network':
			return await self._inspect_network(params)
		if action == 'calculate':
			return self._calculate(params)
		if action in {'done', 'finish', 'answer', 'noop'}:
			return str(self._first(params, 'answer', 'text', 'message', default=action))
		raise ValueError(f'Unsupported browser action: {action!r}')

	async def _scroll(self, params: dict[str, Any]) -> str:
		direction = str(self._first(params, 'direction', default='down')).lower()
		amount_value = self._first(params, 'amount', 'pixels', 'delta')
		if amount_value is None and self._first(params, 'pages') is not None:
			viewport = await self._viewport(self._active_page())
			axis_size = viewport['height'] if direction in {'up', 'down'} else viewport['width']
			amount_value = max(axis_size * 0.5, 300) * float(self._first(params, 'pages'))
		amount = abs(float(amount_value if amount_value is not None else 600))
		delta_x = float(self._first(params, 'delta_x', 'dx', default=0))
		delta_y = float(self._first(params, 'delta_y', 'dy', default=0))
		if not delta_x and not delta_y:
			if direction in {'up', 'down'}:
				delta_y = -amount if direction == 'up' else amount
			else:
				delta_x = -amount if direction == 'left' else amount
		if self._has_target(params):
			if self._has_explicit_selector(params):
				result = await self._target_locator(params).evaluate(
					"""(element, delta) => {
					const permitsScroll = (node, axis) => {
						const style = getComputedStyle(node);
						const overflow = axis === 'x' ? style.overflowX : style.overflowY;
						const hasRange = axis === 'x'
							? node.scrollWidth > node.clientWidth + 1
							: node.scrollHeight > node.clientHeight + 1;
						return hasRange && /(auto|scroll|overlay)/.test(overflow);
					};
					let target = element;
					while (target && !permitsScroll(target, 'x') && !permitsScroll(target, 'y')) {
						target = target.parentElement;
					}
					target = target || document.scrollingElement || document.documentElement;
					const beforeX = target.scrollLeft;
					const beforeY = target.scrollTop;
					target.scrollBy(delta.x, delta.y);
					return {
						tag: target.tagName.toLowerCase(),
						id: target.id || '',
						className: typeof target.className === 'string' ? target.className : '',
						beforeX,
						beforeY,
						afterX: target.scrollLeft,
						afterY: target.scrollTop,
					};
				}""",
					{'x': delta_x, 'y': delta_y},
				)
			else:
				index = self._target_index(params)
				binding = self._binding_for_index(index)
				if self._is_backend_binding(binding):
					await self._backend_scroll(binding, index, delta_x, delta_y)
					return f'scrolled element {index} with Playwright pointer wheel; requested ({delta_x:g}, {delta_y:g})'
				else:
					result = await self._locator_for_index(index).evaluate(
						"""(element, delta) => {
							const permitsScroll = (node, axis) => {
								const style = getComputedStyle(node);
								const overflow = axis === 'x' ? style.overflowX : style.overflowY;
								const hasRange = axis === 'x'
									? node.scrollWidth > node.clientWidth + 1
									: node.scrollHeight > node.clientHeight + 1;
								return hasRange && /(auto|scroll|overlay)/.test(overflow);
							};
							let target = element;
							while (target && !permitsScroll(target, 'x') && !permitsScroll(target, 'y')) {
								target = target.parentElement;
							}
							target = target || document.scrollingElement || document.documentElement;
							const beforeX = target.scrollLeft;
							const beforeY = target.scrollTop;
							target.scrollBy(delta.x, delta.y);
							return {tag: target.tagName.toLowerCase(), id: target.id || '', className: typeof target.className === 'string' ? target.className : '', beforeX, beforeY, afterX: target.scrollLeft, afterY: target.scrollTop};
						}""",
						{'x': delta_x, 'y': delta_y},
					)
			target_name = str(result.get('tag', 'element'))
			if result.get('id'):
				target_name += f'#{result["id"]}'
			return (
				f'scrolled {target_name} from ({result.get("beforeX", 0):g}, {result.get("beforeY", 0):g}) '
				f'to ({result.get("afterX", 0):g}, {result.get("afterY", 0):g}); '
				f'requested ({delta_x:g}, {delta_y:g})'
			)
		else:
			await self._active_page().mouse.wheel(delta_x, delta_y)
		return f'scrolled ({delta_x:g}, {delta_y:g})'

	async def _drag(self, params: dict[str, Any]) -> str:
		if self._first(params, 'source_index', 'from_index') is not None:
			source_index = int(self._first(params, 'source_index', 'from_index'))
			target_value = self._first(params, 'target_index', 'to_index')
			if target_value is None:
				raise ValueError('element drag requires target_index/to_index')
			target_index = int(target_value)
			source_binding = self._binding_for_index(source_index)
			target_binding = self._binding_for_index(target_index)
			if self._is_backend_binding(source_binding) and self._is_backend_binding(target_binding):
				# A drag cannot safely auto-scroll source and target independently: the
				# second scroll could invalidate the first point before mouse-down.
				start_x, start_y = await self._backend_pointer(source_binding, source_index, ensure_visible=False)
				end_x, end_y = await self._backend_pointer(target_binding, target_index, ensure_visible=False)
				viewport = await self._viewport(self._active_page())
				if not (self._point_in_viewport(start_x, start_y, viewport) and self._point_in_viewport(end_x, end_y, viewport)):
					raise ValueError('Both drag endpoints must be visible; scroll and observe before dragging')
				mouse = self._active_page().mouse
				await mouse.move(start_x, start_y)
				await mouse.down()
				await mouse.move(end_x, end_y, steps=12)
				await mouse.up()
			else:
				source = self._locator_for_index(source_index)
				target = self._locator_for_index(target_index)
				await source.drag_to(target, timeout=self.action_timeout_ms)
			return f'dragged element {source_index} to {target_index}'
		start_x = float(self._first(params, 'start_x', 'from_x', 'x1', 'x'))
		start_y = float(self._first(params, 'start_y', 'from_y', 'y1', 'y'))
		end_x = float(self._first(params, 'end_x', 'to_x', 'x2'))
		end_y = float(self._first(params, 'end_y', 'to_y', 'y2'))
		profile = str(self._first(params, 'profile', default='') or '').casefold()
		mouse = self._active_page().mouse
		if profile == 'human':
			from browser_use.webretriever.verification_vision import human_drag_waypoints

			await mouse.move(start_x, start_y)
			await asyncio.sleep(0.08)
			await mouse.down()
			for point_x, point_y, delay_ms in human_drag_waypoints(start_x, start_y, end_x, end_y):
				await mouse.move(point_x, point_y, steps=1)
				if delay_ms > 0:
					await asyncio.sleep(delay_ms / 1000)
			await asyncio.sleep(0.05)
			await mouse.up()
		else:
			await mouse.move(start_x, start_y)
			await mouse.down()
			await mouse.move(end_x, end_y, steps=12)
			await mouse.up()
		return f'dragged ({start_x:g}, {start_y:g}) to ({end_x:g}, {end_y:g})'

	async def _tab_action(self, params: dict[str, Any]) -> str:
		operation = str(self._first(params, 'operation', 'tab_action', 'command', default='switch')).lower()
		pages = self._live_owned_pages()
		if operation in {'new', 'open', 'create'}:
			page = await self.context.new_page()
			self._register_page(page, make_active=True)
			await self._configure_owned_page(page)
			url = self._first(params, 'url')
			if url is not None:
				url = str(url)
				self._validate_navigation_url(url)
				if is_forbidden_search_url(url):
					raise ValueError('External search engines are prohibited by the challenge rules')
				download_started = await self._goto_exact(page, url)
				if download_started:
					self._record_url(url, unless_last=True)
			return f'opened tab {len(self._live_owned_pages()) - 1}'
		if not pages:
			raise RuntimeError('No live tabs')
		if operation in {'close', 'remove'}:
			page = self._select_tab(params, pages)
			await page.close(run_before_unload=False)
			self.page = self._live_owned_pages()[-1] if self._live_owned_pages() else None
			return 'closed tab'
		if operation in {'next', 'previous', 'prev'}:
			current = pages.index(self._active_page())
			delta = 1 if operation == 'next' else -1
			page = pages[(current + delta) % len(pages)]
		else:
			page = self._select_tab(params, pages)
		self.page = page
		self._element_bindings.clear()
		self._legacy_marker_bindings_active = False
		await page.bring_to_front()
		return f'switched to tab {pages.index(page)}: {page.url}'

	async def _read(self, params: dict[str, Any]) -> str:
		if self._has_target(params):
			if self._has_explicit_selector(params):
				text = await self._target_locator(params).inner_text(timeout=self.action_timeout_ms)
				return text[:_MAX_PAGE_TEXT]
			index = self._target_index(params)
			binding = self._binding_for_index(index)
			if self._is_backend_binding(binding):
				return await self._backend_read(binding, index)
			text = await self._locator_for_index(index).inner_text(timeout=self.action_timeout_ms)
			return text[:_MAX_PAGE_TEXT]
		return await self._collect_page_text(self._active_page())

	async def _find(self, params: dict[str, Any]) -> str:
		await self._drain_download_tasks()
		query = str(self._first(params, 'query', 'text', 'value', default='')).strip()
		if not query:
			raise ValueError('find requires query/text')
		matches: list[str] = []
		revealed_control = False
		for frame_index, frame in enumerate(self._active_page().frames):
			try:
				locator = frame.get_by_text(query, exact=False)
				count = min(await locator.count(), 20)
				for match_index in range(count):
					item = locator.nth(match_index)
					if await item.is_visible():
						text = re.sub(r'\s+', ' ', await item.inner_text(timeout=self.action_timeout_ms)).strip()
						metadata = await item.evaluate(
							"""(element, options) => {
								const interactiveSelector = [
									'a[href]', 'button', 'input:not([type="hidden"])', 'textarea', 'select',
									'[contenteditable="true"]', '[role="button"]', '[role="link"]',
									'[role="tab"]', '[tabindex]:not([tabindex="-1"])'
								].join(',');
								const interactive = element.closest(interactiveSelector)
									|| element.querySelector(interactiveSelector);
								let rect = element.getBoundingClientRect();
								const wasInViewport = rect.bottom > 0 && rect.right > 0
									&& rect.top < window.innerHeight && rect.left < window.innerWidth;
								let revealed = false;
								if (options.reveal && interactive && !wasInViewport) {
									interactive.scrollIntoView({block: 'center', inline: 'nearest'});
									rect = interactive.getBoundingClientRect();
									revealed = true;
								}
								const anchor = (interactive && interactive.matches('a[href]') ? interactive : null)
									|| element.closest('a[href]') || element.querySelector('a[href]');
								return {
									tag: (interactive || element).tagName.toLowerCase(),
									id: (interactive || element).id || '',
									href: anchor ? anchor.href : '',
									// Kept empty solely for compatibility with older locator adapters.
									// Element identity is never read from or written to the page DOM.
									interactiveIndex: '',
									inViewport: rect.bottom > 0 && rect.right > 0
										&& rect.top < window.innerHeight && rect.left < window.innerWidth,
									revealed,
								};
							}""",
							{'reveal': not revealed_control},
						)
						if metadata.get('revealed'):
							revealed_control = True
						details = [
							f'tag={metadata.get("tag", "")}',
							f'in_viewport={str(bool(metadata.get("inViewport"))).lower()}',
						]
						if metadata.get('revealed'):
							details.append('revealed_for_next_observation=true')
						if metadata.get('id'):
							details.append(f'id={metadata["id"]}')
						if metadata.get('href'):
							details.append(f'href={metadata["href"]}')
						matches.append(f'frame {frame_index}: {text[:500]} | {" | ".join(details)}')
						if len(matches) == 20:
							break
			except Exception:
				continue
			if len(matches) == 20:
				break
		page = self._active_page()
		page_text = await self._collect_page_text(page)
		documents: list[dict[str, Any]] = [
			{
				'id': 'page:active',
				'source': 'page',
				'field': 'page_text',
				'url': page.url,
				'text': page_text,
			}
		]
		for index, download in enumerate(self.downloads):
			documents.append(
				{
					'id': f'download:{index}',
					'source': 'download',
					'field': 'download.text',
					'filename': str(download.get('filename', download.get('suggested_filename', 'download'))),
					'url': str(download.get('url', '')),
					'text': self._download_text_for_search(download),
				}
			)
		result = await self._run_text_search(query, documents)
		result['dom_matches'] = matches
		rendered_matches = list(matches)
		for item in result.get('results', []):
			if not isinstance(item, Mapping):
				continue
			if str(item.get('source')) == 'download':
				label = f'download {item.get("filename") or "download"}'
			else:
				label = f'page {item.get("url") or page.url}'
			rendered_matches.append(f'{label}: {item.get("text", "")}')
		result['matches'] = rendered_matches[:20]
		if not rendered_matches:
			return f'No visible text matching {query!r}'
		return json.dumps(result, ensure_ascii=False, indent=2)[:_MAX_PAGE_TEXT]

	async def _run_text_search(self, query: str, documents: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
		"""Search current page/download evidence with the task-local Lunr index."""

		script = Path(__file__).with_name('lunr_text_search.js')
		payload = {'query': query, 'documents': list(documents)}
		try:
			process = await asyncio.create_subprocess_exec(
				'node',
				f'--max-old-space-size={_NETWORK_SEARCH_NODE_HEAP_MIB}',
				str(script),
				stdin=asyncio.subprocess.PIPE,
				stdout=asyncio.subprocess.PIPE,
				stderr=asyncio.subprocess.PIPE,
			)
		except OSError as exc:
			return self._substring_text_search(payload, f'{type(exc).__name__}: {exc}')
		input_bytes = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()
		try:
			stdout, stderr = await asyncio.wait_for(
				process.communicate(input_bytes),
				timeout=_TEXT_SEARCH_TIMEOUT_SECONDS,
			)
		except asyncio.CancelledError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			raise
		except TimeoutError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			return self._substring_text_search(payload, 'Lunr text search exceeded 30 seconds')
		if process.returncode != 0:
			reason = stderr.decode('utf-8', errors='replace')[:500] or f'Node search exited {process.returncode}'
			return self._substring_text_search(payload, reason)
		try:
			result = json.loads(stdout)
		except (json.JSONDecodeError, UnicodeDecodeError) as exc:
			return self._substring_text_search(payload, f'Invalid Node text search output: {exc}')
		if not isinstance(result, dict):
			return self._substring_text_search(payload, 'Node text search output was not an object')
		return result

	@staticmethod
	def _substring_text_search(payload: Mapping[str, Any], reason: str) -> dict[str, Any]:
		"""Keep exact search available if the optional Node runtime is unavailable."""

		query = str(payload.get('query', ''))
		query_folded = unicodedata.normalize('NFKC', query).casefold()
		numeric_only = query_folded.isdigit()
		results: list[dict[str, Any]] = []
		for document in payload.get('documents', []):
			if not isinstance(document, Mapping):
				continue
			text = str(document.get('text', ''))
			folded_text = unicodedata.normalize('NFKC', text).casefold()
			start = folded_text.find(query_folded)
			while (
				start >= 0
				and numeric_only
				and (
					(start > 0 and folded_text[start - 1].isdigit())
					or (start + len(query_folded) < len(folded_text) and folded_text[start + len(query_folded)].isdigit())
				)
			):
				start = folded_text.find(query_folded, start + max(1, len(query_folded)))
			if start < 0:
				continue
			context_start = max(0, start - 500)
			context_end = min(len(text), start + len(query) + 1_500)
			results.append(
				{
					'source_id': document.get('id'),
					'source': document.get('source', 'page'),
					'field': document.get('field', 'text'),
					'filename': document.get('filename', ''),
					'url': document.get('url', ''),
					'text': text[context_start:context_end],
					'start': start,
					'end': start + len(query),
					'exact': True,
					'score': None,
					'matched_fields': [str(document.get('field', 'text'))],
				}
			)
			if len(results) == 20:
				break
		return {
			'search_mode': 'substring_fallback',
			'fallback_reason': reason,
			'query': query,
			'query_truncated': False,
			'indexed_documents': len(payload.get('documents', [])),
			'indexed_chunks': 0,
			'exact_match_count': len(results),
			'results': results,
		}

	def _download_text_for_search(self, download: Mapping[str, Any]) -> str:
		"""Return searchable download text, lazily reopening truncated text files.

		Observation previews stay bounded, but ``find_text`` must be able to locate
		rows near the end of a competition CSV.  The file itself was already saved
		through Playwright and capped by ``_MAX_DOWNLOAD_BYTES``.
		"""
		text = str(download.get('text', ''))
		if not download.get('text_truncated'):
			return text
		path_value = download.get('path')
		if not isinstance(path_value, str):
			return text
		path = Path(path_value)
		if path.suffix.lower() not in {'.csv', '.tsv', '.txt', '.json', '.xml'}:
			return text
		try:
			resolved = path.resolve(strict=True)
			if not resolved.is_relative_to(self.download_dir.resolve()):
				return text
			if resolved.stat().st_size > _MAX_DOWNLOAD_BYTES:
				return text
			return resolved.read_text(encoding='utf-8-sig', errors='replace')
		except OSError:
			return text

	async def _inspect_network(self, params: dict[str, Any]) -> str:
		request_id_value = self._first(params, 'request_id')
		query = str(self._first(params, 'query', 'text', default=''))
		cursor = self._first(params, 'network_cursor', 'cursor')
		if cursor is not None and request_id_value is None:
			raise ValueError('inspect_network network_cursor requires request_id')
		if request_id_value is not None:
			request_id = int(request_id_value)
			if query:
				if cursor is not None:
					raise ValueError('inspect_network network_cursor cannot be combined with text')
				return await self._inspect_network_request_search(request_id, query)
			return await self._inspect_network_request(
				request_id,
				cursor=cursor,
			)

		limit = min(max(int(self._first(params, 'limit', default=10)), 1), 50)
		if not query:
			items = []
			for item in self.all_requests[-limit:]:
				output = copy.deepcopy(item)
				request_id = self._request_ids_by_entry.get(id(item))
				if request_id is not None:
					output['request_id'] = request_id
				items.append(output)
			return json.dumps(items, ensure_ascii=False, indent=2)[:_MAX_PAGE_TEXT]

		await self._settle_network_capture_bounded(_NETWORK_SEARCH_TIMEOUT_SECONDS)
		requests = []
		for item in self.all_requests:
			request_id = self._request_ids_by_entry.get(id(item))
			if request_id is None:
				continue
			requests.append(self._network_search_payload_request(request_id, item))
		return await self._render_network_search(query, requests)

	async def _inspect_network_request_search(self, request_id: int, query: str) -> str:
		"""Search one materialized XHR/fetch response using the standard search output."""

		if request_id not in self._request_entries_by_id:
			raise ValueError(f'Unknown inspect_network request_id {request_id}')
		await self._settle_network_capture_bounded(_NETWORK_SEARCH_TIMEOUT_SECONDS)
		snapshot = await self._materialize_inspect_request(request_id)
		request = self._network_search_payload_request(
			request_id,
			snapshot,
			fallback=self._request_entries_by_id[request_id],
		)
		return await self._render_network_search(query, [request])

	@staticmethod
	def _network_search_payload_request(
		request_id: int,
		source: Mapping[str, Any],
		*,
		fallback: Mapping[str, Any] | None = None,
	) -> dict[str, Any]:
		"""Normalize one captured or materialized request for the shared search module."""

		def value(name: str, default: Any = None) -> Any:
			current = source.get(name)
			if current is None and fallback is not None:
				current = fallback.get(name)
			return default if current is None else current

		response_body = source.get('response_body')
		if isinstance(response_body, str):
			response_body_truncated = bool(source.get('response_body_truncated', False))
		elif fallback is not None:
			response_body = fallback.get('response_body')
			response_body_truncated = bool(fallback.get('response_body_truncated', False))
		else:
			response_body_truncated = bool(source.get('response_body_truncated', False))

		return {
			'request_id': request_id,
			'timestamp': value('timestamp'),
			'url': value('url'),
			'method': value('method'),
			'status': value('status'),
			'resource_type': value('resource_type'),
			'post_data': value('post_data'),
			'json_data': value('json_data'),
			'response_headers': value('response_headers'),
			'response_body': response_body,
			'response_body_truncated': response_body_truncated,
		}

	async def _render_network_search(self, query: str, requests: Sequence[Mapping[str, Any]]) -> str:
		payload = {'query': query, 'requests': requests}
		result = await self._run_network_search(payload)
		# Lunr tokenizes a query, so a response containing only a few generic
		# terms can rank highly without containing the complete requested phrase.
		# Count literal response-body matches separately so callers can distinguish
		# relevance from an exact hit. Request metadata is intentionally excluded:
		# a phrase echoed in a URL or request payload is not retrieved evidence.
		result = {
			'exact_match_count': self._exact_network_response_match_count(query, requests),
			**result,
		}
		return json.dumps(result, ensure_ascii=False, indent=2)[:_MAX_PAGE_TEXT]

	@staticmethod
	def _exact_network_response_match_count(query: str, requests: Sequence[Mapping[str, Any]]) -> int:
		"""Count captured response bodies containing ``query`` as one full phrase.

		The count is per captured request (rather than per textual occurrence), and
		uses Unicode NFKC normalization plus case folding. This keeps a title match
		stable across harmless compatibility variants while preserving whitespace and
		word boundaries: token-level or partial Lunr matches do not qualify.
		"""

		normalized_query = unicodedata.normalize('NFKC', query).casefold()
		if not normalized_query:
			return 0
		return sum(
			1
			for request in requests
			if isinstance(request.get('response_body'), str)
			and normalized_query in unicodedata.normalize('NFKC', request['response_body']).casefold()
		)

	async def _inspect_network_request(self, request_id: int, *, cursor: Any = None) -> str:
		if request_id not in self._request_entries_by_id:
			raise ValueError(f'Unknown inspect_network request_id {request_id}')

		metadata_included = cursor is None
		if cursor is None:
			await self._settle_network_capture_bounded(_NETWORK_SEARCH_TIMEOUT_SECONDS)
			snapshot = await self._materialize_inspect_request(request_id)
			offset = 0
			page_number = 1
		else:
			cursor_state = self._network_page_cursors.get(str(cursor))
			if cursor_state is None or cursor_state[0] != request_id:
				raise ValueError('inspect_network cursor is invalid for this request_id')
			snapshot = await self.materialize_network_request(request_id, max_body_bytes=_MAX_DOWNLOAD_BYTES)
			offset = cursor_state[1]
			page_number = cursor_state[3]

		body_state, body_encoding, body_data = self._inspect_response_body(snapshot, request_id)
		body_sha256 = str(snapshot.get('response_body_sha256', ''))
		if not body_sha256 and body_data:
			if body_encoding == 'base64':
				with contextlib.suppress(ValueError):
					body_sha256 = hashlib.sha256(base64.b64decode(body_data, validate=True)).hexdigest()
			else:
				body_sha256 = hashlib.sha256(body_data.encode('utf-8')).hexdigest()

		if cursor is not None and cursor_state[2] != body_sha256:
			raise ValueError('inspect_network cursor no longer matches the response body')

		response = {
			'status': snapshot.get('status'),
			'headers': copy.deepcopy(snapshot.get('response_headers', {})),
			'body_encoding': body_encoding,
			'body_bytes': snapshot.get('response_body_bytes'),
			'body_characters': len(body_data),
			'body_sha256': body_sha256 or None,
			'body_truncated': bool(snapshot.get('response_body_truncated')),
			'error': snapshot.get('response_error'),
		}
		output: dict[str, Any] = {
			'mode': 'request',
			'request_id': request_id,
			'url': snapshot.get('url'),
			'body_state': body_state,
			'response': response,
			'page': {
				'metadata_included': metadata_included,
				'number': page_number,
				'offset': offset,
				'data': '',
				'next_cursor': None,
			},
		}
		if metadata_included:
			response_only = {
				'status',
				'response_headers',
				'response_body',
				'response_json',
				'response_body_base64',
				'response_body_encoding',
				'response_body_bytes',
				'response_body_sha256',
				'response_body_state',
				'response_body_truncated',
				'response_body_omitted',
				'response_error',
				'failure',
			}
			output['request'] = {key: copy.deepcopy(value) for key, value in snapshot.items() if key not in response_only}

		page_length = self._fit_network_body_page(output, body_data, offset)
		page_data = body_data[offset : offset + page_length]
		next_offset = offset + len(page_data)
		next_cursor: str | None = None
		if next_offset < len(body_data):
			next_cursor = secrets.token_urlsafe(24)
			self._network_page_cursors[next_cursor] = (request_id, next_offset, body_sha256, page_number + 1)
		output['page']['data'] = page_data
		output['page']['next_cursor'] = next_cursor
		rendered = json.dumps(output, ensure_ascii=False, indent=2)
		if len(rendered) > _MAX_PAGE_TEXT:
			raise ValueError('inspect_network request metadata exceeds the action output limit')
		return rendered

	@staticmethod
	def _fit_network_body_page(output: dict[str, Any], body_data: str, offset: int) -> int:
		remaining = max(0, len(body_data) - offset)
		length = min(_NETWORK_BODY_PAGE_CHARACTERS, remaining)
		cursor_placeholder = 'x' * 32
		while True:
			output['page']['data'] = body_data[offset : offset + length]
			output['page']['next_cursor'] = cursor_placeholder if length < remaining else None
			rendered_length = len(json.dumps(output, ensure_ascii=False, indent=2))
			if rendered_length <= _MAX_PAGE_TEXT:
				if remaining and length == 0:
					raise ValueError('inspect_network request metadata leaves no room for response data')
				return length
			if length == 0:
				raise ValueError('inspect_network request metadata exceeds the action output limit')
			length = max(0, length - max(1, rendered_length - _MAX_PAGE_TEXT))

	async def _materialize_inspect_request(self, request_id: int) -> dict[str, Any]:
		deadline = asyncio.get_running_loop().time() + _NETWORK_SEARCH_TIMEOUT_SECONDS
		while True:
			snapshot = await self.materialize_network_request(request_id, max_body_bytes=_MAX_DOWNLOAD_BYTES)
			if snapshot.get('response_body_state') != 'pending':
				return snapshot
			remaining = deadline - asyncio.get_running_loop().time()
			if remaining <= 0:
				return snapshot
			await asyncio.sleep(min(0.1, remaining))

	def _inspect_response_body(self, snapshot: Mapping[str, Any], request_id: int) -> tuple[str, str, str]:
		state = str(snapshot.get('response_body_state', 'unavailable'))
		if isinstance(snapshot.get('response_body'), str):
			return ('complete' if state == 'complete' else 'partial', 'text', str(snapshot['response_body']))
		if isinstance(snapshot.get('response_body_base64'), str):
			return ('complete' if state == 'complete' else 'partial', 'base64', str(snapshot['response_body_base64']))

		captured = self._request_entries_by_id[request_id]
		if isinstance(captured.get('response_body'), str):
			return 'partial', 'text', str(captured['response_body'])
		if isinstance(captured.get('response_body_base64'), str):
			return 'partial', 'base64', str(captured['response_body_base64'])
		if state == 'body_too_large':
			return 'body_too_large', 'none', ''
		if state in {'error', 'failed'}:
			return 'error', 'none', ''
		return 'unavailable', 'none', ''

	async def _settle_network_capture_bounded(self, timeout_seconds: float) -> None:
		deadline = asyncio.get_running_loop().time() + max(0.0, timeout_seconds)
		while self._background_tasks:
			remaining = deadline - asyncio.get_running_loop().time()
			if remaining <= 0:
				return
			tasks = tuple(self._background_tasks)
			done, _ = await asyncio.wait(tasks, timeout=remaining)
			if not done:
				return
			self._background_tasks.difference_update(done)
			for task in done:
				with contextlib.suppress(asyncio.CancelledError, Exception):
					task.result()

	async def _run_network_search(self, payload: Mapping[str, Any]) -> dict[str, Any]:
		script = Path(__file__).with_name('lunr_network_search.js')
		try:
			process = await asyncio.create_subprocess_exec(
				'node',
				f'--max-old-space-size={_NETWORK_SEARCH_NODE_HEAP_MIB}',
				str(script),
				stdin=asyncio.subprocess.PIPE,
				stdout=asyncio.subprocess.PIPE,
				stderr=asyncio.subprocess.PIPE,
			)
		except OSError as exc:
			return self._substring_network_search(payload, f'{type(exc).__name__}: {exc}')
		input_bytes = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()
		try:
			stdout, stderr = await asyncio.wait_for(
				process.communicate(input_bytes),
				timeout=_NETWORK_SEARCH_TIMEOUT_SECONDS,
			)
		except asyncio.CancelledError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			raise
		except TimeoutError:
			with contextlib.suppress(ProcessLookupError):
				process.kill()
			await process.wait()
			return self._substring_network_search(payload, 'Node search exceeded 30 seconds')
		if process.returncode != 0:
			reason = stderr.decode('utf-8', errors='replace')[:500] or f'Node search exited {process.returncode}'
			return self._substring_network_search(payload, reason)
		try:
			result = json.loads(stdout)
		except (json.JSONDecodeError, UnicodeDecodeError) as exc:
			return self._substring_network_search(payload, f'Invalid Node search output: {exc}')
		if not isinstance(result, dict):
			return self._substring_network_search(payload, 'Node search output was not an object')
		return result

	@staticmethod
	def _substring_network_search(payload: Mapping[str, Any], reason: str) -> dict[str, Any]:
		query = str(payload.get('query', ''))
		query_lower = query.casefold()
		requests = payload.get('requests')
		if not isinstance(requests, list):
			requests = []
		results = []
		for request in reversed(requests):
			if not isinstance(request, Mapping):
				continue
			serialized = json.dumps(request, ensure_ascii=False)
			start = serialized.casefold().find(query_lower)
			if start < 0:
				continue
			context_start = max(0, start - 500)
			context_end = min(len(serialized), start + len(query) + 1_000)
			results.append(
				{
					'rank': len(results) + 1,
					'request_id': request.get('request_id'),
					'score': None,
					'matched_query_terms': [query],
					'matched_fields': ['serialized_request'],
					'duplicate_count': 1,
					'request': {
						'timestamp': request.get('timestamp'),
						'url': request.get('url'),
						'method': request.get('method'),
						'status': request.get('status'),
						'resource_type': request.get('resource_type'),
						'post_data': request.get('post_data'),
					},
					'matched_chunks': [{'score': None, 'text': serialized[context_start:context_end]}],
				}
			)
			if len(results) == 10:
				break
		return {
			'search_mode': 'substring_fallback',
			'fallback_reason': reason,
			'query': query,
			'indexed_requests': len(requests),
			'pending_response_bodies': sum(
				1 for request in requests if isinstance(request, Mapping) and request.get('status') is None
			),
			'omitted_due_to_budget': 0,
			'results': results,
		}

	@classmethod
	def _calculate(cls, params: dict[str, Any]) -> str:
		"""Deterministically compare a browser-observed numeric series."""
		operation = str(cls._required(params, 'operation'))
		raw_series = cls._required(params, 'text')
		try:
			payload = json.loads(str(raw_series))
		except json.JSONDecodeError as exc:
			raise ValueError(f'calculate text must be valid JSON: {exc.msg}') from exc

		if isinstance(payload, dict):
			raw_points = list(payload.items())
		elif isinstance(payload, list):
			raw_points = []
			for position, item in enumerate(payload):
				if not isinstance(item, Mapping) or 'label' not in item or 'value' not in item:
					raise ValueError(f'calculate series item {position} must contain label and value')
				raw_points.append((item['label'], item['value']))
		else:
			raise ValueError('calculate text must encode an object or a list of label/value objects')
		if not 1 <= len(raw_points) <= 500:
			raise ValueError('calculate series must contain between 1 and 500 points')

		points: list[tuple[str, float]] = []
		for raw_label, raw_value in raw_points:
			if isinstance(raw_value, bool):
				raise ValueError(f'calculate value for {raw_label!r} must be numeric')
			try:
				value = float(raw_value)
			except (TypeError, ValueError) as exc:
				raise ValueError(f'calculate value for {raw_label!r} must be numeric') from exc
			if not math.isfinite(value):
				raise ValueError(f'calculate value for {raw_label!r} must be finite')
			points.append((str(raw_label), value))

		if operation in {'argmax', 'argmin'}:
			reverse = operation == 'argmax'
			ranked = sorted(points, key=lambda point: point[1], reverse=reverse)
			return json.dumps(
				{
					'operation': operation,
					'compared_count': len(points),
					'winner': {'label': ranked[0][0], 'value': ranked[0][1]},
					'top_candidates': [{'label': label, 'value': value} for label, value in ranked[:10]],
				},
				ensure_ascii=False,
				separators=(',', ':'),
			)

		if operation not in {'argmax_difference', 'argmin_difference', 'argmax_growth', 'argmin_growth'}:
			raise ValueError(f'unsupported calculation operation: {operation!r}')
		if len(points) < 2:
			raise ValueError(f'{operation} requires at least two points')

		try:
			ordered = sorted(points, key=lambda point: float(point[0]))
		except ValueError:
			ordered = points
		comparisons: list[dict[str, str | float]] = []
		for (previous_label, previous_value), (label, value) in zip(ordered, ordered[1:], strict=False):
			difference = value - previous_value
			if operation.endswith('growth'):
				if previous_value == 0:
					continue
				metric = difference / abs(previous_value)
			else:
				metric = difference
			comparisons.append(
				{
					'label': label,
					'value': value,
					'previous_label': previous_label,
					'previous_value': previous_value,
					'metric': metric,
				}
			)
		if not comparisons:
			raise ValueError(f'{operation} has no eligible comparisons')
		reverse = operation.startswith('argmax')
		ranked_comparisons = sorted(comparisons, key=lambda item: float(item['metric']), reverse=reverse)
		return json.dumps(
			{
				'operation': operation,
				'point_count': len(points),
				'compared_count': len(comparisons),
				'winner': ranked_comparisons[0],
				'top_candidates': ranked_comparisons[:10],
			},
			ensure_ascii=False,
			separators=(',', ':'),
		)

	async def _enforce_search_policy(
		self,
		previous_page: Page | None = None,
		previous_url: str | None = None,
		known_page_ids: set[int] | None = None,
	) -> None:
		await self._drain_policy_tasks()
		for page in list(self._live_owned_pages()):
			if not is_forbidden_search_url(page.url):
				if page.url and page.url not in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
					self._last_safe_urls[id(page)] = page.url
				continue
			if known_page_ids is not None and id(page) not in known_page_ids:
				self.logger.warning('Closing tab that escaped to prohibited search engine: %s', redact_cdp_url(page.url))
				with contextlib.suppress(Exception):
					await page.close(run_before_unload=False)
				continue
			fallback = self._last_safe_urls.get(id(page))
			if page is previous_page and previous_url and not is_forbidden_search_url(previous_url):
				fallback = previous_url
			await self._rollback_forbidden_page(page, fallback)
		self._active_page()

	async def _rollback_forbidden_page(self, page: Page, fallback: str | None = None) -> None:
		try:
			if page.is_closed() or not is_forbidden_search_url(page.url):
				return
			self.logger.warning('Rolling back prohibited external search-engine navigation: %s', redact_cdp_url(page.url))
			page_safe_url = self._last_safe_urls.get(id(page))
			other_safe_pages = [
				item for item in self._live_owned_pages() if item is not page and not is_forbidden_search_url(item.url)
			]
			if fallback is None and page_safe_url is None and other_safe_pages:
				# A newly opened search-engine popup has no meaningful history to
				# restore. Close it and return focus to the originating safe tab.
				await page.close(run_before_unload=False)
				self.page = other_safe_pages[-1]
				return
			fallback = fallback or page_safe_url or self.website
			with contextlib.suppress(Exception):
				await page.go_back(wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			if not page.is_closed() and is_forbidden_search_url(page.url) and fallback and not is_forbidden_search_url(fallback):
				await page.goto(fallback, wait_until='domcontentloaded', timeout=self.navigation_timeout_ms)
			if not page.is_closed() and not is_forbidden_search_url(page.url):
				self.page = page
				self._last_safe_urls[id(page)] = page.url
		finally:
			self._rollback_pages.discard(id(page))

	async def _collect_elements(self, page: Page) -> list[ElementRef]:
		"""Collect indexed controls with CDP, falling back only when unavailable.

		The normal path never writes an identifier into the page.  ``backendNodeId``
		is retained privately in ``_ElementBinding`` while the model continues to
		use this observation's small, globally unique integer indices.
		"""

		try:
			collected = await collect_interactive_elements(self.context, page, logger=self.logger)
		except CdpCollectionError as exc:
			self.logger.warning('CDP DOM collection unavailable; using legacy marker collector: %s', exc)
			return await self._collect_legacy_elements(page)

		self._element_bindings.clear()
		self._legacy_marker_bindings_active = False
		result: list[ElementRef] = []
		bindings: dict[int, _ElementBinding] = {}
		for index, item in enumerate(collected):
			ref = ElementRef(
				index=index,
				tag=item.tag,
				text=item.text,
				role=item.role,
				name=item.name,
				placeholder=item.placeholder,
				href=item.href,
				input_type=item.input_type,
				checked=item.checked,
				frame_index=item.frame_index,
				frame_url=item.frame_url,
				x=item.x,
				y=item.y,
				width=item.width,
				height=item.height,
				backend_node_id=item.backend_node_id,
				frame_id=item.frame_id,
				signals=item.signals,
			)
			result.append(ref)
			bindings[index] = _ElementBinding(
				frame=item.frame,
				backend_node_id=item.backend_node_id,
				cdp_target=item.cdp_target,
				coordinate_frame=item.coordinate_frame,
				frame_id=item.frame_id,
				read_text=item.read_text,
				options=item.options,
			)
		self._element_bindings = bindings
		return result

	async def _collect_legacy_elements(self, page: Page) -> list[ElementRef]:
		"""Compatibility fallback for test doubles or an unavailable CDP target."""

		self._element_bindings.clear()
		self._legacy_marker_bindings_active = True
		result: list[ElementRef] = []
		next_index = 0
		for frame_index, frame in enumerate(page.frames):
			try:
				items = await frame.evaluate(
					_MARK_ELEMENTS_JS,
					{
						'start': next_index,
						'markerAttribute': _INTERACTIVE_ATTRIBUTE,
						'overlayClass': _OVERLAY_CLASS,
						'frameIndex': frame_index,
						'maxText': _MAX_ELEMENT_TEXT,
					},
				)
			except Exception as exc:
				self.logger.debug('Skipping inaccessible frame %s: %s', redact_cdp_url(frame.url), exc)
				continue
			for item in items or []:
				index = int(item['index'])
				selector = f'[{_INTERACTIVE_ATTRIBUTE}="{index}"]'
				ref = ElementRef(
					index=index,
					tag=str(item.get('tag', 'element')),
					text=str(item.get('text', '')),
					role=str(item.get('role', '')),
					name=str(item.get('name', '')),
					placeholder=str(item.get('placeholder', '')),
					href=str(item.get('href', '')),
					input_type=str(item.get('input_type', '')),
					checked=str(item.get('checked', '')),
					frame_index=frame_index,
					frame_url=frame.url,
					x=float(item.get('x', 0)),
					y=float(item.get('y', 0)),
					width=float(item.get('width', 0)),
					height=float(item.get('height', 0)),
					selector=selector,
				)
				result.append(ref)
				self._element_bindings[index] = _ElementBinding(frame=frame, selector=selector)
				next_index = max(next_index, index + 1)
		return result

	async def _clear_markers(self, *, remove_attributes: bool) -> None:
		# The normal CDP collector never injects DOM state. Avoid even querying for
		# our legacy class in that path: a business page could coincidentally use the
		# same class name, and removing it would violate the non-mutating contract.
		if not self._legacy_marker_bindings_active:
			return
		page = self.page
		if page is None or page.is_closed():
			return
		for frame in page.frames:
			with contextlib.suppress(Exception):
				await frame.evaluate(
					_CLEAR_MARKERS_JS,
					{
						'markerAttribute': _INTERACTIVE_ATTRIBUTE,
						'overlayClass': _OVERLAY_CLASS,
						'removeAttributes': remove_attributes and self._legacy_marker_bindings_active,
					},
				)
		if remove_attributes:
			self._legacy_marker_bindings_active = False

	async def _collect_page_text(self, page: Page) -> str:
		parts: list[str] = []
		remaining = _MAX_PAGE_TEXT
		for frame_index, frame in enumerate(page.frames):
			if remaining <= 0:
				break
			try:
				text = await frame.locator('body').inner_text(timeout=min(self.action_timeout_ms, 5_000))
			except Exception:
				continue
			text = re.sub(r'[ \t]+', ' ', text).strip()
			if not text:
				continue
			prefix = '' if frame_index == 0 else f'\n[Frame {frame_index}: {frame.url}]\n'
			chunk = (prefix + text)[:remaining]
			parts.append(chunk)
			remaining -= len(chunk)
		return '\n'.join(parts)

	async def _viewport(self, page: Page) -> dict[str, int]:
		viewport = page.viewport_size
		if viewport:
			return {'width': int(viewport['width']), 'height': int(viewport['height'])}
		try:
			value = await page.evaluate('() => ({width: window.innerWidth, height: window.innerHeight})')
			return {'width': int(value['width']), 'height': int(value['height'])}
		except Exception:
			return {'width': 0, 'height': 0}

	async def _tabs(self) -> list[dict[str, Any]]:
		active = self._active_page()
		result: list[dict[str, Any]] = []
		for index, page in enumerate(self._live_owned_pages()):
			result.append(
				{
					'index': index,
					'url': page.url,
					'title': await self._page_title(page),
					'active': page is active,
				}
			)
		return result

	async def _page_title(self, page: Page) -> str:
		try:
			return await page.title()
		except Exception:
			return ''

	def _binding_for_index(self, index: int) -> _ElementBinding:
		binding = self._element_bindings.get(index)
		if binding is None:
			raise ValueError(f'Unknown element index {index}; call observe() before interacting')
		return binding

	@staticmethod
	def _is_backend_binding(binding: _ElementBinding) -> bool:
		return binding.backend_node_id > 0 and binding.cdp_target is not None

	def _target_binding(self, params: Mapping[str, Any]) -> _ElementBinding:
		return self._binding_for_index(self._target_index(params))

	@staticmethod
	def _has_explicit_selector(params: Mapping[str, Any]) -> bool:
		return params.get('selector') is not None

	def _stale_element_error(self, index: int) -> ValueError:
		message = f'Element {index} is stale or the page changed; call observe() to obtain current element IDs'
		return ValueError(message)

	@contextlib.asynccontextmanager
	async def _backend_session(self, binding: _ElementBinding, index: int):
		"""Create a short-lived, read-only CDP session for an observed node."""

		if binding.cdp_target is None:
			raise ValueError('Observed element has no CDP backend binding')
		active_page = self._active_page()
		try:
			binding_page = binding.frame.page
		except Exception:
			binding_page = active_page
		if binding_page is not active_page:
			raise self._stale_element_error(index)
		session: CDPSession | None = None
		try:
			session = await self.context.new_cdp_session(binding.cdp_target)
			yield session
		finally:
			if session is not None:
				with contextlib.suppress(Exception):
					await session.detach()

	async def _ensure_coordinate_frame_visible(self, binding: _ElementBinding, index: int) -> None:
		"""Use Playwright to reveal an iframe host before pointer interaction."""

		page = self._active_page()
		frame = binding.coordinate_frame or binding.frame
		if frame is page.main_frame:
			return
		try:
			owner = await frame.frame_element()
			await owner.scroll_into_view_if_needed(timeout=self.action_timeout_ms)
		except Exception as exc:
			raise self._stale_element_error(index) from exc

	@staticmethod
	def _quad_center(quads: Any) -> tuple[float, float] | None:
		best: tuple[float, float, float] | None = None
		for quad in quads or []:
			if not isinstance(quad, Sequence) or isinstance(quad, (str, bytes)) or len(quad) < 8:
				continue
			try:
				xs = [float(quad[position]) for position in range(0, 8, 2)]
				ys = [float(quad[position]) for position in range(1, 8, 2)]
			except (TypeError, ValueError):
				continue
			width = max(xs) - min(xs)
			height = max(ys) - min(ys)
			area = max(0.0, width * height)
			candidate = (area, sum(xs) / len(xs), sum(ys) / len(ys))
			if best is None or candidate[0] > best[0]:
				best = candidate
		return (best[1], best[2]) if best is not None else None

	async def _backend_raw_pointer(self, binding: _ElementBinding, index: int) -> tuple[float, float]:
		"""Read a current CDP quad without changing page state."""

		try:
			async with self._backend_session(binding, index) as session:
				center: tuple[float, float] | None = None
				try:
					quads_result = await session.send('DOM.getContentQuads', {'backendNodeId': binding.backend_node_id})
					center = self._quad_center(quads_result.get('quads') if isinstance(quads_result, Mapping) else None)
				except Exception as exc:
					self.logger.debug('Could not read CDP content quads for element %s: %s', index, exc)
				if center is None:
					try:
						box_result = await session.send('DOM.getBoxModel', {'backendNodeId': binding.backend_node_id})
						model = box_result.get('model') if isinstance(box_result, Mapping) else None
						center = self._quad_center([model.get('content')]) if isinstance(model, Mapping) else None
					except Exception as exc:
						self.logger.debug('Could not read CDP box model for element %s: %s', index, exc)
			if center is None:
				raise RuntimeError('the backend node has no usable content quad')
		except Exception as exc:
			raise self._stale_element_error(index) from exc
		return center

	async def _backend_viewport_point(self, binding: _ElementBinding, index: int) -> tuple[float, float]:
		"""Project an observed node's read-only CDP quad into the top viewport."""

		x, y = await self._backend_raw_pointer(binding, index)
		if binding.coordinate_frame is None:
			return x, y
		try:
			owner = await binding.coordinate_frame.frame_element()
			box = await owner.bounding_box()
			metrics = await owner.evaluate(
				'(element) => ({left: element.clientLeft, top: element.clientTop, '
				'offsetWidth: element.offsetWidth, offsetHeight: element.offsetHeight})'
			)
			if not box:
				raise RuntimeError('iframe owner has no bounding box')
			left = float(metrics.get('left', 0)) if isinstance(metrics, Mapping) else 0.0
			top = float(metrics.get('top', 0)) if isinstance(metrics, Mapping) else 0.0
			offset_width = float(metrics.get('offsetWidth', 0)) if isinstance(metrics, Mapping) else 0.0
			offset_height = float(metrics.get('offsetHeight', 0)) if isinstance(metrics, Mapping) else 0.0
			scale_x = float(box['width']) / offset_width if offset_width > 0 else 1.0
			scale_y = float(box['height']) / offset_height if offset_height > 0 else 1.0
			return float(box['x']) + (left + x) * scale_x, float(box['y']) + (top + y) * scale_y
		except Exception as exc:
			raise self._stale_element_error(index) from exc

	@staticmethod
	def _point_in_viewport(x: float, y: float, viewport: Mapping[str, int]) -> bool:
		return 0 <= x < float(viewport.get('width', 0)) and 0 <= y < float(viewport.get('height', 0))

	async def _scroll_towards_backend_point(
		self,
		binding: _ElementBinding,
		x: float,
		y: float,
		viewport: Mapping[str, int],
	) -> None:
		"""Use Playwright wheel input to reveal a nearby observed target."""

		page = self._active_page()
		if binding.frame is not page.main_frame:
			with contextlib.suppress(Exception):
				owner = await binding.frame.frame_element()
				box = await owner.bounding_box()
				if box:
					await page.mouse.move(float(box['x']) + float(box['width']) / 2, float(box['y']) + float(box['height']) / 2)
		else:
			# Mouse wheel follows the element under the pointer. Keep a root-page
			# reveal from accidentally scrolling an iframe or nested scroller that
			# happened to receive the prior action.
			await page.mouse.move(1, 1)
		width = max(1.0, float(viewport.get('width', 0)))
		height = max(1.0, float(viewport.get('height', 0)))
		delta_x = x - width / 2 if x < 0 or x >= width else 0.0
		delta_y = y - height / 2 if y < 0 or y >= height else 0.0
		if delta_x or delta_y:
			await page.mouse.wheel(delta_x, delta_y)
			await page.wait_for_timeout(50)

	async def _backend_pointer(
		self,
		binding: _ElementBinding,
		index: int,
		*,
		ensure_visible: bool = True,
	) -> tuple[float, float]:
		"""Resolve a node with CDP, but reveal and act through Playwright only."""

		await self._ensure_coordinate_frame_visible(binding, index)
		x, y = await self._backend_viewport_point(binding, index)
		if not ensure_visible:
			return x, y
		for _ in range(2):
			viewport = await self._viewport(self._active_page())
			if self._point_in_viewport(x, y, viewport):
				return x, y
			await self._scroll_towards_backend_point(binding, x, y, viewport)
			x, y = await self._backend_viewport_point(binding, index)
		viewport = await self._viewport(self._active_page())
		if not self._point_in_viewport(x, y, viewport):
			raise ValueError(f'Element {index} is outside the current viewport; scroll and observe before interacting')
		return x, y

	async def _backend_focus(self, binding: _ElementBinding, index: int) -> None:
		# A real Playwright click yields focus for input-capable controls, including
		# controls in a closed shadow tree that have no public locator.
		await self._backend_click(binding, index)

	async def _backend_click(self, binding: _ElementBinding, index: int, *, click_count: int = 1) -> None:
		x, y = await self._backend_pointer(binding, index)
		await self._active_page().mouse.click(x, y, click_count=click_count)

	async def _backend_hover(self, binding: _ElementBinding, index: int) -> None:
		x, y = await self._backend_pointer(binding, index)
		await self._active_page().mouse.move(x, y)

	async def _backend_type(self, binding: _ElementBinding, index: int, text: str, *, submit: bool) -> None:
		await self._backend_focus(binding, index)
		keyboard = self._active_page().keyboard
		await keyboard.press('ControlOrMeta+A')
		await keyboard.press('Backspace')
		insert_text = getattr(keyboard, 'insert_text', None)
		if callable(insert_text):
			await insert_text(text)
		else:
			await keyboard.type(text)
		if submit:
			await keyboard.press('Enter')

	async def _backend_select(self, binding: _ElementBinding, index: int, value: str) -> list[str]:
		matched = next(
			(
				(position, option_value)
				for position, (option_value, label) in enumerate(binding.options)
				if value in {option_value, label}
			),
			None,
		)
		if matched is None:
			# The page text and the CDP snapshot are independent observations. A
			# dynamic/native select can therefore have a complete live option list
			# even when the snapshot captured for this binding omitted an option.
			# Refresh only on cache miss, then retain the exact value/label contract.
			try:
				binding.options = await self._capture_backend_select_options(binding, index)
			except Exception as exc:
				self.logger.debug('Could not refresh live options for element %s: %s', index, exc)
			else:
				matched = next(
					(
						(position, option_value)
						for position, (option_value, label) in enumerate(binding.options)
						if value in {option_value, label}
					),
					None,
				)
		if matched is None:
			raise ValueError(f'No option matching {value!r} exists on element {index}; observe again if choices changed')
		position, selected_value = matched
		await self._backend_focus(binding, index)
		keyboard = self._active_page().keyboard
		await keyboard.press('Home')
		for _ in range(position):
			await keyboard.press('ArrowDown')
		await keyboard.press('Enter')
		return [selected_value]

	async def _backend_read(self, binding: _ElementBinding, index: int) -> str:
		return binding.read_text[:_MAX_PAGE_TEXT]

	async def _backend_scroll(self, binding: _ElementBinding, index: int, delta_x: float, delta_y: float) -> None:
		x, y = await self._backend_pointer(binding, index)
		mouse = self._active_page().mouse
		await mouse.move(x, y)
		await mouse.wheel(delta_x, delta_y)

	def _target_locator(self, params: Mapping[str, Any]) -> Locator:
		selector = self._first(params, 'selector')
		if selector is not None:
			return self._active_page().locator(str(selector)).first
		return self._locator_for_index(self._target_index(params))

	def _locator_for_index(self, index: int) -> Locator:
		binding = self._binding_for_index(index)
		if not binding.selector:
			raise ValueError(f'Element {index} is backend-bound and must be acted on through its current CDP binding')
		return binding.frame.locator(binding.selector).first

	def _target_index(self, params: Mapping[str, Any]) -> int:
		value = self._first(params, 'index', 'element_index', 'target_index', 'element_id', 'target')
		if isinstance(value, str):
			match = re.fullmatch(r'\s*\[?(\d+)\]?\s*', value)
			if match:
				value = match.group(1)
		if value is None:
			raise ValueError('Action requires an element index or selector')
		try:
			return int(value)
		except (TypeError, ValueError) as exc:
			raise ValueError(f'Invalid element index: {value!r}') from exc

	def _has_target(self, params: Mapping[str, Any]) -> bool:
		return self._first(params, 'selector', 'index', 'element_index', 'target_index', 'element_id', 'target') is not None

	def _select_tab(self, params: Mapping[str, Any], pages: Sequence[Page]) -> Page:
		value = self._first(params, 'tab_index', 'index', 'target', default=0)
		index = int(value)
		if not -len(pages) <= index < len(pages):
			raise ValueError(f'Tab index out of range: {index}')
		return pages[index]

	def _active_page(self) -> Page:
		if self.page is not None and not self.page.is_closed() and not is_forbidden_search_url(self.page.url):
			return self.page
		live = [page for page in self._live_owned_pages() if not is_forbidden_search_url(page.url)]
		if not live:
			raise RuntimeError('BrowserRuntime has no live safe page')
		self.page = live[-1]
		return self.page

	async def _active_page_for_observation(self) -> Page:
		page = self._active_page()
		if page.url != _DOWNLOAD_PLACEHOLDER_URL:
			return page

		opener: Page | None = None
		with contextlib.suppress(Exception):
			opener = await page.opener()
		if (
			opener is None
			or opener is page
			or opener.is_closed()
			or opener.url == _DOWNLOAD_PLACEHOLDER_URL
			or is_forbidden_search_url(opener.url)
		):
			opener = next(
				(
					candidate
					for candidate in reversed(self._live_owned_pages())
					if candidate is not page
					and candidate.url != _DOWNLOAD_PLACEHOLDER_URL
					and not is_forbidden_search_url(candidate.url)
				),
				None,
			)

		if opener is None:
			with contextlib.suppress(Exception):
				await page.close(run_before_unload=False)
			self.page = None
			raise RuntimeError("Download placeholder page ':' has no live safe opener or fallback page")

		self.logger.info(
			'Closing download placeholder page and restoring active page: %s',
			redact_cdp_url(opener.url),
		)
		self.page = opener
		with contextlib.suppress(Exception):
			await page.close(run_before_unload=False)
		return opener

	def _live_owned_pages(self) -> list[Page]:
		return [page for page in self._owned_pages if not page.is_closed()]

	def _ensure_started(self) -> None:
		if self._closed:
			raise RuntimeError('BrowserRuntime is closed')
		if not self._started or self.page is None:
			raise RuntimeError('Call BrowserRuntime.start(website) first')

	def _record_url(self, url: str, *, unless_last: bool = False) -> None:
		if not url or url in {'about:blank', _DOWNLOAD_PLACEHOLDER_URL}:
			return
		if unless_last and self.visited_urls and self.visited_urls[-1] == url:
			return
		self.visited_urls.append(url)

	def _spawn_background(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._background_tasks.add(task)

	def _spawn_download(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._download_tasks.add(task)

	def _spawn_policy(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._policy_tasks.add(task)

	def _spawn_dialog(self, coroutine: Coroutine[Any, Any, Any]) -> None:
		task = asyncio.create_task(coroutine)
		self._dialog_tasks.add(task)

	async def _drain_policy_tasks(self) -> None:
		while self._policy_tasks:
			tasks = tuple(self._policy_tasks)
			self._policy_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _drain_dialog_tasks(self) -> None:
		while self._dialog_tasks:
			tasks = tuple(self._dialog_tasks)
			self._dialog_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _drain_background_tasks(self) -> None:
		while self._background_tasks:
			tasks = tuple(self._background_tasks)
			self._background_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	async def _drain_download_tasks(self) -> None:
		while self._download_tasks:
			tasks = tuple(self._download_tasks)
			self._download_tasks.difference_update(tasks)
			await asyncio.gather(*tasks, return_exceptions=True)

	def _cancel_pending_navigation(self, page: Page) -> None:
		pending = self._pending_navigations.pop(id(page), None)
		if pending is None or pending.task.done():
			return
		pending.task.add_done_callback(_consume_detached_task_result)
		pending.task.cancel()

	async def _cancel_pending_navigations(self, *, timeout_seconds: float | None = None) -> int:
		pendings = tuple(self._pending_navigations.values())
		self._pending_navigations.clear()
		for pending in pendings:
			if not pending.task.done():
				pending.task.add_done_callback(_consume_detached_task_result)
				pending.task.cancel()
		if pendings:
			cleanup_timeout = min(1.0, self.navigation_first_observation_timeout_ms / 1000)
			if timeout_seconds is not None:
				cleanup_timeout = min(cleanup_timeout, max(0.0, timeout_seconds))
			_, still_pending = await asyncio.wait([pending.task for pending in pendings], timeout=cleanup_timeout)
			if still_pending:
				self.logger.warning('%d pending navigation task(s) resisted bounded cleanup', len(still_pending))
				for task in still_pending:
					self._detach_cleanup_task(task)
			return len(still_pending)
		return 0

	def _record_navigation_notice(self, page: Page, url: str, status: str, error: str = '') -> None:
		self._navigation_notices.append(
			{
				'page_id': id(page),
				'url': url,
				'status': status,
				'error': error[:500],
				'timestamp': time.time(),
			}
		)
		del self._navigation_notices[:-12]

	def _navigation_observation_text(self, page: Page) -> str:
		notices = [notice for notice in self._navigation_notices if notice.get('page_id') == id(page)][-4:]
		if not notices:
			return ''
		lines = ['[Runtime navigation status]']
		for notice in notices:
			line = f'- {notice.get("status", "unknown")}: {notice.get("url", "")}'
			if notice.get('error'):
				line += f' ({notice["error"]})'
			lines.append(line)
		return '\n'.join(lines) + '\n'

	async def _stop_loading(self, page: Page) -> None:
		"""Best-effort stop for a timed-out HTML navigation while retaining its DOM."""

		with contextlib.suppress(Exception):
			await page.evaluate('window.stop()')

	async def _settle_pending_navigations(self) -> None:
		for page_id, pending in list(self._pending_navigations.items()):
			page = pending.page
			if pending.task.done():
				self._pending_navigations.pop(page_id, None)
				try:
					pending.task.result()
				except PlaywrightError as exc:
					if 'download is starting' in str(exc).casefold():
						self._record_navigation_notice(page, pending.url, 'download_started')
						# Playwright emits the Download event immediately after the
						# navigation error. Yield once through the page event loop so
						# observe() sees the newly spawned blocking task below.
						with contextlib.suppress(PlaywrightError):
							await page.wait_for_timeout(100)
					else:
						await self._stop_loading(page)
						status = 'navigation_timed_out' if isinstance(exc, PlaywrightTimeoutError) else 'navigation_failed'
						self._record_navigation_notice(page, pending.url, status, f'{type(exc).__name__}: {exc}')
				except Exception as exc:
					await self._stop_loading(page)
					self._record_navigation_notice(page, pending.url, 'navigation_failed', f'{type(exc).__name__}: {exc}')
				continue
			if (time.monotonic() - pending.started_at) * 1000 < self.navigation_timeout_ms:
				continue
			self._pending_navigations.pop(page_id, None)
			pending.task.add_done_callback(_consume_detached_task_result)
			pending.task.cancel()
			await self._stop_loading(page)
			self._record_navigation_notice(page, pending.url, 'navigation_timed_out', 'hard navigation deadline elapsed')

	async def _observe_interaction_navigation(self, page: Page) -> None:
		"""Apply the navigation observation deadlines to link/form interactions."""

		if id(page) in self._pending_navigations or page.url == _DOWNLOAD_PLACEHOLDER_URL:
			return
		try:
			await self._start_navigation(
				page,
				page.url,
				page.wait_for_load_state('domcontentloaded', timeout=self.navigation_timeout_ms),
			)
		except PlaywrightError:
			# The interaction already completed. A detached/closed target should not
			# turn the next observation into a second action failure.
			return

	async def _start_navigation(
		self, page: Page, url: str, operation: Coroutine[Any, Any, Any]
	) -> Literal['complete', 'download', 'pending']:
		"""Give one navigation operation the 10-second first-observation contract."""

		self._cancel_pending_navigation(page)
		started_at = time.monotonic()
		task = asyncio.create_task(operation)
		try:
			done, _ = await asyncio.wait({task}, timeout=self.navigation_first_observation_timeout_ms / 1000)
		except BaseException:
			if not task.done():
				task.add_done_callback(_consume_detached_task_result)
				task.cancel()
			raise
		if task in done or task.done():
			try:
				task.result()
				return 'complete'
			except PlaywrightError as exc:
				if 'download is starting' not in str(exc).casefold():
					raise
				self._record_navigation_notice(page, url, 'download_started')
				# A top-level non-HTML navigation is a blocking download. The
				# Download event itself is dispatched just after the navigation
				# error; let it enqueue before draining to terminal state.
				with contextlib.suppress(PlaywrightError):
					await page.wait_for_timeout(100)
				await self._drain_download_tasks()
				return 'download'
		self._pending_navigations[id(page)] = _PendingNavigation(page=page, url=url, started_at=started_at, task=task)
		self._record_navigation_notice(page, url, 'first_observation_ready', 'DOM content is still loading')
		return 'pending'

	async def _goto_exact(self, page: Page, url: str) -> bool:
		"""Start an exact navigation with a 10s first-observation and 120s hard deadline."""

		status = await self._start_navigation(
			page,
			url,
			page.goto(url, wait_until='domcontentloaded', timeout=self.navigation_timeout_ms),
		)
		return status == 'download'

	@staticmethod
	def _is_transient_initial_navigation_error(error: PlaywrightError) -> bool:
		"""Whether the initial document request can safely receive one retry."""

		message = str(error).casefold()
		return any(code in message for code in _TRANSIENT_INITIAL_NAVIGATION_ERROR_CODES)

	@staticmethod
	def _decision_dict(decision: AgentDecision | Mapping[str, Any]) -> dict[str, Any]:
		if isinstance(decision, Mapping):
			return dict(decision)
		action_payload = getattr(decision, 'action_payload', None)
		if callable(action_payload):
			dumped = action_payload()
			if not isinstance(dumped, Mapping):
				raise TypeError('AgentDecision.action_payload() did not return a mapping')
			return {str(key): value for key, value in dumped.items()}
		model_dump = getattr(decision, 'model_dump', None)
		if callable(model_dump):
			dumped = model_dump(exclude_none=True)
			if not isinstance(dumped, Mapping):
				raise TypeError('AgentDecision.model_dump() did not return a mapping')
			return {str(key): value for key, value in dumped.items()}
		if hasattr(decision, '__dict__'):
			return {key: value for key, value in vars(decision).items() if not key.startswith('_') and value is not None}
		raise TypeError(f'Unsupported AgentDecision value: {type(decision).__name__}')

	@staticmethod
	def _normalise_action(data: dict[str, Any]) -> tuple[str, dict[str, Any]]:
		action_value = data.get('action', data.get('name', data.get('kind')))
		params = dict(data)
		if isinstance(action_value, Mapping):
			if len(action_value) != 1:
				raise ValueError('Structured action must contain exactly one operation')
			action_value, nested = next(iter(action_value.items()))
			if isinstance(nested, Mapping):
				params.update(nested)
		elif hasattr(action_value, 'value'):
			action_value = getattr(action_value, 'value')
		if action_value is None:
			raise ValueError('AgentDecision is missing action')
		action = str(action_value).strip().lower().replace('-', '_').replace(' ', '_')
		aliases = {
			'click_xy': 'xy',
			'coordinate_click': 'xy',
			'go_back': 'back',
			'goto': 'navigate',
			'open_url': 'navigate',
			'switch_tab': 'tab',
			'new_tab': 'tab',
			'close_tab': 'tab',
			'read_element': 'read',
			'find_text': 'find',
			'network': 'inspect_network',
		}
		original = action
		action = aliases.get(action, action)
		if original == 'switch_tab':
			params.setdefault('operation', 'switch')
		elif original == 'new_tab':
			params.setdefault('operation', 'new')
		elif original == 'close_tab':
			params.setdefault('operation', 'close')
		return action, params

	@staticmethod
	def _first(values: Mapping[str, Any], *names: str, default: Any = None) -> Any:
		for name in names:
			if name in values and values[name] is not None:
				return values[name]
		return default

	@classmethod
	def _required(cls, values: Mapping[str, Any], name: str) -> Any:
		value = cls._first(values, name)
		if value is None:
			raise ValueError(f'Action requires {name}')
		return value

	@staticmethod
	def _as_bool(value: Any) -> bool:
		if isinstance(value, str):
			return value.strip().lower() in {'1', 'true', 'yes', 'on'}
		return bool(value)

	def _wait_milliseconds(self, params: Mapping[str, Any]) -> int:
		if self._first(params, 'milliseconds', 'ms') is not None:
			milliseconds = float(self._first(params, 'milliseconds', 'ms'))
		else:
			seconds = float(self._first(params, 'seconds', 'duration', 'value', default=1))
			milliseconds = seconds * 1000
		return int(min(max(milliseconds, 0), self.action_timeout_ms))

	@staticmethod
	def _validate_navigation_url(url: str) -> None:
		if not isinstance(url, str) or not url:
			raise ValueError('Navigation URL must be a non-empty string')
		if any(character.isspace() or ord(character) < 32 for character in url):
			raise ValueError('Navigation URL must not contain whitespace or control characters')
		try:
			parsed = urlsplit(url)
		except ValueError as exc:
			raise ValueError(f'Invalid navigation URL: {url!r}') from exc
		if parsed.scheme not in {'http', 'https'} or not parsed.hostname:
			raise ValueError('Navigation URL must be an absolute http(s) URL')
		if parsed.username is not None or parsed.password is not None:
			raise ValueError('Navigation URL must not contain credentials')
		try:
			_ = parsed.port
		except ValueError as exc:
			raise ValueError('Navigation URL contains an invalid port') from exc

	@staticmethod
	def _safe_step_name(step: int) -> str:
		try:
			return str(max(0, int(step)))
		except (TypeError, ValueError) as exc:
			raise ValueError(f'Invalid observation step: {step!r}') from exc

	@staticmethod
	def _content_length(headers: Mapping[str, str]) -> int | None:
		try:
			return int(headers.get('content-length', ''))
		except (TypeError, ValueError):
			return None

	@staticmethod
	def _charset(content_type: str) -> str:
		match = re.search(r'charset=([\w.-]+)', content_type)
		return match.group(1) if match else 'utf-8'

	@staticmethod
	def _is_textual_content(content_type: str, body: bytes) -> bool:
		if any(marker in content_type for marker in ('text/', 'json', 'xml', 'javascript', 'x-www-form-urlencoded')):
			return True
		return b'\x00' not in body[:1024]

	@staticmethod
	def _safe_download_name(name: str) -> str:
		name = Path(str(name).replace('\\', '/')).name
		name = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', '_', name).strip(' .')
		return name[:180] or 'download'

	def _unique_download_path(self, filename: str) -> Path:
		candidate = self.download_dir / filename
		counter = 1
		while candidate.exists() or candidate in self._reserved_download_paths:
			candidate = self.download_dir / f'{Path(filename).stem}_{counter}{Path(filename).suffix}'
			counter += 1
		self._reserved_download_paths.add(candidate)
		return candidate

	@staticmethod
	def _archive_member_parts(filename: str) -> tuple[str, ...] | None:
		assert isinstance(filename, str)
		normalised = filename.replace('\\', '/')
		if normalised.startswith('/') or re.match(r'^[A-Za-z]:', normalised):
			return None
		parts = tuple(part for part in normalised.split('/') if part not in {'', '.'})
		if not parts or any(part == '..' for part in parts):
			return None
		assert all(part not in {'', '.', '..'} for part in parts)
		return parts

	def _extract_download_archive(self, archive_path: Path) -> _ArchiveExtractionResult:
		assert archive_path.is_file()
		assert archive_path.parent.resolve(strict=True) == self.download_dir.resolve(strict=True)
		members: list[_ExtractedArchiveMember] = []
		warnings: list[str] = []
		warning_count = 0
		skipped_count = 0

		def add_warning(message: str) -> None:
			nonlocal warning_count
			warning_count += 1
			if len(warnings) < _MAX_ARCHIVE_WARNINGS:
				warnings.append(message[:500])

		if not zipfile.is_zipfile(archive_path):
			result = _ArchiveExtractionResult(
				status='failed',
				members=[],
				skipped_count=0,
				warnings=['Downloaded file has a .zip suffix but is not a valid ZIP archive'],
			)
			assert result.status == 'failed' and not result.members
			return result

		try:
			with zipfile.ZipFile(archive_path) as archive:
				directory_aliases: dict[tuple[str, ...], Path] = {}
				total_bytes = 0
				for info in archive.infolist():
					mode = info.external_attr >> 16
					file_type = stat.S_IFMT(mode)
					if info.create_system == 3 and file_type not in {0, stat.S_IFREG, stat.S_IFDIR}:
						skipped_count += 1
						add_warning(f'{info.filename}: non-regular archive member skipped')
						continue
					parts = self._archive_member_parts(info.filename)
					if parts is None:
						skipped_count += 1
						add_warning(f'{info.filename}: unsafe archive path skipped')
						continue
					target: Path | None = None
					temporary_path: Path | None = None
					try:
						is_directory = info.is_dir() or (info.create_system == 3 and file_type == stat.S_IFDIR)
						if is_directory:
							self._archive_member_directory(parts, directory_aliases)
							continue
						if len(members) >= _MAX_ARCHIVE_FILES:
							skipped_count += 1
							add_warning(f'{info.filename}: archive exceeds 1,000 file extraction limit')
							continue
						parent = self._archive_member_directory(parts[:-1], directory_aliases)
						if not parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True)):
							raise ValueError('archive member parent escaped the downloads directory')
						target = self._reserve_unique_archive_path(parent / parts[-1], kind='file')
						with tempfile.NamedTemporaryFile(
							dir=parent, prefix=f'.{target.name}.', suffix='.part', delete=False
						) as output:
							temporary_path = Path(output.name)
							with archive.open(info) as source:
								member_bytes = 0
								while chunk := source.read(1024 * 1024):
									projected_member_bytes = member_bytes + len(chunk)
									if projected_member_bytes > _MAX_ARCHIVE_MEMBER_BYTES:
										raise _ArchiveResourceLimitError('member exceeds 100 MiB extraction limit')
									projected_total_bytes = total_bytes + len(chunk)
									if projected_total_bytes > _MAX_ARCHIVE_TOTAL_BYTES:
										raise _ArchiveResourceLimitError('archive exceeds 250 MiB total extraction limit')
									written_bytes = output.write(chunk)
									member_bytes += written_bytes
									total_bytes += written_bytes
									if written_bytes != len(chunk):
										raise OSError(f'partial archive member write: {written_bytes} of {len(chunk)} bytes')
						os.link(temporary_path, target)
						temporary_path.unlink()
						temporary_path = None
						members.append(
							_ExtractedArchiveMember(
								archive_name=info.filename,
								relative_name=target.relative_to(self.download_dir).as_posix(),
								path=target,
								size_bytes=target.stat().st_size,
							)
						)
					except Exception as exc:
						if target is not None:
							self._reserved_download_paths.discard(target)
						skipped_count += 1
						add_warning(f'{info.filename}: extraction failed: {type(exc).__name__}: {exc}')
					finally:
						if temporary_path is not None:
							temporary_path.unlink(missing_ok=True)
		except Exception as exc:
			add_warning(f'ZIP extraction failed: {type(exc).__name__}: {exc}')

		if warning_count > _MAX_ARCHIVE_WARNINGS:
			omitted_count = warning_count - (_MAX_ARCHIVE_WARNINGS - 1)
			warnings[-1] = f'{omitted_count} additional warning(s) omitted'
		status: Literal['success', 'partial', 'failed']
		if warnings and not members:
			status = 'failed'
		elif warnings:
			status = 'partial'
		else:
			status = 'success'
		result = _ArchiveExtractionResult(status=status, members=members, skipped_count=skipped_count, warnings=warnings)
		assert len(result.members) <= _MAX_ARCHIVE_FILES
		assert all(
			member.path.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
			for member in result.members
		)
		return result

	def _archive_member_directory(
		self,
		parts: tuple[str, ...],
		aliases: dict[tuple[str, ...], Path],
	) -> Path:
		assert all(part not in {'', '.', '..'} for part in parts)
		assert self.download_dir.is_dir()
		parent = self.download_dir
		prefix: tuple[str, ...] = ()
		for part in parts:
			prefix += (part,)
			if prefix in aliases:
				parent = aliases[prefix]
				continue
			candidate = parent / part
			if candidate.is_symlink() or (candidate.exists() and not candidate.is_dir()):
				candidate = self._reserve_unique_archive_path(candidate, kind='directory')
				candidate.mkdir()
			else:
				candidate.mkdir(exist_ok=True)
			aliases[prefix] = candidate
			parent = candidate
		assert parent.is_dir()
		assert parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		return parent

	def _reserve_unique_archive_path(self, candidate: Path, *, kind: Literal['directory', 'file']) -> Path:
		assert candidate.parent.is_dir()
		assert candidate.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		counter = 1
		original = candidate
		while candidate.exists() or candidate.is_symlink() or candidate in self._reserved_download_paths:
			name = f'{original.name}_{counter}' if kind == 'directory' else f'{original.stem}_{counter}{original.suffix}'
			candidate = original.with_name(name)
			counter += 1
		self._reserved_download_paths.add(candidate)
		assert candidate in self._reserved_download_paths
		assert candidate.parent.resolve(strict=True).is_relative_to(self.download_dir.resolve(strict=True))
		assert not candidate.exists() and not candidate.is_symlink()
		return candidate

	@staticmethod
	def _archive_extraction_summary(result: _ArchiveExtractionResult) -> str:
		assert result.skipped_count >= 0
		summary = (
			f'ZIP extraction {result.status}: {len(result.members)} file(s) extracted, {result.skipped_count} member(s) skipped.'
		)
		if result.warnings:
			summary += '\nWarnings:\n' + '\n'.join(f'- {warning}' for warning in result.warnings)
		assert summary
		return summary

	@staticmethod
	def _document_response_extension(response: Response) -> str | None:
		try:
			if response.request.resource_type != 'document':
				return None
			if not 200 <= int(getattr(response, 'status', 200)) < 300:
				return None
			headers = response.headers
		except Exception:
			return None
		content_disposition = headers.get('content-disposition', '').lower()
		if 'attachment' in content_disposition:
			# Playwright emits a Download for attachments; using that path avoids
			# saving the same response twice.
			return None
		filename_match = re.search(r'filename\*?\s*=\s*(?:UTF-8\'\')?"?([^";]+)', content_disposition, re.IGNORECASE)
		filename = unquote(filename_match.group(1).strip()) if filename_match else ''
		content_type = headers.get('content-type', '').split(';', 1)[0].strip().lower()
		if content_type in {'text/html', 'application/html', 'application/xhtml+xml'}:
			return None
		for candidate in (filename, unquote(urlsplit(response.url).path)):
			suffix = Path(candidate).suffix.lower()
			if suffix in _DOCUMENT_EXTENSIONS:
				return suffix
		if content_type.endswith('+json'):
			return '.json'
		if content_type in _DOCUMENT_MIME_EXTENSIONS:
			return _DOCUMENT_MIME_EXTENSIONS[content_type]
		for candidate in (filename, unquote(urlsplit(response.url).path)):
			suffix = Path(candidate).suffix.lower()
			if suffix in {'.html', '.htm'}:
				return None
			if re.fullmatch(r'\.[a-z0-9]{1,16}', suffix):
				return suffix
		return '.bin'

	@classmethod
	def _response_filename(cls, response: Response, extension: str) -> str:
		disposition = response.headers.get('content-disposition', '')
		match = re.search(r"filename\*\s*=\s*(?:UTF-8'')?([^;]+)", disposition, re.IGNORECASE)
		if match is None:
			match = re.search(r'filename\s*=\s*"?([^";]+)', disposition, re.IGNORECASE)
		if match is not None:
			name = unquote(match.group(1).strip().strip('"\''))
		else:
			name = unquote(Path(urlsplit(response.url).path).name)
		name = cls._safe_download_name(name or f'document{extension}')
		if Path(name).suffix.lower() != extension:
			name = f'{name}{extension}'
		return name

	@staticmethod
	def _extract_download_text(path: Path) -> tuple[str, bool]:
		if path.stat().st_size > _MAX_DOWNLOAD_BYTES:
			return f'[text extraction omitted: file exceeds {_MAX_DOWNLOAD_BYTES} bytes]', True
		extension = path.suffix.lower()
		text = ''
		if extension == '.pdf':
			from pypdf import PdfReader

			reader = PdfReader(str(path))
			text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
		elif extension == '.docx':
			from docx import Document

			document = Document(str(path))
			parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
			for table in document.tables:
				parts.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
			text = '\n'.join(parts)
		elif extension == '.xlsx':
			text = BrowserRuntime._extract_xlsx_text(path)
		elif extension == '.xls':
			text = BrowserRuntime._extract_xls_text(path)
		elif extension == '.zip':
			text = BrowserRuntime._extract_zip_text(path)
		elif extension in {'.txt', '.text', '.md', '.csv', '.tsv', '.json', '.jsonl', '.xml', '.html', '.htm'}:
			raw = path.read_bytes()
			text = raw.decode('utf-8-sig', errors='replace')
			if extension in {'.html', '.htm'}:
				parser = _VisibleTextParser()
				parser.feed(text)
				text = '\n'.join(parser.parts)
		else:
			return '', False
		truncated = len(text) > _MAX_DOWNLOAD_TEXT
		return text[:_MAX_DOWNLOAD_TEXT], truncated

	@staticmethod
	def _extract_zip_text(path: Path) -> str:
		"""Read bounded text/CSV members from an official data ZIP in memory."""
		parts: list[str] = []
		used = 0
		supported = {'.txt', '.text', '.md', '.csv', '.tsv', '.json', '.jsonl', '.xml'}
		with zipfile.ZipFile(path) as archive:
			members = sorted((item for item in archive.infolist() if not item.is_dir()), key=lambda item: item.filename)
			for item in members[:100]:
				if Path(item.filename).suffix.lower() not in supported or item.file_size > _MAX_DOWNLOAD_BYTES:
					continue
				remaining = _MAX_DOWNLOAD_TEXT - used
				if remaining <= 0:
					break
				with archive.open(item) as member:
					raw = member.read(remaining + 1)
				text = raw.decode('utf-8-sig', errors='replace')
				chunk = f'[{item.filename}]\n{text}'
				parts.append(chunk[:remaining])
				used += min(len(chunk), remaining)
		return '\n'.join(parts)

	@staticmethod
	def _extract_xls_text(path: Path) -> str:
		"""Extract legacy BIFF Excel sheets while preserving row columns."""
		import xlrd

		workbook = xlrd.open_workbook(str(path), on_demand=True)
		lines: list[str] = []
		try:
			for sheet in workbook.sheets():
				lines.append(f'[{sheet.name}]')
				for row_index in range(sheet.nrows):
					values: list[str] = []
					for column_index in range(sheet.ncols):
						cell = sheet.cell(row_index, column_index)
						if cell.ctype == xlrd.XL_CELL_DATE:
							value = xlrd.xldate_as_datetime(float(cell.value), workbook.datemode).isoformat(sep=' ')
						elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
							value = str(int(float(cell.value)))
						else:
							value = str(cell.value)
						values.append(value)
					while values and not values[-1]:
						values.pop()
					lines.append('\t'.join(values))
		finally:
			workbook.release_resources()
		return '\n'.join(lines)

	@staticmethod
	def _extract_xlsx_text(path: Path) -> str:
		namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
		with zipfile.ZipFile(path) as archive:
			shared: list[str] = []
			if 'xl/sharedStrings.xml' in archive.namelist():
				root = ElementTree.fromstring(archive.read('xl/sharedStrings.xml'))
				for item in root.findall(f'{namespace}si'):
					shared.append(''.join(node.text or '' for node in item.iter(f'{namespace}t')))
			lines: list[str] = []
			worksheets = sorted(name for name in archive.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name))
			for worksheet in worksheets:
				lines.append(f'[{Path(worksheet).stem}]')
				root = ElementTree.fromstring(archive.read(worksheet))
				for row in root.iter(f'{namespace}row'):
					values: list[str] = []
					for cell in row.findall(f'{namespace}c'):
						cell_type = cell.get('t')
						value_node = cell.find(f'{namespace}v')
						if cell_type == 'inlineStr':
							inline = cell.find(f'{namespace}is')
							value = (
								''.join(node.text or '' for node in inline.iter(f'{namespace}t')) if inline is not None else ''
							)
						else:
							value = value_node.text if value_node is not None and value_node.text else ''
							if cell_type == 's' and value.isdigit() and int(value) < len(shared):
								value = shared[int(value)]
						values.append(value)
					lines.append('\t'.join(values))
			return '\n'.join(lines)
